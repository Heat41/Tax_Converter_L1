import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from core.reconciliation import ReconciliationResult
from core.finalization import FinalizationInput
from core.legacy_1770_hybrid_l2 import LegacyLampiranIIXlsxRenderer
from core.worksheet_pph_state import WorksheetBupotRow


def _analysis():
    return ReconciliationResult(
        total_harta_sebelumnya=0.0,
        total_harta_berjalan=0.0,
        total_utang_sebelumnya=0.0,
        total_utang_berjalan=0.0,
        naik_turun_harta_utang=0.0,
        biaya_hidup_setahun=0.0,
        pajak_pajak=0.0,
        pengeluaran_lain_lain=0.0,
        kerugian_keuntungan_penjualan_aset=0.0,
        utang_baru_atas_kredit=0.0,
        harta_baru_dari_kredit=0.0,
        total_pengeluaran=0.0,
        penghasilan_bruto_umkm=0.0,
        penambahan_penghasilan_bruto_umkm=0.0,
        margin_usaha=0.0,
        jumlah_penghasilan_bruto_final=0.0,
        jumlah_penghasilan_bukan_objek=0.0,
        penghasilan_netto=0.0,
        selisih_pengeluaran_vs_penghasilan=0.0,
    )


def _input():
    return FinalizationInput(
        npwp="3275017007900006",
        nama_wp="Lisa Vinatalia",
        tahun_pajak=2025,
        harta_current_rows=[],
        harta_original_hash="",
        bupot_rows=[
            WorksheetBupotRow(
                jenis="BP21",
                no_bupot="2501LZ5NR",
                tahun="2025",
                jenis_pph="Pasal 21",
                pph_dipotong=46519.0,
                npwp_pemotong="0013058839038000",
                nama_pemotong="STAR COSMOS",
                tanggal_pemotongan="31/03/2025",
                npwp_pemberi_kerja="0013058839038000",
            ),
            WorksheetBupotRow(
                jenis="BP21",
                no_bupot="2503ADYB7",
                tahun="2025",
                jenis_pph="Pasal 21",
                pph_dipotong=21350.0,
                npwp_pemotong="0018767327123000",
                nama_pemotong="HOKINDA CITRALESTARI",
                tanggal_pemotongan="30/06/2025",
                npwp_pemberi_kerja="0018767327123000",
            ),
        ],
        pph_components={},
        umkm_state={},
        penghasilan_lainnya={},
        zakat=0.0,
        status_ptkp="TK/0",
        pph_calc_result={},
        analisis_result=_analysis(),
    )


class TestLegacyLampiranIIXlsx(unittest.TestCase):
    def test_renderer_maps_bupot_into_1770_ii_legacy_table(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "lampiran_ii.xlsx"
            wb = Workbook()
            ws = wb.active
            ws.title = "04 Legacy Lamp II"
            LegacyLampiranIIXlsxRenderer().render(ws, _input())
            wb.save(path)

            check = load_workbook(path, data_only=False)["04 Legacy Lamp II"]
            text = " ".join(
                str(check.cell(row, col).value or "")
                for row in range(1, check.max_row + 1)
                for col in range(1, check.max_column + 1)
            )

            self.assertIn("1770 - II", text)
            self.assertIn("LAMPIRAN - II", text)
            self.assertIn("STAR COSMOS", text)
            self.assertIn("0013058839038000", text)
            self.assertIn("2501LZ5NR", text)
            self.assertIn("31/03/2025", text)
            self.assertIn("PPh Pasal 21", text)
            self.assertIn("HOKINDA CITRALESTARI", text)
            self.assertIn("2503ADYB7", text)
            self.assertIn("JUMLAH BAGIAN A", text)
            total_cell = next(
                check.cell(row, 9)
                for row in range(1, check.max_row + 1)
                if check.cell(row, 1).value == "JUMLAH BAGIAN A"
            )
            self.assertEqual(total_cell.value, "=SUM(I14:I28)")
            self.assertEqual(str(check.page_setup.paperSize), str(check.PAPERSIZE_LEGAL))
            self.assertEqual(check.page_setup.orientation, "portrait")
            self.assertEqual(check.page_setup.fitToWidth, 1)
            self.assertEqual(check.page_setup.fitToHeight, 1)


if __name__ == "__main__":
    unittest.main()
