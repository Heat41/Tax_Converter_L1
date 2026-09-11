from pathlib import Path

from openpyxl import load_workbook

from core.reverse_coretax_excel import ReverseCoretaxExcelExporter
from core.reverse_coretax_mapping import (
    CATEGORY_ORDER,
    ReverseCoretaxPackage,
    ReverseCoretaxRow,
)


def _row(category: str, code: str, no: int = 1):
    return ReverseCoretaxRow(
        nomor=no,
        kategori=category,
        kode_harta=code,
        nama_harta=f"HARTA {category}",
        tahun_perolehan=2025,
        nilai=1_250_000,
        nomor_akun_keterangan=f"DETAIL-{category}",
        atas_nama="WP TEST",
        nama_bank="BANK TEST",
        source_eform_code="011",
    )


def _package():
    package = ReverseCoretaxPackage(
        npwp="1111222233334444",
        nama_wp="WP TEST",
        tahun_pajak=2025,
        revision=1,
        snapshot_hash="abc",
    )
    rows = {
        "KAS": _row("KAS", "0101"),
        "PIUTANG": _row("PIUTANG", "0201"),
        "INVESTASI": _row("INVESTASI", "0301"),
        "BERGERAK": _row("BERGERAK", "0401"),
        "HTB": _row("HTB", "0501"),
        "LAINNYA": _row("LAINNYA", "0601"),
    }
    for category, row in rows.items():
        package.rows_by_category[category].append(row)
    return package


def test_exporter_creates_exactly_six_xlsx_files(tmp_path):
    result = ReverseCoretaxExcelExporter().export_package(
        _package(),
        tmp_path,
    )

    assert result.ok is True
    assert set(result.files) == set(CATEGORY_ORDER)
    assert len(list(tmp_path.glob("*.xlsx"))) == 6
    assert all(result.row_counts[category] == 1 for category in CATEGORY_ORDER)


def test_exported_headers_match_internal_import_contract(tmp_path):
    result = ReverseCoretaxExcelExporter().export_package(_package(), tmp_path)

    expected = {
        "KAS": ["kode_harta", "nama_harta", "tahun_perolehan", "harga_perolehan", "keterangan"],
        "PIUTANG": ["kode_harta", "nama_peminjam", "npwp_peminjam", "tahun_perolehan", "harga_perolehan", "keterangan"],
        "INVESTASI": ["kode_harta", "nama_harta", "penerbit_saham", "tahun_perolehan", "harga_perolehan", "keterangan"],
        "BERGERAK": ["kode_harta", "nama_harta", "merek_type", "tahun_perolehan", "harga_perolehan", "keterangan"],
        "HTB": ["kode_harta", "jenis_harta", "lokasi_alamat", "tahun_perolehan", "harga_perolehan", "keterangan"],
        "LAINNYA": ["kode_harta", "nama_harta", "tahun_perolehan", "harga_perolehan", "keterangan"],
    }

    for category in CATEGORY_ORDER:
        wb = load_workbook(result.files[category], read_only=True, data_only=True)
        ws = wb.active
        headers = [cell.value for cell in ws[1]]
        assert headers == expected[category]
        wb.close()


def test_empty_category_still_gets_valid_header_only_file(tmp_path):
    package = _package()
    package.rows_by_category["PIUTANG"] = []

    result = ReverseCoretaxExcelExporter().export_package(package, tmp_path)

    assert result.ok is True
    wb = load_workbook(result.files["PIUTANG"], read_only=True, data_only=True)
    ws = wb.active
    assert ws.max_row == 1
    assert [cell.value for cell in ws[1]][0] == "kode_harta"
    wb.close()


def test_numeric_value_is_exported_as_number_not_text(tmp_path):
    result = ReverseCoretaxExcelExporter().export_package(_package(), tmp_path)

    wb = load_workbook(result.files["KAS"], read_only=True, data_only=True)
    ws = wb.active
    headers = [cell.value for cell in ws[1]]
    value_column = headers.index("harga_perolehan") + 1
    assert ws.cell(2, value_column).value == 1_250_000
    assert isinstance(ws.cell(2, value_column).value, (int, float))
    wb.close()
