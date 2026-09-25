import unittest

from openpyxl import Workbook

from core.finalization import FinalizationInput
from core.legacy_1770_hybrid_l3 import LegacyLampiranIIIXlsxRenderer
from core.reconciliation import ReconciliationResult


def _analysis(non_object_total=0.0):
    return ReconciliationResult(
        total_harta_sebelumnya=0.0,
        total_harta_berjalan=0.0,
        total_utang_sebelumnya=0.0,
        total_utang_berjalan=0.0,
        naik_turun_harta_utang=0.0,
        biaya_hidup_setahun=0.0,
        pajak_pajak=0.0,
        pengeluaran_lain_lain=0.0,
        kerugian_keuntungan_penjualan_aset=0.0,
        utang_baru_atas_kredit=0.0,
        harta_baru_dari_kredit=0.0,
        total_pengeluaran=0.0,
        penghasilan_bruto_umkm=0.0,
        penambahan_penghasilan_bruto_umkm=0.0,
        margin_usaha=0.0,
        jumlah_penghasilan_bruto_final=0.0,
        jumlah_penghasilan_bukan_objek=non_object_total,
        penghasilan_netto=0.0,
        selisih_pengeluaran_vs_penghasilan=0.0,
    )


def _input(non_object_rows=None, aggregate=0.0):
    return FinalizationInput(
        npwp="3275017007900006",
        nama_wp="Lisa Vinatalia",
        tahun_pajak=2025,
        harta_current_rows=[],
        harta_original_hash="",
        bupot_rows=[],
        pph_components={},
        umkm_state={},
        penghasilan_lainnya={"non_object_rows": non_object_rows or []},
        zakat=0.0,
        status_ptkp="TK/0",
        pph_calc_result={},
        analisis_result=_analysis(aggregate),
    )


class TestLegacyLampiranIIIStage8E4C(unittest.TestCase):
    def _render(self, data):
        wb = Workbook()
        ws = wb.active
        LegacyLampiranIIIXlsxRenderer().render(ws, data)
        return ws

    def _bagian_b_value(self, ws, no):
        label = LegacyLampiranIIIXlsxRenderer.NON_OBJECT_LABELS[no - 1]
        row = next(
            r for r in range(1, ws.max_row + 1)
            if ws.cell(r, 2).value == label
        )
        return ws.cell(row, 8).value

    def test_detail_non_object_rows_map_without_double_counting_aggregate(self):
        ws = self._render(
            _input(
                [
                    {"keterangan": "Bantuan/Sumbangan", "dpp": 2_000_000},
                    {"keterangan": "Warisan", "dpp": 3_000_000},
                ],
                aggregate=6_500_000,
            )
        )
        self.assertEqual(self._bagian_b_value(ws, 1), 2_000_000)
        self.assertEqual(self._bagian_b_value(ws, 2), 3_000_000)
        self.assertEqual(self._bagian_b_value(ws, 6), 1_500_000)

    def test_ambiguous_hibah_warisan_stays_in_other_non_object_row(self):
        ws = self._render(
            _input(
                [{"keterangan": "HIBAH / WARISAN", "dpp": 4_000_000}],
                aggregate=4_000_000,
            )
        )
        self.assertEqual(self._bagian_b_value(ws, 1), 0)
        self.assertEqual(self._bagian_b_value(ws, 2), 0)
        self.assertEqual(self._bagian_b_value(ws, 6), 4_000_000)

    def test_aggregate_only_keeps_existing_fallback(self):
        ws = self._render(_input([], aggregate=6_832_000))
        self.assertEqual(self._bagian_b_value(ws, 6), 6_832_000)


if __name__ == "__main__":
    unittest.main()
