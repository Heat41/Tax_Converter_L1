from pathlib import Path

from openpyxl import Workbook, load_workbook

from core.coretax_official_schema import OFFICIAL_CORETAX_SCHEMAS
from core.reverse_coretax_mapping import ReverseCoretaxPackage, ReverseCoretaxRow
from core.reverse_coretax_official_excel import (
    DATA_START_ROW,
    OfficialCoretaxExcelExporter,
)


CATEGORIES = (
    "KAS",
    "PIUTANG",
    "INVESTASI",
    "BERGERAK",
    "HTB",
    "LAINNYA",
)


def _create_template(directory: Path, category: str) -> Path:
    schema = OFFICIAL_CORETAX_SCHEMAS[category]
    path = directory / f"{schema.excel_filename_hint}.xlsx"

    wb = Workbook()
    ws = wb.active
    ws.title = "DATA"

    ws["A1"] = "NPWP"
    ws["B1"] = ""
    ws["A2"] = "Tahun Pajak"
    ws["B2"] = ""

    for column_index, header in enumerate(schema.excel_headers, start=1):
        ws.cell(3, column_index).value = header

    ws.column_dimensions["A"].width = 27
    wb.save(path)
    return path


def _create_all_templates(directory: Path):
    directory.mkdir(parents=True, exist_ok=True)
    for category in CATEGORIES:
        _create_template(directory, category)


def _package():
    package = ReverseCoretaxPackage(
        npwp="6101015612710001",
        nama_wp="EVY BACHTIAR",
        tahun_pajak=2025,
        revision=1,
    )

    package.rows_by_category["KAS"].append(
        ReverseCoretaxRow(
            nomor=1,
            kategori="KAS",
            kode_harta="0102",
            nama_harta="Tabungan",
            tahun_perolehan=2025,
            nilai=391742658,
            nomor_akun_keterangan="116801005138508",
            atas_nama="DR EVY BACHTIAR SPOG",
            nama_bank="BRI",
            official_metadata={
                "code": "0102",
                "account_number": "116801005138508",
                "account_on_behalf_of": "DR EVY BACHTIAR SPOG",
                "bank_name": "BRI",
                "country": "Indonesia",
                "year": 2025,
                "balance": 391742658,
                "remarks": "",
            },
        )
    )

    return package


def test_stage8d4c_exports_only_populated_category_files(tmp_path):
    templates = tmp_path / "templates"
    output = tmp_path / "output"
    _create_all_templates(templates)

    result = OfficialCoretaxExcelExporter(templates).export_package(
        _package(),
        output,
    )

    assert result.ok
    assert set(result.files) == {"KAS"}
    assert result.row_counts == {"KAS": 1}
    assert result.files["KAS"].exists()


def test_stage8d4c_preserves_official_sheet_and_headers(tmp_path):
    templates = tmp_path / "templates"
    output = tmp_path / "output"
    _create_all_templates(templates)

    result = OfficialCoretaxExcelExporter(templates).export_package(
        _package(),
        output,
    )

    wb = load_workbook(result.files["KAS"])
    ws = wb["DATA"]
    schema = OFFICIAL_CORETAX_SCHEMAS["KAS"]

    headers = tuple(
        ws.cell(3, column).value
        for column in range(1, len(schema.excel_headers) + 1)
    )

    assert headers == schema.excel_headers
    wb.close()


def test_stage8d4c_writes_npwp_and_tax_year_metadata(tmp_path):
    templates = tmp_path / "templates"
    output = tmp_path / "output"
    _create_all_templates(templates)

    result = OfficialCoretaxExcelExporter(templates).export_package(
        _package(),
        output,
    )

    wb = load_workbook(result.files["KAS"])
    ws = wb["DATA"]

    assert ws["B1"].value == "6101015612710001"
    assert ws["B2"].value == 2025

    wb.close()


def test_stage8d4c_cash_row_matches_official_schema(tmp_path):
    templates = tmp_path / "templates"
    output = tmp_path / "output"
    _create_all_templates(templates)

    result = OfficialCoretaxExcelExporter(templates).export_package(
        _package(),
        output,
    )

    wb = load_workbook(result.files["KAS"])
    ws = wb["DATA"]

    values = tuple(
        ws.cell(DATA_START_ROW, column).value
        for column in range(1, 9)
    )

    assert values == (
        "0102",
        "116801005138508",
        "DR EVY BACHTIAR SPOG",
        "BRI",
        "Indonesia",
        2025,
        391742658,
        None,
    )

    wb.close()


def test_stage8d4c_does_not_modify_source_template(tmp_path):
    templates = tmp_path / "templates"
    output = tmp_path / "output"
    _create_all_templates(templates)

    source = next(path for path in templates.glob("*Kas Setara Kas*.xlsx"))
    original_bytes = source.read_bytes()

    OfficialCoretaxExcelExporter(templates).export_package(
        _package(),
        output,
    )

    assert source.read_bytes() == original_bytes


def test_stage8d4c_preserves_template_layout(tmp_path):
    templates = tmp_path / "templates"
    output = tmp_path / "output"
    _create_all_templates(templates)

    result = OfficialCoretaxExcelExporter(templates).export_package(
        _package(),
        output,
    )

    wb = load_workbook(result.files["KAS"])
    ws = wb["DATA"]

    assert ws.column_dimensions["A"].width == 27

    wb.close()


def test_stage8d4c_requires_template_only_for_populated_categories(tmp_path):
    templates = tmp_path / "templates"
    output = tmp_path / "output"
    templates.mkdir(parents=True, exist_ok=True)
    _create_template(templates, "KAS")

    result = OfficialCoretaxExcelExporter(templates).export_package(
        _package(),
        output,
    )

    assert result.ok
    assert set(result.files) == {"KAS"}


def test_stage8d4h_missing_category_specific_values_stay_blank(tmp_path):
    templates = tmp_path / "templates"
    output = tmp_path / "output"
    templates.mkdir(parents=True, exist_ok=True)
    _create_template(templates, "INVESTASI")

    package = ReverseCoretaxPackage(
        npwp="6101015612710001",
        nama_wp="EVY BACHTIAR",
        tahun_pajak=2025,
        revision=1,
    )
    package.rows_by_category["INVESTASI"].append(
        ReverseCoretaxRow(
            nomor=1,
            kategori="INVESTASI",
            kode_harta="0305",
            nama_harta="Investasi",
            tahun_perolehan=2025,
            nilai=250000000,
            nomor_akun_keterangan="ACC-01",
            atas_nama="EVY BACHTIAR",
            nama_bank="BANK CONTOH",
            official_metadata={
                "country": "Indonesia",
                "institution_name": "BANK CONTOH",
                "account_number": "ACC-01",
                "year": 2025,
                "current_balance": 250000000,
            },
        )
    )

    result = OfficialCoretaxExcelExporter(templates).export_package(
        package,
        output,
    )

    assert result.ok
    wb = load_workbook(result.files["INVESTASI"])
    ws = wb["DATA"]

    # Biaya Perolehan tidak tersedia di sumber. Jangan menyalin Nilai Saat Ini
    # ke kolom tersebut hanya untuk membuat baris tampak lengkap.
    assert ws.cell(DATA_START_ROW, 6).value is None
    assert ws.cell(DATA_START_ROW, 8).value == 250000000

    wb.close()
