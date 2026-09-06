import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from config.constants import KategoriL1
from core.exporters.simulasi_harta_writer import SimulasiHartaExcelWriter
from core.mapping.eform_reference import find_eform_reference
from core.mapping.worksheet_harta_mapper import WorksheetHartaMapper
from core.models.asset import HartaL1Item


class TestRealWorksheetStage3F(unittest.TestCase):
    def setUp(self):
        self.mapper = WorksheetHartaMapper()

    def test_reference_maps_ct_to_eform(self):
        self.assertEqual(find_eform_reference("0102").kode_eform, "012")
        self.assertEqual(find_eform_reference("0302").kode_eform, "032")
        self.assertEqual(find_eform_reference("0701").kode_eform, "051")
        self.assertEqual(find_eform_reference("0501").kode_eform, "061")
        self.assertEqual(find_eform_reference("0712").kode_eform, "019")

    def test_ambiguous_0499_uses_description_hint(self):
        self.assertEqual(find_eform_reference("0499", description_hint="kendaraan lain").kode_eform, "049")
        self.assertEqual(find_eform_reference("0499", description_hint="hewan ternak").kode_eform, "059")

    def test_kas_matches_real_worksheet_columns(self):
        item = HartaL1Item(
            wp_id=1,
            kategori_l1=KategoriL1.KAS,
            kode_harta="0102",
            tahun_perolehan=2024,
            nomor_akun="0668222229",
            atas_nama="LISA VINATALIA",
            nama_bank_institusi="BCA",
            saldo_current=1521825000,
        )
        rows = self.mapper.map_items([item])
        row = rows[0]
        self.assertEqual(row.kode_eform, "012")
        self.assertEqual(row.kode_ct, "0102")
        self.assertEqual(row.nama_harta, "Tabungan (Bank/Lembaga Keuangan)")
        self.assertEqual(row.nomor_akun_keterangan, "0668222229")
        self.assertEqual(row.atas_nama, "LISA VINATALIA")
        self.assertEqual(row.nama_bank, "BCA")
        self.assertEqual(row.nilai_tahun_berjalan, 1521825000.0)

    def test_previous_year_item_is_matched(self):
        old = HartaL1Item(
            wp_id=1,
            kategori_l1=KategoriL1.KAS,
            kode_harta="0102",
            tahun_perolehan=2024,
            nomor_akun="0668222229",
            saldo_current=75211685,
        )
        current = HartaL1Item(
            wp_id=1,
            kategori_l1=KategoriL1.KAS,
            kode_harta="0102",
            tahun_perolehan=2024,
            nomor_akun="0668222229",
            saldo_current=1521825000,
        )
        row = self.mapper.map_items([current], previous_items=[old])[0]
        self.assertEqual(row.nilai_tahun_sebelumnya, 75211685.0)
        self.assertEqual(row.nilai_tahun_berjalan, 1521825000.0)

    def test_unknown_ct_code_is_error(self):
        item = HartaL1Item(
            wp_id=1,
            kategori_l1=KategoriL1.LAINNYA,
            kode_harta="9999",
            tahun_perolehan=2025,
            biaya_perolehan_current=1000,
        )
        with self.assertRaises(ValueError):
            self.mapper.map_items([item])


class TestSimulasiHartaWriterStage3F(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        root = Path(self.temp.name)
        self.template = root / "template.xlsx"
        self.output = root / "output.xlsx"
        self._create_template()

    def tearDown(self):
        self.temp.cleanup()

    def _create_template(self):
        wb = Workbook()
        ws = wb.active
        ws.title = "SIMULASI I"
        headers = [
            "NO", "KODE EFORM", "KODE CT", "NAMA HARTA",
            "NOMOR AKUN / KETERANGAN", "ATAS NAMA", "NAMA BANK", "TH PEROLEHAN",
            2024, 2025, None, "Kas Setara Kas", "Piutang", "Investasi",
            "Harta Bergerak", "Harta Tidak Bergerak", "Harta Lainnya",
        ]
        for col, value in enumerate(headers, start=1):
            ws.cell(20, col).value = value
        for row in range(21, 24):
            ws.cell(row, 1).value = row - 20
        ws.cell(24, 1).value = "Grand Total"
        ws.cell(24, 9).value = "=SUBTOTAL(9,I21:I23)"
        ws.cell(24, 10).value = "=SUBTOTAL(9,J21:J23)"
        wb.create_sheet("2025")["A1"] = "TETAP"
        wb.save(self.template)

    def _mapped_rows(self):
        mapper = WorksheetHartaMapper()
        items = [
            HartaL1Item(
                wp_id=1, kategori_l1=KategoriL1.KAS, kode_harta="0102",
                tahun_perolehan=2024, nomor_akun="0668222229", atas_nama="LISA VINATALIA",
                nama_bank_institusi="BCA", saldo_current=1521825000,
            ),
            HartaL1Item(
                wp_id=1, kategori_l1=KategoriL1.INVESTASI, kode_harta="0302",
                tahun_perolehan=2019, nomor_akun_bukti="AKTA - No. 34",
                nama_institusi="PT. Maximus Graha Sapta", biaya_perolehan_current=500000000,
            ),
        ]
        return mapper.map_items(items)

    def test_writer_matches_simulasi_layout(self):
        writer = SimulasiHartaExcelWriter()
        result = writer.write(self.template, self.output, self._mapped_rows(), current_year=2025)
        self.assertEqual(result.header_row, 20)
        self.assertEqual(result.start_row, 21)
        self.assertEqual(result.written_rows, 2)

        wb = load_workbook(self.output, data_only=False)
        ws = wb["SIMULASI I"]
        self.assertEqual(ws["B21"].value, "012")
        self.assertEqual(ws["C21"].value, "0102")
        self.assertEqual(ws["D21"].value, "Tabungan (Bank/Lembaga Keuangan)")
        self.assertEqual(ws["E21"].value, "0668222229")
        self.assertEqual(ws["F21"].value, "LISA VINATALIA")
        self.assertEqual(ws["G21"].value, "BCA")
        self.assertEqual(ws["I20"].value, 2024)
        self.assertEqual(ws["J20"].value, 2025)

    def test_category_formulas_are_rebuilt_without_ref_errors(self):
        writer = SimulasiHartaExcelWriter()
        writer.write(self.template, self.output, self._mapped_rows(), current_year=2025)
        wb = load_workbook(self.output, data_only=False)
        ws = wb["SIMULASI I"]
        self.assertIn('LEFT($C21,2)="01"', ws["L21"].value)
        self.assertIn('LEFT($C21,2)="02"', ws["M21"].value)
        self.assertNotIn("#REF!", ws["L21"].value)
        self.assertNotIn("#REF!", ws["M21"].value)

    def test_grand_total_and_other_sheet_are_preserved(self):
        writer = SimulasiHartaExcelWriter()
        writer.write(self.template, self.output, self._mapped_rows(), current_year=2025)
        wb = load_workbook(self.output, data_only=False)
        ws = wb["SIMULASI I"]
        self.assertEqual(ws["A24"].value, "Grand Total")
        self.assertEqual(ws["I24"].value, "=SUBTOTAL(9,I21:I23)")
        self.assertEqual(wb["2025"]["A1"].value, "TETAP")


if __name__ == "__main__":
    unittest.main()
