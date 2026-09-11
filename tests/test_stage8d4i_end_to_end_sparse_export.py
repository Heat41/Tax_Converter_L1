import json
from pathlib import Path

from openpyxl import Workbook, load_workbook

from config.database import init_database
from core.coretax_official_schema import OFFICIAL_CORETAX_SCHEMAS
from core.evy_reconciliation import EvyReconciliationResult
from core.finalization import FinalizationInput, FinalizationService
from core.mapping.worksheet_harta_mapper import WorksheetHartaRow
from core.reverse_coretax_official_package import OfficialCoretaxPackageExporter
from core.reverse_coretax_official_package_validator import (
    OfficialCoretaxPackageValidator,
)


ACTIVE_CATEGORIES = ("KAS", "HTB", "LAINNYA")


def _analysis():
    return EvyReconciliationResult(
        total_harta_sebelumnya=0.0,
        total_harta_berjalan=579_742_658.0,
        total_utang_sebelumnya=0.0,
        total_utang_berjalan=0.0,
        naik_turun_harta_utang=579_742_658.0,
        biaya_hidup_setahun=0.0,
        pajak_pajak=0.0,
        pengeluaran_lain_lain=0.0,
        kerugian_keuntungan_penjualan_aset=0.0,
        utang_baru_atas_kredit=0.0,
        harta_baru_dari_kredit=0.0,
        total_pengeluaran=0.0,
        penghasilan_bruto_umkm=0.0,
        penambahan_penghasilan_bruto_umkm=0.0,
        margin_usaha=0.0,
        jumlah_penghasilan_bruto_final=0.0,
        jumlah_penghasilan_bukan_objek=0.0,
        penghasilan_netto=0.0,
        selisih_pengeluaran_vs_penghasilan=0.0,
    )


def _rows():
    return [
        WorksheetHartaRow(
            nomor=1,
            kode_eform="011",
            kode_ct="0102",
            nama_harta="Tabungan",
            nomor_akun_keterangan="116801005138508",
            atas_nama="DR EVY BACHTIAR SPOG",
            nama_bank="BRI",
            tahun_perolehan=2025,
            nilai_tahun_sebelumnya=0.0,
            nilai_tahun_berjalan=391_742_658.0,
            coretax_metadata={
                "category": "KAS",
                "code": "0102",
                "account_number": "116801005138508",
                "account_on_behalf_of": "DR EVY BACHTIAR SPOG",
                "bank_name": "BRI",
                "country": "Indonesia",
                "year": 2025,
                "balance": 391_742_658.0,
                "remarks": "",
            },
        ),
        WorksheetHartaRow(
            nomor=2,
            kode_eform="061",
            kode_ct="0501",
            nama_harta="Tanah",
            nomor_akun_keterangan="SHM-001; Pontianak",
            atas_nama="-",
            nama_bank="-",
            tahun_perolehan=2021,
            nilai_tahun_sebelumnya=0.0,
            nilai_tahun_berjalan=88_000_000.0,
            coretax_metadata={
                "category": "HTB",
                "code": "0501",
                "location_of_asset": "Pontianak",
                "property_size_land": "120 M2",
                "property_size_building": "",
                "source_of_ownership": "Own Income",
                "certificate_number": "SHM-001",
                "year": 2021,
                # cost_of_acquisition sengaja tidak tersedia.
                "fair_market_value": 88_000_000.0,
                "remarks": "",
            },
        ),
        WorksheetHartaRow(
            nomor=3,
            kode_eform="069",
            kode_ct="0601",
            nama_harta="Harta Lainnya",
            nomor_akun_keterangan="DOC-001",
            atas_nama="-",
            nama_bank="-",
            tahun_perolehan=2024,
            nilai_tahun_sebelumnya=0.0,
            nilai_tahun_berjalan=100_000_000.0,
            coretax_metadata={
                "category": "LAINNYA",
                "code": "0601",
                "year": 2024,
                "account_number": "DOC-001",
                "additional_information": "Data yang tersedia dari WP",
                # cost_of_acquisition sengaja tidak tersedia.
                "current_value": 100_000_000.0,
                "remarks": "",
            },
        ),
    ]


def _finalization_input():
    return FinalizationInput(
        npwp="6101015612710001",
        nama_wp="EVY BACHTIAR",
        tahun_pajak=2025,
        harta_current_rows=_rows(),
        harta_original_hash="e2e-source-hash",
        bupot_rows=[],
        pph_components={},
        umkm_state={},
        penghasilan_lainnya={},
        zakat=0.0,
        status_ptkp="TK/0",
        pph_calc_result={},
        analisis_result=_analysis(),
    )


def _create_active_templates(directory: Path):
    directory.mkdir(parents=True, exist_ok=True)

    for category in ACTIVE_CATEGORIES:
        schema = OFFICIAL_CORETAX_SCHEMAS[category]
        path = directory / f"{schema.excel_filename_hint}.xlsx"

        wb = Workbook()
        ws = wb.active
        ws.title = schema.excel_sheet
        ws["A1"] = "NPWP"
        ws["B1"] = ""
        ws["A2"] = "Tahun Pajak"
        ws["B2"] = ""

        for index, header in enumerate(schema.excel_headers, start=1):
            ws.cell(3, index).value = header

        wb.save(path)


def test_stage8d4i_end_to_end_sparse_final_to_validated_official_package(tmp_path):
    db_path = tmp_path / "stage8d4i.db"
    init_database(db_path)

    finalization = FinalizationService(db_path=db_path)
    final_result = finalization.finalize(_finalization_input())

    assert final_result.success
    assert final_result.revision == 1
    assert final_result.snapshot_hash

    templates = tmp_path / "templates"
    output = tmp_path / "official-package"
    _create_active_templates(templates)

    export_result = OfficialCoretaxPackageExporter(
        templates
    ).export_active_final(
        "6101015612710001",
        2025,
        output,
        db_path=db_path,
    )

    assert export_result.ok
    assert set(export_result.excel_result.files) == set(ACTIVE_CATEGORIES)
    assert set(export_result.xml_result.files) == {"KAS", "HTB"}
    assert export_result.xml_result.unsupported_categories == ["LAINNYA"]

    validation = OfficialCoretaxPackageValidator().validate(output)

    assert validation.ok
    assert not validation.errors

    manifest = json.loads(
        (output / "manifest.json").read_text(encoding="utf-8")
    )
    assert manifest["active_categories"] == ["KAS", "HTB", "LAINNYA"]
    assert manifest["total_rows"] == 3
    assert manifest["revision"] == 1
    assert manifest["snapshot_hash"] == final_result.snapshot_hash


def test_stage8d4i_missing_values_remain_blank_through_final_and_export(tmp_path):
    db_path = tmp_path / "stage8d4i_blank.db"
    init_database(db_path)

    final_result = FinalizationService(
        db_path=db_path
    ).finalize(
        _finalization_input()
    )
    assert final_result.success

    templates = tmp_path / "templates"
    output = tmp_path / "official-package"
    _create_active_templates(templates)

    export_result = OfficialCoretaxPackageExporter(
        templates
    ).export_active_final(
        "6101015612710001",
        2025,
        output,
        db_path=db_path,
    )
    assert export_result.ok

    htb = load_workbook(export_result.excel_result.files["HTB"])
    htb_ws = htb["DATA"]
    assert htb_ws.cell(4, 8).value is None
    assert htb_ws.cell(4, 9).value == 88_000_000
    htb.close()

    lainnya = load_workbook(export_result.excel_result.files["LAINNYA"])
    lainnya_ws = lainnya["DATA"]
    assert lainnya_ws.cell(4, 5).value is None
    assert lainnya_ws.cell(4, 6).value == 100_000_000
    lainnya.close()


def test_stage8d4i_does_not_require_unused_templates(tmp_path):
    db_path = tmp_path / "stage8d4i_templates.db"
    init_database(db_path)
    assert FinalizationService(
        db_path=db_path
    ).finalize(
        _finalization_input()
    ).success

    templates = tmp_path / "templates"
    _create_active_templates(templates)

    assert not any("Piutang" in path.name for path in templates.glob("*.xlsx"))
    assert not any("Investasi" in path.name for path in templates.glob("*.xlsx"))
    assert not any("Bergerak" in path.name for path in templates.glob("*.xlsx"))

    result = OfficialCoretaxPackageExporter(
        templates
    ).export_active_final(
        "6101015612710001",
        2025,
        tmp_path / "official-package",
        db_path=db_path,
    )

    assert result.ok
    assert set(result.excel_result.files) == set(ACTIVE_CATEGORIES)
