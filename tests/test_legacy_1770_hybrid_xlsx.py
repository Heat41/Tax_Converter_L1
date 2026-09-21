import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from core.reconciliation import ReconciliationResult
from core.finalization import FinalizationInput
from core.legacy_1770_hybrid_xlsx import (
    DATA_BUPOT_SHEET,
    DATA_HARTA_SHEET,
    META_SHEET,
    Legacy1770HybridXlsxService,
)
from core.mapping.worksheet_harta_mapper import WorksheetHartaRow
from core.worksheet_pph_state import WorksheetBupotRow


def _analysis():
    return ReconciliationResult(
        total_harta_sebelumnya=10000000.0,
        total_harta_berjalan=15000000.0,
        total_utang_sebelumnya=0.0,
        total_utang_berjalan=0.0,
        naik_turun_harta_utang=5000000.0,
        biaya_hidup_setahun=19200000.0,
        pajak_pajak=0.0,
        pengeluaran_lain_lain=0.0,
        kerugian_keuntungan_penjualan_aset=0.0,
        utang_baru_atas_kredit=0.0,
        harta_baru_dari_kredit=0.0,
        total_pengeluaran=24200000.0,
        penghasilan_bruto_umkm=0.0,
        penambahan_penghasilan_bruto_umkm=0.0,
        margin_usaha=0.0,
        jumlah_penghasilan_bruto_final=0.0,
        jumlah_penghasilan_bukan_objek=0.0,
        penghasilan_netto=24200000.0,
        selisih_pengeluaran_vs_penghasilan=0.0,
    )


def _input():
    return FinalizationInput(
        npwp="1234567890123456",
        nama_wp="Budi",
        tahun_pajak=2025,
        harta_current_rows=[
            WorksheetHartaRow(
                nomor=1,
                kode_eform="011",
                kode_ct="0101",
                nama_harta="Kas",
                nomor_akun_keterangan="KAS-01",
                atas_nama="BUDI",
                nama_bank="-",
                tahun_perolehan=2020,
                nilai_tahun_sebelumnya=10000000.0,
                nilai_tahun_berjalan=15000000.0,
            )
        ],
        harta_original_hash="original-hash",
        bupot_rows=[
            WorksheetBupotRow(
                jenis="BP21",
                no_bupot="BP-001",
                masa="01",
                tahun="2025",
                sifat="TIDAK FINAL",
                status="NORMAL",
                npwp_penerima="1234567890123456",
                nama_penerima="Budi",
                fasilitas="",
                jenis_pph="Pasal 21",
                kop="21-100-07",
                bruto=500000.0,
                dpp_persen=50.0,
                tarif=5.0,
                pengurang=250000.0,
                pph_dipotong=12500.0,
                bukti="Surat Tagihan",
                no_bukti="DOC-1",
                tanggal_bukti="04/06/2025",
                npwp_pemotong="0072103856707000",
                nama_pemotong="PEMOTONG",
                tanggal_pemotongan="04/06/2025",
                mekanisme_sp2d="",
                no_sp2d="",
                npwp_pemberi_kerja="0072103856707000",
            )
        ],
        pph_components={"pph_terutang": 1000000.0},
        umkm_state={},
        penghasilan_lainnya={},
        zakat=0.0,
        status_ptkp="TK/0",
        pph_calc_result={"pph_terutang": 1000000.0},
        analisis_result=_analysis(),
    )


class TestLegacy1770HybridXlsx(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.path = Path(self.temp.name) / "hybrid.xlsx"
        self.service = Legacy1770HybridXlsxService()

    def tearDown(self):
        self.temp.cleanup()

    def test_export_builds_hybrid_and_roundtrip_sheets(self):
        result = self.service.export(
            _input(), self.path, revision=3, snapshot_hash="snapshot-abc"
        )
        self.assertTrue(result.ok)

        wb = load_workbook(self.path, data_only=False)
        self.assertEqual(
            wb.sheetnames[:6],
            [
                "01 eForm Induk H1",
                "02 eForm Induk H2",
                "03 Legacy Lamp I H2",
                "04 Legacy Lamp II",
                "05 Legacy Lamp III",
                "06 Legacy Lamp IV",
            ],
        )
        self.assertIn(DATA_HARTA_SHEET, wb.sheetnames)
        self.assertIn(DATA_BUPOT_SHEET, wb.sheetnames)
        self.assertEqual(wb[DATA_HARTA_SHEET].sheet_state, "veryHidden")
        self.assertEqual(wb[DATA_BUPOT_SHEET].sheet_state, "veryHidden")
        self.assertEqual(wb[META_SHEET].sheet_state, "veryHidden")
        self.assertEqual(
            wb["01 eForm Induk H1"]["A2"].value,
            "SPT TAHUNAN PAJAK PENGHASILAN (PPh) WAJIB PAJAK ORANG PRIBADI",
        )
        self.assertEqual(wb["01 eForm Induk H1"]["I3"].value, "HALAMAN 1")
        self.assertEqual(wb["02 eForm Induk H2"]["I3"].value, "HALAMAN 2")
        self.assertEqual(
            wb["03 Legacy Lamp I H2"]["D1"].value,
            "SPT TAHUNAN PPh WAJIB PAJAK ORANG PRIBADI",
        )
        self.assertEqual(wb["03 Legacy Lamp I H2"]["A3"].value, "LAMPIRAN - I")
        self.assertEqual(wb["03 Legacy Lamp I H2"]["I1"].value, "FORMULIR\n1770 - I")

    def test_first_two_pages_follow_new_eform_sections(self):
        self.service.export(_input(), self.path, revision=3, snapshot_hash="snapshot-abc")
        wb = load_workbook(self.path, data_only=False)
        page1 = wb["01 eForm Induk H1"]
        page2 = wb["02 eForm Induk H2"]

        page1_values = [str(page1.cell(row, 1).value or "") for row in range(1, page1.max_row + 1)]
        page2_values = [str(page2.cell(row, 1).value or "") for row in range(1, page2.max_row + 1)]

        self.assertIn("A. IDENTITAS WAJIB PAJAK", page1_values)
        self.assertIn("B. IKHTISAR PENGHASILAN NETO", page1_values)
        self.assertIn("C. PERHITUNGAN PPh TERUTANG", page1_values)
        self.assertIn("D. KREDIT PAJAK", page1_values)
        self.assertIn("E. PPh KURANG/LEBIH BAYAR", page2_values)
        self.assertIn("F. PEMBETULAN (DIISI JIKA STATUS SPT ADALAH PEMBETULAN)", page2_values)
        self.assertIn("H. ANGSURAN PPh PASAL 25 TAHUN PAJAK BERIKUTNYA", page2_values)
        self.assertIn("I. PERNYATAAN TRANSAKSI LAINNYA", page2_values)
        self.assertIn("J. LAMPIRAN TAMBAHAN", page2_values)
        self.assertIn("K. PERNYATAAN", page2_values)

    def test_new_eform_pages_are_print_ready_legal_single_page(self):
        self.service.export(_input(), self.path, revision=3, snapshot_hash="snapshot-abc")
        wb = load_workbook(self.path, data_only=False)

        for sheet_name in ("01 eForm Induk H1", "02 eForm Induk H2"):
            ws = wb[sheet_name]
            self.assertEqual(ws.page_setup.orientation, "portrait")
            self.assertEqual(str(ws.page_setup.paperSize), str(ws.PAPERSIZE_LEGAL))
            self.assertEqual(ws.page_setup.fitToWidth, 1)
            self.assertEqual(ws.page_setup.fitToHeight, 1)
            self.assertTrue(ws.sheet_properties.pageSetUpPr.fitToPage)
            self.assertFalse(ws.sheet_view.showGridLines)
            self.assertTrue(str(ws.print_area).startswith("'") or str(ws.print_area).startswith("$A$1"))

    def test_all_visible_output_sheets_use_legal_paper(self):
        self.service.export(_input(), self.path, revision=3, snapshot_hash="snapshot-abc")
        wb = load_workbook(self.path, data_only=False)

        expected = (
            "01 eForm Induk H1",
            "02 eForm Induk H2",
            "03 Legacy Lamp I H2",
            "04 Legacy Lamp II",
            "05 Legacy Lamp III",
            "06 Legacy Lamp IV",
        )
        for sheet_name in expected:
            ws = wb[sheet_name]
            self.assertEqual(str(ws.page_setup.paperSize), str(ws.PAPERSIZE_LEGAL), sheet_name)

    def test_legacy_lampiran_i_h2_renders_bagian_b_c_d_from_final_data(self):
        data = _input()
        data.pph_components["penghasilan_neto_lainnya"] = 54001000.0
        data.pph_calc_result["penghasilan_neto_lainnya"] = 54001000.0

        self.service.export(data, self.path, revision=3, snapshot_hash="snapshot-abc")
        ws = load_workbook(self.path, data_only=False)["03 Legacy Lamp I H2"]
        values = [str(ws.cell(row, 1).value or "") for row in range(1, ws.max_row + 1)]

        self.assertIn(
            "BAGIAN B : PENGHASILAN NETO DALAM NEGERI DARI USAHA DAN/ATAU PEKERJAAN BEBAS",
            values,
        )
        self.assertIn(
            "BAGIAN C : PENGHASILAN NETO DALAM NEGERI SEHUBUNGAN DENGAN PEKERJAAN",
            values,
        )
        self.assertIn("BAGIAN D : PENGHASILAN NETO DALAM NEGERI LAINNYA", values)

        text = " ".join(
            str(ws.cell(row, col).value or "")
            for row in range(1, ws.max_row + 1)
            for col in range(1, 11)
        )
        self.assertIn("PEMOTONG", text)
        self.assertIn("0072103856707000", text)
        self.assertIn("250000", text)
        self.assertIn("54001000", text)
        self.assertEqual(str(ws.page_setup.paperSize), str(ws.PAPERSIZE_LEGAL))
        self.assertEqual(ws.page_setup.orientation, "portrait")
        self.assertEqual(ws.page_setup.fitToHeight, 1)

    def test_hybrid_pages_keep_compact_column_widths(self):
        self.service.export(_input(), self.path, revision=3, snapshot_hash="snapshot-abc")
        wb = load_workbook(self.path, data_only=False)

        for sheet_name in ("01 eForm Induk H1", "02 eForm Induk H2", "03 Legacy Lamp I H2"):
            ws = wb[sheet_name]
            widths = [
                float(ws.column_dimensions[col].width or 0)
                for col in ("A", "B", "C", "D", "E", "F", "G", "H", "I", "J")
            ]
            self.assertLessEqual(max(widths), 13.0, sheet_name)
            self.assertLessEqual(widths[0], 5.0, sheet_name)

    def test_hybrid_format_boundary_is_locked_after_page_two(self):
        self.service.export(_input(), self.path, revision=3, snapshot_hash="snapshot-abc")
        wb = load_workbook(self.path, data_only=False)

        self.assertEqual(wb["01 eForm Induk H1"]["I3"].value, "HALAMAN 1")
        self.assertEqual(wb["02 eForm Induk H2"]["I3"].value, "HALAMAN 2")
        self.assertEqual(wb["03 Legacy Lamp I H2"]["A3"].value, "LAMPIRAN - I")

        lamp2 = wb["04 Legacy Lamp II"]
        self.assertEqual(lamp2["C1"].value, "LAMPIRAN - II")
        self.assertIn("1770 - II", str(lamp2["A1"].value or ""))
        self.assertEqual(lamp2["C2"].value, "SPT TAHUNAN PPh WAJIB PAJAK ORANG PRIBADI")

        lamp3 = wb["05 Legacy Lamp III"]
        self.assertEqual(lamp3["C1"].value, "LAMPIRAN - III")
        self.assertIn("1770 - III", str(lamp3["A1"].value or ""))
        self.assertEqual(lamp3["C2"].value, "SPT TAHUNAN PPh WAJIB PAJAK ORANG PRIBADI")

        lamp4 = wb["06 Legacy Lamp IV"]
        self.assertEqual(lamp4["C1"].value, "LAMPIRAN - IV")
        self.assertIn("1770 - IV", str(lamp4["A1"].value or ""))
        self.assertEqual(lamp4["C2"].value, "SPT TAHUNAN PPh WAJIB PAJAK ORANG PRIBADI")

        for sheet_name in (
            "03 Legacy Lamp I H2",
            "04 Legacy Lamp II",
            "05 Legacy Lamp III",
            "06 Legacy Lamp IV",
        ):
            ws = wb[sheet_name]
            self.assertEqual(str(ws.page_setup.paperSize), str(ws.PAPERSIZE_LEGAL))

    def test_roundtrip_reads_user_edits_and_keeps_provenance(self):
        self.service.export(_input(), self.path, revision=7, snapshot_hash="snapshot-final-7")

        wb = load_workbook(self.path)
        harta = wb[DATA_HARTA_SHEET]
        bupot = wb[DATA_BUPOT_SHEET]
        harta["D2"] = "Kas Koreksi"
        harta["J2"] = 17500000
        bupot["M2"] = 600000
        bupot["P2"] = 300000
        bupot["Q2"] = 15000
        wb.save(self.path)

        result = self.service.import_revision(
            self.path,
            expected_npwp="1234567890123456",
            expected_year=2025,
            expected_snapshot_hash="snapshot-final-7",
        )
        self.assertTrue(result.ok)
        self.assertEqual(result.base_revision, 7)
        self.assertEqual(result.base_snapshot_hash, "snapshot-final-7")
        self.assertEqual(result.harta_rows[0].nama_harta, "Kas Koreksi")
        self.assertEqual(result.harta_rows[0].nilai_tahun_berjalan, 17500000)
        self.assertEqual(result.bupot_rows[0].bruto, 600000)
        self.assertEqual(result.bupot_rows[0].pengurang, 300000)
        self.assertEqual(result.bupot_rows[0].pph_dipotong, 15000)

    def test_roundtrip_rejects_wrong_wp_or_snapshot(self):
        self.service.export(_input(), self.path, revision=2, snapshot_hash="snapshot-2")

        wrong_wp = self.service.import_revision(
            self.path,
            expected_npwp="9999999999999999",
            expected_year=2025,
        )
        self.assertFalse(wrong_wp.ok)
        self.assertTrue(any(i.code == "LX_105" for i in wrong_wp.errors))

        wrong_hash = self.service.import_revision(
            self.path,
            expected_npwp="1234567890123456",
            expected_year=2025,
            expected_snapshot_hash="other-hash",
        )
        self.assertFalse(wrong_hash.ok)
        self.assertTrue(any(i.code == "LX_107" for i in wrong_hash.errors))


if __name__ == "__main__":
    unittest.main()
