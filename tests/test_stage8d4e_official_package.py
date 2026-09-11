import json
from pathlib import Path

from openpyxl import Workbook

from core.coretax_official_schema import OFFICIAL_CORETAX_SCHEMAS
from core.reverse_coretax_mapping import ReverseCoretaxPackage, ReverseCoretaxRow
from core.reverse_coretax_official_package import OfficialCoretaxPackageExporter


CATEGORIES = (
    "KAS",
    "PIUTANG",
    "INVESTASI",
    "BERGERAK",
    "HTB",
    "LAINNYA",
)


def _templates(directory: Path):
    directory.mkdir(parents=True, exist_ok=True)

    for category in CATEGORIES:
        schema = OFFICIAL_CORETAX_SCHEMAS[category]
        path = directory / f"{schema.excel_filename_hint}.xlsx"

        wb = Workbook()
        ws = wb.active
        ws.title = "DATA"
        ws["A1"] = "NPWP"
        ws["B1"] = ""
        ws["A2"] = "Tahun Pajak"
        ws["B2"] = ""

        for index, header in enumerate(schema.excel_headers, start=1):
            ws.cell(3, index).value = header

        wb.save(path)


def _package():
    package = ReverseCoretaxPackage(
        npwp="6101015612710001",
        nama_wp="EVY BACHTIAR",
        tahun_pajak=2025,
        revision=3,
        snapshot_hash="abc123",
    )

    package.rows_by_category["KAS"].append(
        ReverseCoretaxRow(
            nomor=1,
            kategori="KAS",
            kode_harta="0102",
            nama_harta="Tabungan",
            tahun_perolehan=2025,
            nilai=391742658,
            nomor_akun_keterangan="116801005138508",
            atas_nama="DR EVY BACHTIAR SPOG",
            nama_bank="BRI",
            official_metadata={
                "account_number": "116801005138508",
                "account_on_behalf_of": "DR EVY BACHTIAR SPOG",
                "bank_name": "BRI",
                "country": "Indonesia",
                "year": 2025,
                "balance": 391742658,
            },
        )
    )

    return package


def test_stage8d4e_builds_excel_xml_and_manifest(tmp_path):
    templates = tmp_path / "templates"
    output = tmp_path / "package"
    _templates(templates)

    result = OfficialCoretaxPackageExporter(
        templates
    ).export_package(
        _package(),
        output,
    )

    assert result.ok
    assert result.excel_result is not None
    assert set(result.excel_result.files) == {"KAS"}
    assert result.xml_result is not None
    assert set(result.xml_result.files) == {"KAS"}
    assert result.manifest_path == output / "manifest.json"
    assert result.manifest_path.is_file()


def test_stage8d4e_separates_excel_and_xml_directories(tmp_path):
    templates = tmp_path / "templates"
    output = tmp_path / "package"
    _templates(templates)

    result = OfficialCoretaxPackageExporter(
        templates
    ).export_package(
        _package(),
        output,
    )

    assert all(path.parent == output / "excel" for path in result.excel_result.files.values())
    assert all(path.parent == output / "xml" for path in result.xml_result.files.values())


def test_stage8d4e_manifest_records_snapshot_identity(tmp_path):
    templates = tmp_path / "templates"
    output = tmp_path / "package"
    _templates(templates)

    result = OfficialCoretaxPackageExporter(
        templates
    ).export_package(
        _package(),
        output,
    )

    manifest = json.loads(result.manifest_path.read_text(encoding="utf-8"))

    assert manifest["stage"] == "8D.4E"
    assert manifest["npwp"] == "6101015612710001"
    assert manifest["nama_wp"] == "EVY BACHTIAR"
    assert manifest["tahun_pajak"] == 2025
    assert manifest["revision"] == 3
    assert manifest["snapshot_hash"] == "abc123"
    assert manifest["total_rows"] == 1


def test_stage8d4e_manifest_records_supported_and_unsupported_outputs(tmp_path):
    templates = tmp_path / "templates"
    output = tmp_path / "package"
    _templates(templates)

    result = OfficialCoretaxPackageExporter(
        templates
    ).export_package(
        _package(),
        output,
    )

    manifest = json.loads(result.manifest_path.read_text(encoding="utf-8"))

    assert manifest["excel"]["file_count"] == 1
    assert set(manifest["excel"]["files"]) == {"KAS"}
    assert manifest["xml"]["file_count"] == 1
    assert set(manifest["xml"]["files"]) == {"KAS"}
    assert manifest["xml"]["unsupported_categories"] == []


def test_stage8d4e_manifest_contains_file_hashes(tmp_path):
    templates = tmp_path / "templates"
    output = tmp_path / "package"
    _templates(templates)

    result = OfficialCoretaxPackageExporter(
        templates
    ).export_package(
        _package(),
        output,
    )

    manifest = json.loads(result.manifest_path.read_text(encoding="utf-8"))

    for section in ("excel", "xml"):
        for item in manifest[section]["files"].values():
            assert len(item["sha256"]) == 64
            assert item["path"]
            assert item["filename"]


def test_stage8d4e_only_needs_templates_for_categories_with_data(tmp_path):
    templates = tmp_path / "templates"
    output = tmp_path / "package"
    templates.mkdir(parents=True, exist_ok=True)

    schema = OFFICIAL_CORETAX_SCHEMAS["KAS"]
    path = templates / f"{schema.excel_filename_hint}.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "DATA"
    ws["A1"] = "NPWP"
    ws["B1"] = ""
    ws["A2"] = "Tahun Pajak"
    ws["B2"] = ""
    for index, header in enumerate(schema.excel_headers, start=1):
        ws.cell(3, index).value = header
    wb.save(path)

    result = OfficialCoretaxPackageExporter(
        templates
    ).export_package(
        _package(),
        output,
    )

    assert result.ok
    assert set(result.excel_result.files) == {"KAS"}
    assert set(result.xml_result.files) == {"KAS"}
