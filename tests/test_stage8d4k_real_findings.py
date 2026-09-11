from pathlib import Path

from openpyxl import Workbook

from core.mapping.harta_mapper import CoretaxHartaMapper
from core.physical_reconciliation import PhysicalSourceExportReconciler


def test_investment_identity_maps_to_institution_field():
    item = CoretaxHartaMapper().map_row(
        "INVESTASI",
        {
            "Kode *": "0305",
            "Lokasi Harta *": "Indonesia",
            "Nomor Identitas *": "6101015612710001",
            "Nama Bank/Institusi/Penerima Investasi *": "EVY BACHTIAR",
            "Bukti Kepemilikan/Nomor Akun *": "1971197117",
            "Biaya Perolehan *": "1000000000",
            "Tahun Perolehan *": "2023",
            "Nilai Saat Ini *": "1000000000",
            "Keterangan": "",
        },
        wp_id=0,
    )

    assert item.nomor_identitas_institusi == "6101015612710001"
    assert item.nomor_identitas_pihak_ketiga is None


def test_physical_reconciliation_resolves_simple_absolute_cell_formula(tmp_path):
    source = tmp_path / "source"
    exported = tmp_path / "export"
    source.mkdir()
    exported.mkdir()

    source_path = source / "PIT L1 Harta Bergerak.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "DATA"
    headers = (
        "Kode *",
        "Merk/Model *",
        "Nomor Polisi/Registrasi *",
        "Kepemilikan*",
        "NPWP Pemilik*",
        "Nama Pemilik *",
        "Tahun Perolehan *",
        "Biaya Perolehan *",
        "Nilai Saat Ini *",
        "Keterangan",
    )
    ws["A1"] = "NPWP"
    ws["B1"] = "6101015612710001"
    for col, header in enumerate(headers, start=1):
        ws.cell(3, col).value = header
    values = (
        "0402",
        "HONDA BEAT",
        "KB 2183 T",
        "TAXPAYER",
        "=$B$1",
        "EVY BACHTIAR",
        2009,
        11000000,
        3000000,
        "",
    )
    for col, value in enumerate(values, start=1):
        ws.cell(4, col).value = value
    wb.save(source_path)

    export_path = exported / "PIT L1 Harta Bergerak.xlsx"
    wb2 = Workbook()
    ws2 = wb2.active
    ws2.title = "DATA"
    for col, header in enumerate(headers, start=1):
        ws2.cell(3, col).value = header
    export_values = list(values)
    export_values[4] = "6101015612710001"
    for col, value in enumerate(export_values, start=1):
        ws2.cell(4, col).value = value
    wb2.save(export_path)

    result = PhysicalSourceExportReconciler().reconcile(
        source,
        exported,
    )

    assert result.ok
    assert result.categories["BERGERAK"].excel_match is True
