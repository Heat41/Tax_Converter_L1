import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from core.evy_reconciliation import EvyReconciliationResult
from core.finalization import FinalizationInput
from core.legacy_1770_hybrid_l4 import LegacyLampiranIVXlsxRenderer
from core.mapping.worksheet_harta_mapper import WorksheetHartaRow


def _analysis():
    return EvyReconciliationResult(
        total_harta_sebelumnya=10000000.0,
        total_harta_berjalan=15000000.0,
        total_utang_sebelumnya=0.0,
        total_utang_berjalan=0.0,
        naik_turun_harta_utang=5000000.0,
        biaya_hidup_setahun=0.0,
        pajak_pajak=0.0,
        pengeluaran_lain_lain=0.0,
        kerugian_keuntungan_penjualan_aset=0.0,
        utang_baru_atas_kredit=0.0,
        harta_baru_dari_kredit=0.0,
        total_pengeluaran=0.0,
        penghasilan_bruto_umkm=0.0,
        penambahan_penghasilan_bruto_umkm=0.0,
        margin_usaha=0.0,
        jumlah_penghasilan_bruto_final=0.0,
        jumlah_penghasilan_bukan_objek=0.0,
        penghasilan_netto=0.0,
        selisih_pengeluaran_vs_penghasilan=0.0,
    )


def _input():
    return FinalizationInput(
        npwp="3275017007900006",
        nama_wp="Lisa Vinatalia",
        tahun_pajak=2025,
        harta_current_rows=[
            WorksheetHartaRow(
                nomor=1,
                kode_eform="011",
                kode_ct="0101",
                nama_harta="Kas",
                nomor_akun_keterangan="KAS-01",
                atas_nama="LISA VINATALIA",
                nama_bank="-",
                tahun_perolehan=2020,
                nilai_tahun_sebelumnya=10000000.0,
                nilai_tahun_berjalan=15000000.0,
            )
        ],
        harta_original_hash="",
        bupot_rows=[],
        pph_components={},
        umkm_state={},
        penghasilan_lainnya={},
        zakat=0.0,
        status_ptkp="TK/0",
        pph_calc_result={},
        analisis_result=_analysis(),
    )


class TestLegacyLampiranIVXlsx(unittest.TestCase):
    def test_renderer_maps_harta_and_keeps_empty_utang_family_sections(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "lampiran_iv.xlsx"
            wb = Workbook()
            ws = wb.active
            ws.title = "06 Legacy Lamp IV"
            LegacyLampiranIVXlsxRenderer().render(ws, _input())
            wb.save(path)

            check = load_workbook(path, data_only=False)["06 Legacy Lamp IV"]
            text = " ".join(
                str(check.cell(row, col).value or "")
                for row in range(1, check.max_row + 1)
                for col in range(1, check.max_column + 1)
            )

            self.assertIn("1770 - IV", text)
            self.assertIn("LAMPIRAN - IV", text)
            self.assertIn("BAGIAN A : HARTA PADA AKHIR TAHUN", text)
            self.assertIn("BAGIAN B : KEWAJIBAN/UTANG PADA AKHIR TAHUN", text)
            self.assertIn("BAGIAN C : DAFTAR SUSUNAN ANGGOTA KELUARGA", text)
            self.assertIn("Kas", text)
            self.assertIn("011", text)
            self.assertIn("KAS-01", text)
            self.assertIn("LISA VINATALIA", text)

            numeric_values = {
                check.cell(row, col).value
                for row in range(1, check.max_row + 1)
                for col in range(1, check.max_column + 1)
                if isinstance(check.cell(row, col).value, (int, float))
            }
            self.assertIn(15000000, numeric_values)
            self.assertEqual(str(check.page_setup.paperSize), str(check.PAPERSIZE_LEGAL))
            self.assertEqual(check.page_setup.orientation, "portrait")
            self.assertEqual(check.page_setup.fitToWidth, 1)
            self.assertEqual(check.page_setup.fitToHeight, 1)


if __name__ == "__main__":
    unittest.main()
