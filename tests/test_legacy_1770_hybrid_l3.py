import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from core.reconciliation import ReconciliationResult
from core.finalization import FinalizationInput
from core.legacy_1770_hybrid_l3 import LegacyLampiranIIIXlsxRenderer


def _analysis():
    return ReconciliationResult(
        total_harta_sebelumnya=0.0,
        total_harta_berjalan=0.0,
        total_utang_sebelumnya=0.0,
        total_utang_berjalan=0.0,
        naik_turun_harta_utang=0.0,
        biaya_hidup_setahun=0.0,
        pajak_pajak=0.0,
        pengeluaran_lain_lain=0.0,
        kerugian_keuntungan_penjualan_aset=0.0,
        utang_baru_atas_kredit=0.0,
        harta_baru_dari_kredit=0.0,
        total_pengeluaran=0.0,
        penghasilan_bruto_umkm=150000000.0,
        penambahan_penghasilan_bruto_umkm=0.0,
        margin_usaha=0.0,
        jumlah_penghasilan_bruto_final=160000000.0,
        jumlah_penghasilan_bukan_objek=6832000.0,
        penghasilan_netto=0.0,
        selisih_pengeluaran_vs_penghasilan=0.0,
    )


def _input():
    return FinalizationInput(
        npwp="3275017007900006",
        nama_wp="Lisa Vinatalia",
        tahun_pajak=2025,
        harta_current_rows=[],
        harta_original_hash="",
        bupot_rows=[],
        pph_components={},
        umkm_state={
            "bruto_bulanan": {"01": 100000000.0, "02": 50000000.0},
            "pph_setor": 750000.0,
        },
        penghasilan_lainnya={
            "final_other_rows": [
                {"keterangan": "DIVIDEN", "dpp": 10000000.0, "pph": 1000000.0},
            ],
        },
        zakat=0.0,
        status_ptkp="TK/0",
        pph_calc_result={},
        analisis_result=_analysis(),
    )


class TestLegacyLampiranIIIXlsx(unittest.TestCase):
    def test_renderer_maps_final_and_non_object_income_to_1770_iii(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "lampiran_iii.xlsx"
            wb = Workbook()
            ws = wb.active
            ws.title = "05 Legacy Lamp III"
            LegacyLampiranIIIXlsxRenderer().render(ws, _input())
            wb.save(path)

            check = load_workbook(path, data_only=False)["05 Legacy Lamp III"]
            text = " ".join(
                str(check.cell(row, col).value or "")
                for row in range(1, check.max_row + 1)
                for col in range(1, check.max_column + 1)
            )

            self.assertIn("1770 - III", text)
            self.assertIn("LAMPIRAN - III", text)
            self.assertIn("BAGIAN A", text)
            self.assertIn("BAGIAN B", text)
            self.assertIn("BAGIAN C", text)
            self.assertIn("DIVIDEN", text)
            self.assertIn("PENGHASILAN LAIN YANG DIKENAKAN PAJAK FINAL", text)
            self.assertIn("PENGHASILAN LAIN YANG TIDAK TERMASUK OBJEK PAJAK", text)

            numeric_values = {
                check.cell(row, col).value
                for row in range(1, check.max_row + 1)
                for col in range(1, check.max_column + 1)
                if isinstance(check.cell(row, col).value, (int, float))
            }
            self.assertIn(10000000, numeric_values)
            self.assertIn(1000000, numeric_values)
            self.assertIn(150000000, numeric_values)
            self.assertIn(750000, numeric_values)
            self.assertIn(6832000, numeric_values)

            jumlah_a_row = next(
                row
                for row in range(1, check.max_row + 1)
                if check.cell(row, 1).value == "17. JUMLAH (1 s.d. 16)"
            )
            self.assertTrue(str(check.cell(jumlah_a_row, 7).value).startswith("=SUM(G"))
            self.assertTrue(str(check.cell(jumlah_a_row, 9).value).startswith("=SUM(I"))

            jumlah_b_row = next(
                row
                for row in range(1, check.max_row + 1)
                if check.cell(row, 1).value == "JUMLAH BAGIAN B"
            )
            self.assertTrue(str(check.cell(jumlah_b_row, 8).value).startswith("=SUM(H"))

            self.assertEqual(str(check.page_setup.paperSize), str(check.PAPERSIZE_LEGAL))
            self.assertEqual(check.page_setup.orientation, "portrait")
            self.assertEqual(check.page_setup.fitToWidth, 1)
            self.assertEqual(check.page_setup.fitToHeight, 1)


if __name__ == "__main__":
    unittest.main()
