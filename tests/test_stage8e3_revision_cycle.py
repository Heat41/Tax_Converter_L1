import json

from dataclasses import replace

from openpyxl import load_workbook

from config.database import get_db_connection, init_database
from core.finalization import FinalizationInput, FinalizationService
from core.legacy_1770_hybrid_xlsx import Legacy1770HybridXlsxService, META_SHEET
from core.mapping.worksheet_harta_mapper import WorksheetHartaRow
from core.reconciliation import ReconciliationResult, calculate_reconciliation
from core.worksheet_pph_state import WorksheetBupotRow


def _analysis(*, harta_now=15_000_000.0, utang_now=5_000_000.0):
    return calculate_reconciliation(
        total_harta_sebelumnya=10_000_000.0,
        total_harta_berjalan=harta_now,
        total_utang_sebelumnya=0.0,
        total_utang_berjalan=utang_now,
        status_ptkp="TK/0",
        netto_bupot=24_000_000.0,
        penghasilan_netto_override=24_000_000.0,
    )


def _input():
    return FinalizationInput(
        npwp="1234567890123456",
        nama_wp="Budi",
        tahun_pajak=2025,
        harta_current_rows=[
            WorksheetHartaRow(
                nomor=1,
                kode_eform="011",
                kode_ct="0101",
                nama_harta="Kas",
                nomor_akun_keterangan="KAS-01",
                atas_nama="BUDI",
                nama_bank="-",
                tahun_perolehan=2020,
                nilai_tahun_sebelumnya=10_000_000.0,
                nilai_tahun_berjalan=15_000_000.0,
            )
        ],
        harta_original_hash="original-hash",
        bupot_rows=[
            WorksheetBupotRow(
                jenis="BP21",
                no_bupot="BP-001",
                masa="01",
                tahun="2025",
                sifat="TIDAK FINAL",
                status="NORMAL",
                npwp_penerima="1234567890123456",
                nama_penerima="Budi",
                fasilitas="",
                jenis_pph="Pasal 21",
                kop="21-100-07",
                bruto=25_000_000.0,
                dpp_persen=100.0,
                tarif=5.0,
                pengurang=1_000_000.0,
                pph_dipotong=1_250_000.0,
                bukti="",
                no_bukti="",
                tanggal_bukti="",
                npwp_pemotong="0072103856707000",
                nama_pemotong="PEMOTONG",
                tanggal_pemotongan="04/06/2025",
                mekanisme_sp2d="",
                no_sp2d="",
                npwp_pemberi_kerja="0072103856707000",
            )
        ],
        pph_components={
            "utang_rows": [
                {
                    "kode_utang": "101",
                    "nama_pemberi_pinjaman": "BANK DUMMY",
                    "alamat_pemberi_pinjaman": "PONTIANAK",
                    "tahun_pinjaman": 2024,
                    "jumlah": 5_000_000.0,
                }
            ],
            "pph_terutang": 1_000_000.0,
        },
        umkm_state={},
        penghasilan_lainnya={},
        zakat=0.0,
        status_ptkp="TK/0",
        pph_calc_result={"pph_terutang": 1_000_000.0},
        analisis_result=_analysis(),
    )


def _find_row(ws, text):
    return next(
        row
        for row in range(1, ws.max_row + 1)
        if str(ws.cell(row, 1).value or "").strip() == text
    )


def test_stage8e3_full_revision_cycle_creates_next_revision_and_reexports(tmp_path):
    db_path = tmp_path / "stage8e3.db"
    init_database(db_path)

    finalizer = FinalizationService(db_path=db_path)
    xlsx = Legacy1770HybridXlsxService()
    original = _input()

    # 1. Finalisasi awal.
    rev1 = finalizer.finalize(original)
    assert rev1.success is True
    assert rev1.revision == 1

    # 2. Export workbook revisi dari Revision 1.
    rev1_book = tmp_path / "revision-1.xlsx"
    exported = xlsx.export(
        original,
        rev1_book,
        revision=rev1.revision,
        snapshot_hash=rev1.snapshot_hash,
    )
    assert exported.ok

    # 3. User mengubah Harta, Utang, dan Bupot pada sheet visual.
    wb = load_workbook(rev1_book)

    lamp4 = wb["06 Legacy Lamp IV"]
    harta_section = _find_row(lamp4, "BAGIAN A : HARTA PADA AKHIR TAHUN")
    harta_row = harta_section + 2
    lamp4.cell(harta_row, 3).value = "Kas Revisi"
    lamp4.cell(harta_row, 6).value = 17_000_000

    utang_section = _find_row(
        lamp4, "BAGIAN B : KEWAJIBAN/UTANG PADA AKHIR TAHUN"
    )
    utang_row = utang_section + 2
    lamp4.cell(utang_row, 2).value = "BANK REVISI"
    lamp4.cell(utang_row, 9).value = 7_500_000

    lamp2 = wb["04 Legacy Lamp II"]
    bupot_header = next(
        row
        for row in range(1, lamp2.max_row + 1)
        if str(lamp2.cell(row, 1).value or "").strip().upper() == "NO"
    )
    bupot_row = bupot_header + 1
    lamp2.cell(bupot_row, 2).value = "PEMOTONG REVISI"
    lamp2.cell(bupot_row, 9).value = 1_500_000
    wb.save(rev1_book)

    # 4. Import revisi harus mengikat workbook ke Revision 1.
    imported = xlsx.import_revision(
        rev1_book,
        expected_npwp=original.npwp,
        expected_year=original.tahun_pajak,
        expected_snapshot_hash=rev1.snapshot_hash,
    )
    assert imported.ok
    assert imported.base_revision == 1
    assert imported.base_snapshot_hash == rev1.snapshot_hash
    assert imported.harta_rows[0].nama_harta == "Kas Revisi"
    assert imported.harta_rows[0].nilai_tahun_berjalan == 17_000_000
    assert imported.utang_rows[0]["nama_pemberi_pinjaman"] == "BANK REVISI"
    assert imported.utang_rows[0]["jumlah"] == 7_500_000
    assert imported.bupot_rows[0].nama_pemotong == "PEMOTONG REVISI"
    assert imported.bupot_rows[0].pph_dipotong == 1_500_000

    # 5. Buka kembali/VOID Revision 1 lalu bentuk Edited / Current hasil revisi.
    assert finalizer.void_snapshot(rev1.snapshot_id, "Stage 8E.3 revision cycle")

    revised_components = dict(original.pph_components)
    revised_components["utang_rows"] = list(imported.utang_rows)

    revised = replace(
        original,
        harta_current_rows=list(imported.harta_rows),
        bupot_rows=list(imported.bupot_rows),
        pph_components=revised_components,
        analisis_result=_analysis(
            harta_now=17_000_000.0,
            utang_now=7_500_000.0,
        ),
    )

    # 6. Finalisasi ulang harus menjadi Revision 2 dengan snapshot berbeda.
    rev2 = finalizer.finalize(revised)
    assert rev2.success is True
    assert rev2.revision == 2
    assert rev2.snapshot_hash != rev1.snapshot_hash

    history = finalizer.list_snapshots(original.npwp, original.tahun_pajak)
    assert [row["revision"] for row in history[:2]] == [2, 1]
    assert history[0]["status"] == "FINAL"
    assert history[1]["status"] == "VOID"

    # 7. Snapshot Revision 2 harus menyimpan hasil revisi, bukan data Revision 1.
    conn = get_db_connection(db_path)
    try:
        stored = conn.execute(
            "SELECT snapshot_json FROM worksheet_final_snapshots WHERE id = ?",
            (rev2.snapshot_id,),
        ).fetchone()
    finally:
        conn.close()

    payload = json.loads(stored["snapshot_json"])
    assert payload["harta_current_rows"][0]["nama_harta"] == "Kas Revisi"
    assert payload["harta_current_rows"][0]["nilai_tahun_berjalan"] == 17_000_000
    assert payload["pph_components"]["utang_rows"][0]["jumlah"] == 7_500_000
    assert payload["bupot_rows"][0]["nama_pemotong"] == "PEMOTONG REVISI"
    assert payload["bupot_rows"][0]["pph_dipotong"] == 1_500_000

    # 8. Re-export Revision 2 harus membawa data hasil revisi dan provenance baru.
    rev2_book = tmp_path / "revision-2.xlsx"
    exported2 = xlsx.export(
        revised,
        rev2_book,
        revision=rev2.revision,
        snapshot_hash=rev2.snapshot_hash,
    )
    assert exported2.ok

    wb2 = load_workbook(rev2_book, data_only=False)
    lamp4_2 = wb2["06 Legacy Lamp IV"]
    harta_section_2 = _find_row(lamp4_2, "BAGIAN A : HARTA PADA AKHIR TAHUN")
    utang_section_2 = _find_row(
        lamp4_2, "BAGIAN B : KEWAJIBAN/UTANG PADA AKHIR TAHUN"
    )
    assert lamp4_2.cell(harta_section_2 + 2, 3).value == "Kas Revisi"
    assert lamp4_2.cell(harta_section_2 + 2, 6).value == 17_000_000
    assert lamp4_2.cell(utang_section_2 + 2, 2).value == "BANK REVISI"
    assert lamp4_2.cell(utang_section_2 + 2, 9).value == 7_500_000

    meta = wb2[META_SHEET]
    meta_values = {
        str(meta.cell(row, 1).value): meta.cell(row, 2).value
        for row in range(1, meta.max_row + 1)
    }
    assert int(meta_values["base_revision"]) == 2
    assert str(meta_values["base_snapshot_hash"]) == rev2.snapshot_hash
