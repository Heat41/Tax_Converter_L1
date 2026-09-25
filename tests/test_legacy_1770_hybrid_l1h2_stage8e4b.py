import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from core.finalization import FinalizationInput
from core.legacy_1770_hybrid_l1h2 import LegacyLampiranIH2XlsxRenderer
from core.reconciliation import ReconciliationResult


def _analysis():
    return ReconciliationResult(
        total_harta_sebelumnya=0.0,
        total_harta_berjalan=0.0,
        total_utang_sebelumnya=0.0,
        total_utang_berjalan=0.0,
        naik_turun_harta_utang=0.0,
        biaya_hidup_setahun=0.0,
        pajak_pajak=0.0,
        pengeluaran_lain_lain=0.0,
        kerugian_keuntungan_penjualan_aset=0.0,
        utang_baru_atas_kredit=0.0,
        harta_baru_dari_kredit=0.0,
        total_pengeluaran=0.0,
        penghasilan_bruto_umkm=0.0,
        penambahan_penghasilan_bruto_umkm=0.0,
        margin_usaha=0.0,
        jumlah_penghasilan_bruto_final=0.0,
        jumlah_penghasilan_bukan_objek=0.0,
        penghasilan_netto=0.0,
        selisih_pengeluaran_vs_penghasilan=0.0,
    )


def _input(pekerjaan_bebas=48_000_000.0):
    return FinalizationInput(
        npwp="3275017007900006",
        nama_wp="Lisa Vinatalia",
        tahun_pajak=2025,
        harta_current_rows=[],
        harta_original_hash="",
        bupot_rows=[],
        pph_components={},
        umkm_state={},
        penghasilan_lainnya={
            "pekerjaan_bebas_dpp": pekerjaan_bebas,
        },
        zakat=0.0,
        status_ptkp="TK/0",
        pph_calc_result={},
        analisis_result=_analysis(),
    )


class TestLegacyLampiranIH2Stage8E4B(unittest.TestCase):
    def test_pekerjaan_bebas_uses_explicit_source_without_guessing_norma_or_neto(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "lampiran_i_h2.xlsx"
            wb = Workbook()
            ws = wb.active
            ws.title = "03 Legacy Lamp I H2"
            LegacyLampiranIH2XlsxRenderer().render(ws, _input())
            wb.save(path)

            check = load_workbook(path, data_only=False)["03 Legacy Lamp I H2"]

            pekerjaan_row = next(
                row
                for row in range(1, check.max_row + 1)
                if check.cell(row, 2).value == "PEKERJAAN BEBAS"
            )
            self.assertEqual(check.cell(pekerjaan_row, 5).value, 48_000_000)
            self.assertIn(check.cell(pekerjaan_row, 7).value, (None, ""))
            self.assertIn(check.cell(pekerjaan_row, 9).value, (None, ""))

            for label in ("DAGANG", "INDUSTRI", "JASA"):
                row = next(
                    r
                    for r in range(1, check.max_row + 1)
                    if check.cell(r, 2).value == label
                )
                self.assertIn(check.cell(row, 5).value, (None, ""))
                self.assertIn(check.cell(row, 7).value, (None, ""))
                self.assertIn(check.cell(row, 9).value, (None, ""))

    def test_pekerjaan_bebas_falls_back_to_component_other_income_contract(self):
        data = _input(0)
        data.penghasilan_lainnya = {}
        data.pph_components = {
            "other_income": {
                "pekerjaan_bebas_dpp": 12_500_000.0,
            }
        }

        wb = Workbook()
        ws = wb.active
        LegacyLampiranIH2XlsxRenderer().render(ws, data)

        pekerjaan_row = next(
            row
            for row in range(1, ws.max_row + 1)
            if ws.cell(row, 2).value == "PEKERJAAN BEBAS"
        )
        self.assertEqual(ws.cell(pekerjaan_row, 5).value, 12_500_000)
        self.assertIn(ws.cell(pekerjaan_row, 7).value, (None, ""))
        self.assertIn(ws.cell(pekerjaan_row, 9).value, (None, ""))


if __name__ == "__main__":
    unittest.main()
