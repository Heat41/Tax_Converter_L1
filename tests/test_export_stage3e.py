import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from core.exporters.legacy_1770iv_writer import Legacy1770IVExcelWriter
from core.mapping.legacy_1770iv import Legacy1770IVRow


class TestLegacy1770IVExcelWriterStage3E(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.root = Path(self.temp_dir.name)
        self.template = self.root / "template.xlsx"
        self.output = self.root / "hasil.xlsx"
        self.writer = Legacy1770IVExcelWriter()
        self._create_template()

    def tearDown(self):
        self.temp_dir.cleanup()

    def _create_template(self):
        wb = Workbook()
        ws = wb.active
        ws.title = "1770-IV"
        other = wb.create_sheet("LAIN")
        other["A1"] = "JANGAN BERUBAH"

        ws["A1"] = "FORMULIR 1770-IV"
        headers = [
            "Kode Harta",
            "Nama Harta",
            "Tahun Perolehan",
            "Harga Perolehan",
            "Keterangan",
        ]
        for idx, header in enumerate(headers, start=2):
            cell = ws.cell(row=4, column=idx, value=header)
            cell.font = Font(bold=True)

        thin = Side(style="thin")
        for row in range(5, 7):
            for col in range(2, 7):
                cell = ws.cell(row=row, column=col)
                cell.fill = PatternFill(fill_type="solid", fgColor="FFF2CC")
                cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
            ws.cell(row=row, column=5).number_format = '#,##0'

        ws["D7"] = "TOTAL"
        ws["E7"] = "=SUM(E5:E6)"
        ws["D7"].font = Font(bold=True)
        ws["E7"].font = Font(bold=True)

        wb.save(self.template)

    @staticmethod
    def _rows(count=2):
        base = [
            Legacy1770IVRow("0101", "BANK A", 2020, 125000000, "Kas"),
            Legacy1770IVRow("0401", "TOYOTA AVANZA", 2019, 190000000, "Mobil"),
            Legacy1770IVRow("0501", "RUMAH PONTIANAK", 2018, 850000000, "Rumah"),
        ]
        return base[:count]

    def test_write_detects_sheet_and_headers(self):
        result = self.writer.write(self.template, self.output, self._rows(2))

        self.assertEqual(result.sheet_name, "1770-IV")
        self.assertEqual(result.header_row, 4)
        self.assertEqual(result.start_row, 5)
        self.assertEqual(result.written_rows, 2)
        self.assertEqual(result.total_acquisition_cost, 315000000.0)
        self.assertTrue(self.output.exists())

    def test_written_values_match_canonical_rows(self):
        self.writer.write(self.template, self.output, self._rows(2))
        wb = load_workbook(self.output, data_only=False)
        ws = wb["1770-IV"]

        self.assertEqual(ws["B5"].value, "0101")
        self.assertEqual(ws["C5"].value, "BANK A")
        self.assertEqual(ws["D5"].value, 2020)
        self.assertEqual(ws["E5"].value, 125000000)
        self.assertEqual(ws["F5"].value, "Kas")

        self.assertEqual(ws["B6"].value, "0401")
        self.assertEqual(ws["E6"].value, 190000000)

    def test_existing_footer_formula_is_preserved(self):
        self.writer.write(self.template, self.output, self._rows(2))
        wb = load_workbook(self.output, data_only=False)
        ws = wb["1770-IV"]

        self.assertEqual(ws["D7"].value, "TOTAL")
        self.assertEqual(ws["E7"].value, "=SUM(E5:E6)")

    def test_writer_inserts_rows_before_footer_when_needed(self):
        self.writer.write(self.template, self.output, self._rows(3))
        wb = load_workbook(self.output, data_only=False)
        ws = wb["1770-IV"]

        self.assertEqual(ws["B7"].value, "0501")
        self.assertEqual(ws["C7"].value, "RUMAH PONTIANAK")
        self.assertEqual(ws["D8"].value, "TOTAL")
        self.assertEqual(ws["E8"].value, "=SUM(E5:E7)")

    def test_inserted_row_inherits_template_style(self):
        self.writer.write(self.template, self.output, self._rows(3))
        wb = load_workbook(self.output)
        ws = wb["1770-IV"]

        self.assertEqual(ws["B7"].fill.fill_type, ws["B6"].fill.fill_type)
        self.assertEqual(ws["B7"].fill.fgColor.rgb, ws["B6"].fill.fgColor.rgb)
        self.assertEqual(ws["E7"].number_format, ws["E6"].number_format)

    def test_unrelated_sheet_is_preserved(self):
        self.writer.write(self.template, self.output, self._rows(2))
        wb = load_workbook(self.output)

        self.assertIn("LAIN", wb.sheetnames)
        self.assertEqual(wb["LAIN"]["A1"].value, "JANGAN BERUBAH")

    def test_explicit_sheet_name_supported(self):
        result = self.writer.write(
            self.template,
            self.output,
            self._rows(1),
            sheet_name="1770-IV",
        )
        self.assertEqual(result.sheet_name, "1770-IV")

    def test_unknown_sheet_name_is_error(self):
        with self.assertRaises(ValueError):
            self.writer.write(
                self.template,
                self.output,
                self._rows(1),
                sheet_name="TIDAK ADA",
            )

    def test_missing_template_is_error(self):
        with self.assertRaises(FileNotFoundError):
            self.writer.write(
                self.root / "missing.xlsx",
                self.output,
                self._rows(1),
            )

    def test_missing_headers_is_error(self):
        broken = self.root / "broken.xlsx"
        wb = Workbook()
        wb.active["A1"] = "BUKAN TEMPLATE 1770-IV"
        wb.save(broken)

        with self.assertRaises(ValueError):
            self.writer.write(broken, self.output, self._rows(1))


if __name__ == "__main__":
    unittest.main()
