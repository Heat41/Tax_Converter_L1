import json
from pathlib import Path

from openpyxl import Workbook

from core.coretax_official_schema import OFFICIAL_CORETAX_SCHEMAS
from core.reverse_coretax_mapping import ReverseCoretaxPackage, ReverseCoretaxRow
from core.reverse_coretax_official_package import OfficialCoretaxPackageExporter
from core.reverse_coretax_official_package_validator import (
    OfficialCoretaxPackageValidator,
)


CATEGORIES = (
    "KAS",
    "PIUTANG",
    "INVESTASI",
    "BERGERAK",
    "HTB",
    "LAINNYA",
)


def _templates(directory: Path):
    directory.mkdir(parents=True, exist_ok=True)
    for category in CATEGORIES:
        schema = OFFICIAL_CORETAX_SCHEMAS[category]
        path = directory / f"{schema.excel_filename_hint}.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.title = "DATA"
        ws["A1"] = "NPWP"
        ws["B1"] = ""
        ws["A2"] = "Tahun Pajak"
        ws["B2"] = ""
        for index, header in enumerate(schema.excel_headers, start=1):
            ws.cell(3, index).value = header
        wb.save(path)


def _package():
    package = ReverseCoretaxPackage(
        npwp="6101015612710001",
        nama_wp="EVY BACHTIAR",
        tahun_pajak=2025,
        revision=3,
        snapshot_hash="abc123",
    )
    package.rows_by_category["KAS"].append(
        ReverseCoretaxRow(
            nomor=1,
            kategori="KAS",
            kode_harta="0102",
            nama_harta="Tabungan",
            tahun_perolehan=2025,
            nilai=391742658,
            nomor_akun_keterangan="116801005138508",
            atas_nama="DR EVY BACHTIAR SPOG",
            nama_bank="BRI",
            official_metadata={
                "account_number": "116801005138508",
                "account_on_behalf_of": "DR EVY BACHTIAR SPOG",
                "bank_name": "BRI",
                "country": "Indonesia",
                "year": 2025,
                "balance": 391742658,
            },
        )
    )
    return package


def _export(tmp_path):
    templates = tmp_path / "templates"
    output = tmp_path / "package"
    _templates(templates)

    result = OfficialCoretaxPackageExporter(
        templates
    ).export_package(
        _package(),
        output,
    )
    assert result.ok
    return output


def test_stage8d4f_accepts_clean_official_package(tmp_path):
    output = _export(tmp_path)

    result = OfficialCoretaxPackageValidator().validate(output)

    assert result.ok
    assert len(result.checked_files) == 11
    assert not result.errors


def test_stage8d4f_detects_modified_excel_hash(tmp_path):
    output = _export(tmp_path)
    manifest = json.loads(
        (output / "manifest.json").read_text(encoding="utf-8")
    )
    relative = manifest["excel"]["files"]["KAS"]["path"]
    path = output / relative

    with path.open("ab") as stream:
        stream.write(b"modified")

    result = OfficialCoretaxPackageValidator().validate(output)

    assert not result.ok
    assert any(
        issue.code == "RCX4F_EXCEL_HASH"
        for issue in result.errors
    )


def test_stage8d4f_detects_missing_xml_file(tmp_path):
    output = _export(tmp_path)
    manifest = json.loads(
        (output / "manifest.json").read_text(encoding="utf-8")
    )
    relative = manifest["xml"]["files"]["KAS"]["path"]
    (output / relative).unlink()

    result = OfficialCoretaxPackageValidator().validate(output)

    assert not result.ok
    assert any(
        issue.code == "RCX4F_XML_MISSING"
        for issue in result.errors
    )


def test_stage8d4f_detects_manifest_stage_change(tmp_path):
    output = _export(tmp_path)
    manifest_path = output / "manifest.json"
    manifest = json.loads(
        manifest_path.read_text(encoding="utf-8")
    )
    manifest["stage"] = "OTHER"
    manifest_path.write_text(
        json.dumps(manifest, indent=2),
        encoding="utf-8",
    )

    result = OfficialCoretaxPackageValidator().validate(output)

    assert not result.ok
    assert any(
        issue.code == "RCX4F_011"
        for issue in result.errors
    )


def test_stage8d4f_detects_excel_row_count_change_even_if_hash_updated(tmp_path):
    output = _export(tmp_path)
    manifest_path = output / "manifest.json"
    manifest = json.loads(
        manifest_path.read_text(encoding="utf-8")
    )
    manifest["excel"]["files"]["KAS"]["rows"] = 99
    manifest_path.write_text(
        json.dumps(manifest, indent=2),
        encoding="utf-8",
    )

    result = OfficialCoretaxPackageValidator().validate(output)

    assert not result.ok
    assert any(
        issue.code == "RCX4F_EXCEL_ROWS"
        for issue in result.errors
    )
