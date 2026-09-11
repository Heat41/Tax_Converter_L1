from openpyxl import load_workbook

from core.reverse_coretax_excel import ReverseCoretaxExcelExporter
from core.reverse_coretax_mapping import (
    ReverseCoretaxPackage,
    ReverseCoretaxRow,
)
from core.reverse_coretax_roundtrip import ReverseCoretaxRoundtripValidator


def _package():
    package = ReverseCoretaxPackage(
        npwp="1111222233334444",
        nama_wp="WP TEST",
        tahun_pajak=2025,
        revision=1,
        snapshot_hash="abc",
    )
    samples = (
        ("KAS", "0101"),
        ("PIUTANG", "0201"),
        ("INVESTASI", "0301"),
        ("BERGERAK", "0401"),
        ("HTB", "0501"),
        ("LAINNYA", "0601"),
    )
    for category, code in samples:
        package.rows_by_category[category].append(
            ReverseCoretaxRow(
                nomor=1,
                kategori=category,
                kode_harta=code,
                nama_harta=f"HARTA {category}",
                tahun_perolehan=2025,
                nilai=12_345_678,
                nomor_akun_keterangan=f"DETAIL-{category}",
                atas_nama="WP TEST",
                nama_bank="BANK TEST",
                source_eform_code="011",
            )
        )
    return package


def test_roundtrip_passes_for_unmodified_six_file_export(tmp_path):
    package = _package()
    export = ReverseCoretaxExcelExporter().export_package(package, tmp_path)
    assert export.ok is True

    result = ReverseCoretaxRoundtripValidator().validate_package(
        package,
        tmp_path,
    )

    assert result.ok is True
    assert len(result.checked_files) == 6
    assert all(result.expected_counts[key] == 1 for key in package.rows_by_category)
    assert all(result.actual_counts[key] == 1 for key in package.rows_by_category)


def test_roundtrip_accepts_header_only_empty_category(tmp_path):
    package = _package()
    package.rows_by_category["PIUTANG"] = []
    ReverseCoretaxExcelExporter().export_package(package, tmp_path)

    result = ReverseCoretaxRoundtripValidator().validate_package(
        package,
        tmp_path,
    )

    assert result.ok is True
    assert result.expected_counts["PIUTANG"] == 0
    assert result.actual_counts["PIUTANG"] == 0


def test_roundtrip_detects_changed_numeric_value(tmp_path):
    package = _package()
    export = ReverseCoretaxExcelExporter().export_package(package, tmp_path)

    path = export.files["KAS"]
    wb = load_workbook(path)
    ws = wb.active
    headers = [cell.value for cell in ws[1]]
    value_col = headers.index("harga_perolehan") + 1
    ws.cell(2, value_col, 99_999_999)
    wb.save(path)

    result = ReverseCoretaxRoundtripValidator().validate_package(
        package,
        tmp_path,
    )

    assert result.ok is False
    assert any(
        issue.code == "RCR_105" and issue.category == "KAS"
        for issue in result.errors
    )


def test_roundtrip_detects_missing_file(tmp_path):
    package = _package()
    export = ReverseCoretaxExcelExporter().export_package(package, tmp_path)
    export.files["LAINNYA"].unlink()

    result = ReverseCoretaxRoundtripValidator().validate_package(
        package,
        tmp_path,
    )

    assert result.ok is False
    assert any(
        issue.code == "RCR_101" and issue.category == "LAINNYA"
        for issue in result.errors
    )
