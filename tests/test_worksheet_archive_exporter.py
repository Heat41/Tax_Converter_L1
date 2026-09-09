from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from core.finalization import FinalizationInput
from core.mapping.worksheet_harta_mapper import WorksheetHartaRow
from core.worksheet_archive_exporter_polished import PolishedWorksheetArchiveExporter
from core.worksheet_pph_state import WorksheetBupotRow
from core.worksheet_workbook_generic import GenericWorksheetWorkbookImporter


class _Analysis:
    total_utang_sebelumnya = 0.0
    total_utang_berjalan = 0.0
    naik_turun_harta_utang = 100.0
    biaya_hidup_setahun = 10.0
    pajak_pajak = 5.0
    pengeluaran_lain_lain = 0.0
    kerugian_keuntungan_penjualan_aset = 0.0
    utang_baru_atas_kredit = 0.0
    harta_baru_dari_kredit = 0.0
    total_pengeluaran = 115.0
    penghasilan_netto = 115.0
    selisih_pengeluaran_vs_penghasilan = 0.0


def _input(*, dirty=False):
    return FinalizationInput(
        npwp="6101015612710001",
        nama_wp="WP ARSIP TEST",
        tahun_pajak=2025,
        harta_current_rows=[
            WorksheetHartaRow(
                nomor=1,
                kode_eform="014",
                kode_ct="0104",
                nama_harta="Deposito",
                nomor_akun_keterangan="001234567890123456",
                atas_nama="WP ARSIP TEST",
                nama_bank="BANK TEST",
                tahun_perolehan=2024,
                nilai_tahun_sebelumnya=100_000_000,
                nilai_tahun_berjalan=125_000_000,
            )
        ],
        harta_original_hash="hash",
        bupot_rows=[
            WorksheetBupotRow(
                jenis="BP21",
                npwp_pemberi_kerja="0011050945093000",
                no_bupot="BUPOT-001",
                bruto=10_000_000,
                pengurang=1_000_000,
            )
        ],
        pph_components={"ptkp": 54_000_000, "kredit_pajak": 500_000, "pph25": 0},
        umkm_state={"bruto_bulanan": [0.0] * 12, "pph_setor_bulanan": [0.0] * 12},
        penghasilan_lainnya={
            "domestic_other_enabled": False,
            "domestic_other_dpp": 0.0,
            "zakat": 0.0,
            "final_other_rows": [],
        },
        zakat=0.0,
        status_ptkp="TK/0",
        pph_calc_result={"pph_terutang": 750_000},
        analisis_result=_Analysis(),
        is_harta_dirty=dirty,
        is_pph_dirty=False,
        is_analisis_dirty=False,
    )


def test_export_excel_is_round_trip_compatible(tmp_path: Path):
    path = tmp_path / "arsip.xlsx"
    result = PolishedWorksheetArchiveExporter().export_excel(_input(), path)

    assert result.output_path == path
    assert path.exists()
    assert result.bupot_count == 1
    assert result.harta_count == 1

    wb = load_workbook(path)
    annual = wb["2025"]
    simulasi = wb["SIMULASI I"]
    assert "TblBupotArsip" in annual.tables
    assert "TblUmkmArsip" in annual.tables
    assert "TblHartaArsip" in simulasi.tables
    assert annual.sheet_view.showGridLines is False
    assert simulasi.sheet_view.showGridLines is False

    # Identifiers tidak boleh berubah menjadi scientific notation / angka.
    assert annual["B5"].value == "6101015612710001"
    assert annual["B5"].number_format == "@"
    assert annual["C9"].value == "0011050945093000"
    assert annual["C9"].number_format == "@"
    assert simulasi["E9"].value == "001234567890123456"
    assert simulasi["E9"].number_format == "@"
    assert simulasi.column_dimensions["D"].width >= 30
    assert simulasi.column_dimensions["E"].width >= 25

    imported = GenericWorksheetWorkbookImporter().parse(path)
    assert imported.is_valid
    assert imported.npwp == "6101015612710001"
    assert imported.tahun_pajak == 2025
    assert len(imported.bupot_rows) == 1
    assert imported.bupot_rows[0].no_bupot == "BUPOT-001"
    assert imported.bupot_rows[0].npwp_pemberi_kerja == "0011050945093000"
    assert len(imported.harta_rows) == 1
    assert imported.harta_rows[0].kode_ct == "0104"
    assert imported.harta_rows[0].nomor_akun_keterangan == "001234567890123456"
    assert imported.harta_rows[0].nilai_tahun_berjalan == 125_000_000


def test_export_is_blocked_when_saved_state_is_dirty(tmp_path: Path):
    try:
        PolishedWorksheetArchiveExporter().export_excel(_input(dirty=True), tmp_path / "dirty.xlsx")
    except ValueError as exc:
        assert "Simpan perubahan" in str(exc)
    else:
        raise AssertionError("Export seharusnya diblokir untuk state dirty")
