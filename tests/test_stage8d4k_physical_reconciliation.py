from pathlib import Path
from xml.etree import ElementTree as ET

from openpyxl import Workbook

from core.coretax_official_schema import OFFICIAL_CORETAX_SCHEMAS
from core.physical_reconciliation import PhysicalSourceExportReconciler


def _write_excel(directory: Path, category: str, rows):
    directory.mkdir(parents=True, exist_ok=True)
    schema = OFFICIAL_CORETAX_SCHEMAS[category]
    path = directory / f"{schema.excel_filename_hint}.xlsx"

    wb = Workbook()
    ws = wb.active
    ws.title = schema.excel_sheet
    for col, header in enumerate(schema.excel_headers, start=1):
        ws.cell(3, col).value = header
    for row_index, values in enumerate(rows, start=4):
        for col, value in enumerate(values, start=1):
            ws.cell(row_index, col).value = value
    wb.save(path)
    return path


def _write_xml(directory: Path, category: str, rows):
    directory.mkdir(parents=True, exist_ok=True)
    schema = OFFICIAL_CORETAX_SCHEMAS[category]
    root = ET.Element(schema.xml_root)
    tin = ET.SubElement(root, schema.xml_tin_field)
    tin.text = "6101015612710001"
    year = ET.SubElement(root, schema.xml_year_field)
    year.text = "2025"
    container = ET.SubElement(root, schema.xml_list)
    for values in rows:
        item = ET.SubElement(container, schema.xml_item)
        for field_name, value in zip(schema.xml_fields, values):
            node = ET.SubElement(item, field_name)
            node.text = "" if value is None else str(value)
    path = directory / f"{schema.xml_root}.xml"
    ET.ElementTree(root).write(path, encoding="utf-8", xml_declaration=True)
    return path


def test_stage8d4k_source_and_export_match(tmp_path):
    source = tmp_path / "source"
    exported = tmp_path / "export"

    excel_rows = [
        (
            "0102",
            "12345",
            "EVY",
            "BANK",
            "Indonesia",
            2025,
            1000000,
            "",
        )
    ]
    xml_rows = [
        (
            "0102",
            "12345",
            "EVY",
            "BANK",
            "Indonesia",
            2025,
            1000000,
        )
    ]

    _write_excel(source, "KAS", excel_rows)
    _write_excel(exported, "KAS", excel_rows)
    _write_xml(source, "KAS", xml_rows)
    _write_xml(exported, "KAS", xml_rows)

    result = PhysicalSourceExportReconciler().reconcile(
        source, exported
    )

    assert result.ok
    assert result.categories["KAS"].excel_match is True
    assert result.categories["KAS"].xml_match is True
    assert not result.errors


def test_stage8d4k_numeric_representation_does_not_create_false_difference(tmp_path):
    source = tmp_path / "source"
    exported = tmp_path / "export"

    source_rows = [
        (
            "0102",
            "12345",
            "EVY",
            "BANK",
            "Indonesia",
            "2025",
            "1000000.0",
            "",
        )
    ]
    export_rows = [
        (
            "0102",
            "12345",
            "EVY",
            "BANK",
            "Indonesia",
            2025,
            1000000,
            None,
        )
    ]

    _write_excel(source, "KAS", source_rows)
    _write_excel(exported, "KAS", export_rows)

    result = PhysicalSourceExportReconciler().reconcile(
        source, exported
    )

    assert result.ok
    assert result.categories["KAS"].excel_match is True


def test_stage8d4k_reports_exact_different_field(tmp_path):
    source = tmp_path / "source"
    exported = tmp_path / "export"

    source_rows = [
        (
            "0102",
            "12345",
            "EVY",
            "BANK A",
            "Indonesia",
            2025,
            1000000,
            "",
        )
    ]
    export_rows = [
        (
            "0102",
            "12345",
            "EVY",
            "BANK B",
            "Indonesia",
            2025,
            1000000,
            "",
        )
    ]

    _write_excel(source, "KAS", source_rows)
    _write_excel(exported, "KAS", export_rows)

    result = PhysicalSourceExportReconciler().reconcile(
        source, exported
    )

    assert not result.ok
    assert result.categories["KAS"].excel_match is False
    assert any(
        issue.code == "RCX4K_EXCEL_VALUE"
        and issue.category == "KAS"
        and issue.row_number == 1
        and issue.field_name == "NAMA BANK/ INSTITUSI*"
        for issue in result.errors
    )


def test_stage8d4k_absent_category_on_both_sides_is_ignored(tmp_path):
    source = tmp_path / "source"
    exported = tmp_path / "export"

    rows = [
        (
            "0501",
            "Pontianak",
            100,
            50,
            "Own Income",
            "SHM-001",
            2021,
            80000000,
            100000000,
            "",
        )
    ]

    _write_excel(source, "HTB", rows)
    _write_excel(exported, "HTB", rows)

    result = PhysicalSourceExportReconciler().reconcile(
        source, exported
    )

    assert result.ok
    assert set(result.categories) == {"HTB"}
    assert "KAS" not in result.categories
    assert "PIUTANG" not in result.categories
