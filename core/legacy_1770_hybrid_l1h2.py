from __future__ import annotations

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


class LegacyLampiranIH2XlsxRenderer:
    """Renderer XLSX Lampiran I Halaman 2 format lama (Bagian B/C/D)."""

    # Format lama DJP: dominan putih/monokrom dengan area isian kuning muda.
    # Hindari warna biru e-Form agar Lampiran I H2 tetap jelas legacy.
    FILL = PatternFill("solid", fgColor="FFFFFF")
    LIGHT = PatternFill("solid", fgColor="FFFFFF")
    VALUE_FILL = PatternFill("solid", fgColor="FFF2CC")
    THIN = Side(style="thin", color="7F8C8D")
    BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
    MONEY = '#,##0;[Red]-#,##0;-'

    @staticmethod
    def _num(value) -> float:
        try:
            return float(value or 0)
        except (TypeError, ValueError):
            return 0.0

    @staticmethod
    def _money(value) -> int:
        try:
            return int(round(float(value or 0)))
        except (TypeError, ValueError):
            return 0

    def render(self, ws, data) -> None:
        row = self._header(ws, data)
        row = self._bagian_b(ws, row)
        row = self._bagian_c(ws, row, data)
        row = self._bagian_d(ws, row, data)
        self._finish(ws, row)

    def _header(self, ws, data) -> int:
        ws.sheet_view.showGridLines = False
        ws.merge_cells("A1:C2")
        ws["A1"] = "KEMENTERIAN KEUANGAN RI\nDIREKTORAT JENDERAL PAJAK"
        ws["A1"].font = Font(size=8, bold=True)
        ws["A1"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        ws.merge_cells("D1:H2")
        ws["D1"] = "SPT TAHUNAN PPh WAJIB PAJAK ORANG PRIBADI"
        ws["D1"].font = Font(size=11, bold=True)
        ws["D1"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        ws.merge_cells("I1:J2")
        ws["I1"] = "FORMULIR\n1770 - I"
        ws["I1"].font = Font(size=10, bold=True)
        ws["I1"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        ws.merge_cells("A3:G3")
        ws["A3"] = "LAMPIRAN - I"
        ws["A3"].font = Font(size=10, bold=True)
        ws["A3"].alignment = Alignment(horizontal="left")
        ws.merge_cells("H3:J3")
        ws["H3"] = "HALAMAN 2"
        ws["H3"].font = Font(size=9, bold=True)
        ws["H3"].alignment = Alignment(horizontal="right")

        ws.merge_cells("A4:J4")
        ws["A4"] = (
            "PENGHITUNGAN PENGHASILAN NETO DALAM NEGERI DARI USAHA DAN/ATAU "
            "PEKERJAAN BEBAS, PEKERJAAN, DAN PENGHASILAN DALAM NEGERI LAINNYA"
        )
        ws["A4"].font = Font(size=9, bold=True)
        ws["A4"].alignment = Alignment(horizontal="center", wrap_text=True)
        ws.row_dimensions[4].height = 30

        ws.merge_cells("A6:E6")
        ws["A6"] = f"NPWP : {data.npwp}"
        ws.merge_cells("F6:J6")
        ws["F6"] = f"NAMA WAJIB PAJAK : {str(data.nama_wp or '').upper()}"
        ws["A6"].font = Font(bold=True)
        ws["F6"].font = Font(bold=True)

        ws.merge_cells("A7:C7")
        ws["A7"] = f"TAHUN PAJAK : {data.tahun_pajak}"
        ws.merge_cells("D7:F7")
        ws["D7"] = "01 s.d 12"
        ws.merge_cells("G7:J7")
        ws["G7"] = "PEMBUKUAN / PENCATATAN"
        for key in ("A7", "D7", "G7"):
            ws[key].font = Font(bold=True)
            ws[key].alignment = Alignment(horizontal="center")
        return 9

    def _section(self, ws, row: int, title: str, note: str) -> int:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=10)
        cell = ws.cell(row, 1)
        cell.value = title
        cell.font = Font(size=9, bold=True)
        cell.fill = self.FILL
        cell.border = self.BORDER
        cell.alignment = Alignment(wrap_text=True)
        row += 1
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=10)
        ws.cell(row, 1).value = note
        ws.cell(row, 1).font = Font(size=8, italic=True)
        ws.cell(row, 1).alignment = Alignment(wrap_text=True)
        return row + 1

    def _table_header(self, ws, row: int, headers, spans) -> int:
        for value, (c1, c2) in zip(headers, spans):
            if c1 != c2:
                ws.merge_cells(start_row=row, start_column=c1, end_row=row, end_column=c2)
            cell = ws.cell(row, c1)
            cell.value = value
            cell.font = Font(size=8, bold=True)
            cell.fill = self.LIGHT
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            for col in range(c1, c2 + 1):
                ws.cell(row, col).border = self.BORDER
        ws.row_dimensions[row].height = 38
        return row + 1

    def _write_spanned(self, ws, row, values, spans, numeric_starts=()):
        for value, (c1, c2) in zip(values, spans):
            if c1 != c2:
                ws.merge_cells(start_row=row, start_column=c1, end_row=row, end_column=c2)
            cell = ws.cell(row, c1)
            cell.value = value
            cell.alignment = Alignment(
                horizontal="right" if c1 in numeric_starts else "left",
                vertical="center",
                wrap_text=True,
            )
            if c1 in numeric_starts:
                cell.number_format = self.MONEY
                cell.fill = self.VALUE_FILL
            for col in range(c1, c2 + 1):
                ws.cell(row, col).border = self.BORDER

    def _bagian_b(self, ws, row: int) -> int:
        row = self._section(
            ws, row,
            "BAGIAN B : PENGHASILAN NETO DALAM NEGERI DARI USAHA DAN/ATAU PEKERJAAN BEBAS",
            "Pindahkan Jumlah Bagian B Kolom (5) ke Formulir 1770 Angka 1.",
        )
        spans = ((1, 1), (2, 4), (5, 6), (7, 8), (9, 10))
        row = self._table_header(
            ws, row,
            ("NO.", "JENIS USAHA", "PEREDARAN USAHA (Rupiah)", "NORMA (%)", "PENGHASILAN NETO (Rupiah)"),
            spans,
        )
        first_data_row = row
        for idx, jenis in enumerate(("DAGANG", "INDUSTRI", "JASA", "PEKERJAAN BEBAS"), start=1):
            self._write_spanned(ws, row, (idx, jenis, 0, "", 0), spans, {5, 9})
            row += 1

        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
        ws.cell(row, 1).value = "JUMLAH BAGIAN B"
        ws.cell(row, 1).font = Font(bold=True)
        ws.cell(row, 1).alignment = Alignment(horizontal="right")
        ws.merge_cells(start_row=row, start_column=9, end_row=row, end_column=10)
        ws.cell(row, 9).value = f"=SUM(I{first_data_row}:I{row - 1})"
        ws.cell(row, 9).number_format = self.MONEY
        ws.cell(row, 9).font = Font(bold=True)
        ws.cell(row, 9).fill = self.VALUE_FILL
        for col in range(1, 11):
            ws.cell(row, col).border = self.BORDER
        return row + 2

    def _bagian_c(self, ws, row: int, data) -> int:
        row = self._section(
            ws, row,
            "BAGIAN C : PENGHASILAN NETO DALAM NEGERI SEHUBUNGAN DENGAN PEKERJAAN",
            "(Tidak termasuk penghasilan yang dikenakan PPh bersifat final). Pindahkan Jumlah Bagian C Kolom (5) ke Formulir 1770 Angka 2.",
        )
        spans = ((1, 1), (2, 4), (5, 6), (7, 8), (9, 10))
        row = self._table_header(
            ws, row,
            ("NO.", "NAMA DAN NPWP PEMBERI KERJA", "PENGHASILAN BRUTO (Rupiah)", "PENGURANGAN PENGHASILAN BRUTO/BIAYA (Rupiah)", "PENGHASILAN NETO (Rupiah)"),
            spans,
        )

        entries = []
        for item in (data.bupot_rows or []):
            bruto = self._num(item.bruto)
            pengurang = self._num(item.pengurang)
            npwp = str(item.npwp_pemotong or item.npwp_pemberi_kerja or "")
            nama = str(item.nama_pemotong or "")
            if any((bruto, pengurang, npwp, nama, str(item.no_bupot or ""))):
                entries.append(item)

        total_bruto = total_pengurang = total_netto = 0.0
        if not entries:
            entries = [None]

        first_data_row = row
        for idx, item in enumerate(entries, start=1):
            if item is None:
                values = ("", "", 0, 0, 0)
            else:
                bruto = self._num(item.bruto)
                pengurang = self._num(item.pengurang)
                netto = bruto - pengurang
                npwp = str(item.npwp_pemotong or item.npwp_pemberi_kerja or "")
                nama = str(item.nama_pemotong or "").strip()
                values = (idx, "\n".join(v for v in (nama, npwp) if v), bruto, pengurang, netto)
                total_bruto += bruto
                total_pengurang += pengurang
                total_netto += netto
            self._write_spanned(ws, row, values, spans, {5, 7, 9})
            # Neto mengikuti koreksi Bruto/Pengurang langsung di Excel.
            ws.cell(row, 9).value = f"=E{row}-G{row}"
            ws.row_dimensions[row].height = 32
            row += 1

        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
        ws.cell(row, 1).value = "JUMLAH BAGIAN C"
        ws.cell(row, 1).font = Font(bold=True)
        ws.cell(row, 1).alignment = Alignment(horizontal="right")
        for c1, c2, source_col in ((5, 6, "E"), (7, 8, "G"), (9, 10, "I")):
            ws.merge_cells(start_row=row, start_column=c1, end_row=row, end_column=c2)
            ws.cell(row, c1).value = f"=SUM({source_col}{first_data_row}:{source_col}{row - 1})"
            ws.cell(row, c1).number_format = self.MONEY
            ws.cell(row, c1).font = Font(bold=True)
            ws.cell(row, c1).fill = self.VALUE_FILL
            ws.cell(row, c1).alignment = Alignment(horizontal="right")
        for col in range(1, 11):
            ws.cell(row, col).border = self.BORDER
        return row + 2

    def _bagian_d(self, ws, row: int, data) -> int:
        row = self._section(
            ws, row,
            "BAGIAN D : PENGHASILAN NETO DALAM NEGERI LAINNYA",
            "(Tidak termasuk penghasilan yang dikenakan PPh bersifat final). Pindahkan Jumlah Bagian D ke Formulir 1770 Angka 3.",
        )
        spans = ((1, 2), (3, 7), (8, 10))
        row = self._table_header(ws, row, ("NO.", "JENIS PENGHASILAN", "PENGHASILAN NETO (Rupiah)"), spans)

        calc = data.pph_calc_result or {}
        components = data.pph_components or {}
        domestic_other = self._num(calc.get("penghasilan_neto_lainnya", components.get("penghasilan_neto_lainnya", 0)))
        kinds = (
            "BUNGA",
            "ROYALTI",
            "SEWA",
            "PENGHARGAAN DAN HADIAH",
            "KEUNTUNGAN DARI PENJUALAN/PENGALIHAN HARTA",
            "PENGHASILAN LAINNYA",
        )
        first_data_row = row
        for idx, kind in enumerate(kinds, start=1):
            value = domestic_other if kind == "PENGHASILAN LAINNYA" else 0
            self._write_spanned(ws, row, (idx, kind, self._money(value)), spans, {8})
            row += 1

        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=7)
        ws.cell(row, 1).value = "JUMLAH BAGIAN D"
        ws.cell(row, 1).font = Font(bold=True)
        ws.cell(row, 1).alignment = Alignment(horizontal="right")
        ws.merge_cells(start_row=row, start_column=8, end_row=row, end_column=10)
        ws.cell(row, 8).value = f"=SUM(H{first_data_row}:H{row - 1})"
        ws.cell(row, 8).number_format = self.MONEY
        ws.cell(row, 8).font = Font(bold=True)
        ws.cell(row, 8).fill = self.VALUE_FILL
        ws.cell(row, 8).alignment = Alignment(horizontal="right")
        for col in range(1, 11):
            ws.cell(row, col).border = self.BORDER
        return row

    def _finish(self, ws, row: int) -> None:
        # Proporsi mengikuti karakter form legacy asli:
        # nomor sempit, uraian identitas/jenis penghasilan lebih lebar,
        # kolom angka cukup untuk nilai rupiah tanpa menyisakan ruang berlebih.
        widths = {
            "A": 4.0,
            "B": 5.0,
            "C": 12.0,
            "D": 12.0,
            "E": 11.0,
            "F": 11.0,
            "G": 11.0,
            "H": 11.0,
            "I": 11.0,
            "J": 11.0,
        }
        for col, width in widths.items():
            ws.column_dimensions[col].width = width
        ws.page_setup.orientation = "portrait"
        ws.page_setup.paperSize = ws.PAPERSIZE_LEGAL
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 1
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_margins.left = 0.2
        ws.page_margins.right = 0.2
        ws.page_margins.top = 0.3
        ws.page_margins.bottom = 0.3
        ws.print_area = f"A1:J{row}"
