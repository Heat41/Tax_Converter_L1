from copy import copy
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable, List, Optional, Tuple

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from core.mapping.worksheet_harta_mapper import WorksheetHartaRow


@dataclass(frozen=True)
class SimulasiHartaWriteResult:
    output_path: Path
    sheet_name: str
    header_row: int
    start_row: int
    written_rows: int
    previous_year: int
    current_year: int


class SimulasiHartaExcelWriter:
    """Writer untuk tabel harta workbook kertas kerja nyata (sheet SIMULASI I)."""

    REQUIRED_HEADERS = (
        "NO",
        "KODE EFORM",
        "KODE CT",
        "NAMA HARTA",
        "NOMOR AKUN / KETERANGAN",
        "ATAS NAMA",
        "NAMA BANK",
        "TH PEROLEHAN",
    )

    def write(
        self,
        template_path: Path | str,
        output_path: Path | str,
        rows: Iterable[WorksheetHartaRow],
        *,
        current_year: int,
        previous_year: Optional[int] = None,
        sheet_name: str = "SIMULASI I",
    ) -> SimulasiHartaWriteResult:
        template = Path(template_path).resolve()
        output = Path(output_path).resolve()
        if not template.exists():
            raise FileNotFoundError(f"Kertas kerja tidak ditemukan: '{template}'")

        previous_year = int(previous_year if previous_year is not None else current_year - 1)
        data = list(rows)

        wb = load_workbook(template)
        if sheet_name not in wb.sheetnames:
            raise ValueError(f"Sheet '{sheet_name}' tidak ditemukan.")
        ws = wb[sheet_name]

        header_row = self._find_header_row(ws)
        start_row = header_row + 1
        total_row = self._find_grand_total_row(ws, start_row)
        if total_row is None:
            raise ValueError("Baris Grand Total tabel harta tidak ditemukan.")

        capacity = total_row - start_row
        if len(data) > capacity:
            extra = len(data) - capacity
            style_source = total_row - 1
            ws.insert_rows(total_row, amount=extra)
            self._copy_row_style(ws, style_source, total_row, extra, max_col=17)
            total_row += extra

        ws.cell(header_row, 9).value = previous_year
        ws.cell(header_row, 10).value = current_year

        self._clear_data_area(ws, start_row, total_row - 1)
        for offset, item in enumerate(data):
            row_idx = start_row + offset
            values = (
                item.nomor,
                item.kode_eform,
                item.kode_ct,
                item.nama_harta,
                item.nomor_akun_keterangan,
                item.atas_nama,
                item.nama_bank,
                item.tahun_perolehan,
                item.nilai_tahun_sebelumnya,
                item.nilai_tahun_berjalan,
            )
            for col_idx, value in enumerate(values, start=1):
                ws.cell(row_idx, col_idx).value = value
            self._write_category_formulas(ws, row_idx)

        # Sisakan nomor urut pada row kosong agar perilaku template tetap familiar.
        for row_idx in range(start_row + len(data), total_row):
            ws.cell(row_idx, 1).value = row_idx - start_row + 1
            self._write_category_formulas(ws, row_idx)

        ws.cell(total_row, 1).value = "Grand Total"
        ws.cell(total_row, 9).value = f"=SUBTOTAL(9,I{start_row}:I{total_row - 1})"
        ws.cell(total_row, 10).value = f"=SUBTOTAL(9,J{start_row}:J{total_row - 1})"

        output.parent.mkdir(parents=True, exist_ok=True)
        wb.save(output)
        return SimulasiHartaWriteResult(
            output_path=output,
            sheet_name=sheet_name,
            header_row=header_row,
            start_row=start_row,
            written_rows=len(data),
            previous_year=previous_year,
            current_year=int(current_year),
        )

    def _find_header_row(self, ws: Worksheet) -> int:
        expected = [self._norm(value) for value in self.REQUIRED_HEADERS]
        for row_idx in range(1, min(ws.max_row, 100) + 1):
            found = [self._norm(ws.cell(row_idx, col).value) for col in range(1, 9)]
            if found == expected:
                return row_idx
        raise ValueError("Header tabel harta pada sheet SIMULASI I tidak ditemukan.")

    @staticmethod
    def _find_grand_total_row(ws: Worksheet, start_row: int) -> Optional[int]:
        for row_idx in range(start_row, ws.max_row + 1):
            value = str(ws.cell(row_idx, 1).value or "").strip().lower()
            if value == "grand total":
                return row_idx
        return None

    @staticmethod
    def _clear_data_area(ws: Worksheet, start_row: int, end_row: int) -> None:
        for row_idx in range(start_row, end_row + 1):
            for col_idx in range(1, 18):
                ws.cell(row_idx, col_idx).value = None

    @staticmethod
    def _write_category_formulas(ws: Worksheet, row_idx: int) -> None:
        detail = f'$D{row_idx}&"; "&$E{row_idx}&"; "&$F{row_idx}&"; "&$G{row_idx}'
        ws.cell(row_idx, 12).value = f'=IF(LEFT($C{row_idx},2)="01",{detail},"")'
        ws.cell(row_idx, 13).value = f'=IF(LEFT($C{row_idx},2)="02",{detail},"")'
        ws.cell(row_idx, 14).value = f'=IF(LEFT($C{row_idx},2)="03",{detail},"")'
        ws.cell(row_idx, 15).value = f'=IF(LEFT($C{row_idx},2)="04",{detail},"")'
        ws.cell(row_idx, 16).value = f'=IF(LEFT($C{row_idx},2)="05",{detail},"")'
        ws.cell(row_idx, 17).value = (
            f'=IF(OR(LEFT($C{row_idx},2)="06",LEFT($C{row_idx},2)="07"),{detail},"")'
        )

    @staticmethod
    def _copy_row_style(ws: Worksheet, source_row: int, target_start: int, count: int, *, max_col: int) -> None:
        for offset in range(count):
            target_row = target_start + offset
            for col in range(1, max_col + 1):
                source = ws.cell(source_row, col)
                target = ws.cell(target_row, col)
                if source.has_style:
                    target._style = copy(source._style)
                if source.number_format:
                    target.number_format = source.number_format
            if source_row in ws.row_dimensions:
                ws.row_dimensions[target_row].height = ws.row_dimensions[source_row].height

    @staticmethod
    def _norm(value: object) -> str:
        return " ".join(str(value or "").replace("\n", " ").strip().upper().split())
