from __future__ import annotations

from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from PySide6.QtCore import Qt, QRectF
from PySide6.QtGui import QBrush, QColor, QFont, QPainter, QPen, QPdfWriter

from core.finalization import FinalizationInput
from core.worksheet_archive_exporter import (
    WorksheetArchiveExporter,
    WorksheetArchiveExportResult,
)


class StyledWorksheetArchiveExporter(WorksheetArchiveExporter):
    """Presentation layer Stage 8B.2 untuk arsip Excel/PDF yang lebih rapi.

    Data dan aturan export tetap memakai ``WorksheetArchiveExporter``. Kelas ini
    hanya memperkaya hasil dengan tabel, border, format angka, section header,
    serta layout PDF yang terstruktur. Excel tetap kompatibel dengan importer 8B.1.
    """

    THIN = Side(style="thin", color="B7C3CF")
    BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
    HEADER_FILL = PatternFill("solid", fgColor="D9EAF7")
    SECTION_FILL = PatternFill("solid", fgColor="EAF2F8")
    TOTAL_FILL = PatternFill("solid", fgColor="EEF2F5")
    MONEY_FORMAT = '#,##0;[Red]-#,##0'

    def export_excel(
        self,
        data: FinalizationInput,
        output_path: str | Path,
    ) -> WorksheetArchiveExportResult:
        result = super().export_excel(data, output_path)
        self._style_excel(result.output_path, data)
        return result

    def _style_excel(self, path: Path, data: FinalizationInput) -> None:
        wb = load_workbook(path)
        annual = wb[str(data.tahun_pajak)]
        simulasi = wb["SIMULASI I"]
        ref = wb["REF"]

        self._style_annual(annual, data)
        self._style_simulasi(simulasi, data)
        self._style_ref(ref)
        wb.save(path)

    def _style_annual(self, ws, data: FinalizationInput) -> None:
        ws.sheet_view.showGridLines = False
        ws.merge_cells("A1:G1")
        ws["A1"].alignment = Alignment(horizontal="center")
        ws["A1"].font = Font(bold=True, size=15)
        ws["A2"].font = Font(bold=True, size=11)

        for row in (4, 5):
            for col in range(1, 3):
                ws.cell(row, col).border = self.BORDER
                ws.cell(row, col).alignment = Alignment(vertical="center")
            ws.cell(row, 1).font = Font(bold=True)
            ws.cell(row, 1).fill = self.SECTION_FILL

        bupot_count = len(data.bupot_rows)
        bupot_header = 8
        bupot_last = bupot_header + bupot_count
        if bupot_count:
            self._add_table(ws, f"A{bupot_header}:G{bupot_last}", "TblBupotArsip")
        self._style_grid(ws, bupot_header, bupot_last + 1, 1, 7, money_cols={5, 6, 7})
        for cell in ws[bupot_header]:
            if cell.column <= 7:
                cell.font = Font(bold=True)
                cell.fill = self.HEADER_FILL
        total_row = bupot_last + 1
        for col in range(1, 8):
            ws.cell(total_row, col).fill = self.TOTAL_FILL
            ws.cell(total_row, col).font = Font(bold=True)

        umkm_section = total_row + 3
        umkm_header = umkm_section + 1
        umkm_last = umkm_header + 12
        self._section_row(ws, umkm_section, 1, 6)
        self._add_table(ws, f"A{umkm_header}:F{umkm_last}", "TblUmkmArsip")
        self._style_grid(ws, umkm_header, umkm_last, 1, 6, money_cols={4, 6})

        # Rapikan seluruh area penghasilan/PPh yang ditulis setelah UMKM.
        for row in range(umkm_last + 1, ws.max_row + 1):
            label = str(ws.cell(row, 1).value or ws.cell(row, 2).value or "").strip()
            if label:
                for col in range(1, 8):
                    ws.cell(row, col).border = self.BORDER
                    ws.cell(row, col).alignment = Alignment(vertical="center", wrap_text=True)
                if label.upper() in {"PENGHASILAN LAINNYA", "PENGURANG PENGHASILAN NETO"}:
                    self._section_row(ws, row, 1, 7)
            for col in (5, 6):
                ws.cell(row, col).number_format = self.MONEY_FORMAT

        ws.auto_filter.ref = f"A8:G{max(total_row, 8)}"
        ws.freeze_panes = "A8"
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_title_rows = "1:8"

    def _style_simulasi(self, ws, data: FinalizationInput) -> None:
        ws.sheet_view.showGridLines = False
        ws.merge_cells("A1:J1")
        ws["A1"].alignment = Alignment(horizontal="center")
        ws["A1"].font = Font(bold=True, size=15)

        for row in (3, 4):
            for col in range(1, 4):
                ws.cell(row, col).border = self.BORDER
                ws.cell(row, col).alignment = Alignment(vertical="center")
            ws.cell(row, 1).fill = self.SECTION_FILL
            ws.cell(row, 1).font = Font(bold=True)

        harta_count = len(data.harta_current_rows)
        harta_header = 8
        harta_last = harta_header + harta_count
        if harta_count:
            self._add_table(ws, f"A{harta_header}:J{harta_last}", "TblHartaArsip")
        self._style_grid(ws, harta_header, harta_last + 1, 1, 10, money_cols={9, 10})
        total_row = harta_last + 1
        for col in range(1, 11):
            ws.cell(total_row, col).fill = self.TOTAL_FILL
            ws.cell(total_row, col).font = Font(bold=True)

        for row in range(total_row + 1, ws.max_row + 1):
            text = " ".join(str(ws.cell(row, col).value or "") for col in range(1, 3)).strip()
            if not text:
                continue
            for col in range(1, 11):
                ws.cell(row, col).border = self.BORDER
                ws.cell(row, col).alignment = Alignment(vertical="center", wrap_text=True)
            ws.cell(row, 10).number_format = self.MONEY_FORMAT
            if text.upper().startswith("UTANG") or text.upper().startswith("PERHITUNGAN PENGHASILAN"):
                self._section_row(ws, row, 1, 10)

        ws.freeze_panes = "A8"
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_title_rows = "1:8"

    def _style_ref(self, ws) -> None:
        ws.sheet_view.showGridLines = False
        for row in range(1, ws.max_row + 1):
            for col in range(1, min(ws.max_column, 2) + 1):
                ws.cell(row, col).border = self.BORDER
                ws.cell(row, col).alignment = Alignment(vertical="top", wrap_text=True)
        ws["A1"].fill = self.HEADER_FILL
        ws["A1"].font = Font(bold=True, size=14)

    def _style_grid(self, ws, start_row, end_row, start_col, end_col, *, money_cols=None):
        money_cols = set(money_cols or set())
        for row in range(start_row, end_row + 1):
            for col in range(start_col, end_col + 1):
                cell = ws.cell(row, col)
                cell.border = self.BORDER
                cell.alignment = Alignment(vertical="center", wrap_text=True)
                if col in money_cols and row > start_row:
                    cell.number_format = self.MONEY_FORMAT
        ws.row_dimensions[start_row].height = 30

    def _section_row(self, ws, row: int, start_col: int, end_col: int) -> None:
        for col in range(start_col, end_col + 1):
            cell = ws.cell(row, col)
            cell.fill = self.SECTION_FILL
            cell.font = Font(bold=True)
            cell.border = self.BORDER

    @staticmethod
    def _add_table(ws, ref: str, name: str) -> None:
        if name in ws.tables:
            return
        table = Table(displayName=name, ref=ref)
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        ws.add_table(table)

    # ------------------------------------------------------------------ PDF
    def _paint_pdf(self, painter: QPainter, writer: QPdfWriter, data: FinalizationInput) -> None:
        page = writer.pageLayout().paintRectPixels(writer.resolution())
        margin = 28.0
        usable_width = float(page.width()) - margin * 2
        bottom = float(page.height()) - margin
        y = margin

        pen = QPen(QColor(90, 100, 110))
        pen.setWidthF(0.7)
        painter.setPen(pen)

        def new_page():
            nonlocal y
            writer.newPage()
            y = margin

        def ensure(height: float):
            if y + height > bottom:
                new_page()

        def title(text: str, size=14, height=28):
            nonlocal y
            ensure(height)
            font = QFont("Arial", size)
            font.setBold(True)
            painter.setFont(font)
            painter.drawText(QRectF(margin, y, usable_width, height), Qt.AlignCenter | Qt.AlignVCenter, text)
            y += height

        def section(text: str):
            nonlocal y
            ensure(24)
            rect = QRectF(margin, y, usable_width, 22)
            painter.fillRect(rect, QBrush(QColor(232, 238, 243)))
            painter.drawRect(rect)
            font = QFont("Arial", 9)
            font.setBold(True)
            painter.setFont(font)
            painter.drawText(rect.adjusted(6, 0, -4, 0), Qt.AlignLeft | Qt.AlignVCenter, text)
            y += 22

        def table(headers: Iterable[str], rows: Iterable[Iterable[Any]], widths: Iterable[float], *, row_height=21.0):
            nonlocal y
            headers = [str(v) for v in headers]
            rows = [["" if v is None else str(v) for v in row] for row in rows]
            fractions = list(widths)
            total = sum(fractions) or 1.0
            widths_px = [usable_width * value / total for value in fractions]

            def draw_header():
                nonlocal y
                ensure(row_height)
                x = margin
                painter.setFont(QFont("Arial", 7, QFont.Bold))
                for idx, value in enumerate(headers):
                    rect = QRectF(x, y, widths_px[idx], row_height)
                    painter.fillRect(rect, QBrush(QColor(217, 234, 247)))
                    painter.drawRect(rect)
                    painter.drawText(rect.adjusted(3, 1, -3, -1), Qt.AlignCenter | Qt.AlignVCenter, value)
                    x += widths_px[idx]
                y += row_height

            draw_header()
            painter.setFont(QFont("Arial", 7))
            for row in rows:
                if y + row_height > bottom:
                    new_page()
                    draw_header()
                    painter.setFont(QFont("Arial", 7))
                x = margin
                for idx, value in enumerate(row):
                    rect = QRectF(x, y, widths_px[idx], row_height)
                    painter.drawRect(rect)
                    align = Qt.AlignRight | Qt.AlignVCenter if self._looks_numeric(value) else Qt.AlignLeft | Qt.AlignVCenter
                    painter.drawText(rect.adjusted(3, 1, -3, -1), align, self._clip_text(value, widths_px[idx]))
                    x += widths_px[idx]
                y += row_height

        title("KERTAS KERJA SPT TAHUNAN")
        title(f"TAHUN PAJAK {data.tahun_pajak}", 10, 20)
        table(
            ["IDENTITAS", "KETERANGAN"],
            [["Nama Wajib Pajak", data.nama_wp], ["NPWP", data.npwp]],
            [1.1, 3.4],
            row_height=20,
        )
        y += 8

        section("BUKTI POTONG / PENGHASILAN PEKERJAAN")
        bupot_rows = []
        total_bruto = total_pengurang = 0.0
        for no, item in enumerate(data.bupot_rows, start=1):
            bruto = self._money(item.bruto)
            pengurang = self._money(item.pengurang)
            total_bruto += bruto
            total_pengurang += pengurang
            bupot_rows.append([
                no,
                item.jenis,
                item.npwp_pemberi_kerja,
                item.no_bupot,
                self._money_text(bruto),
                self._money_text(pengurang),
                self._money_text(bruto - pengurang),
            ])
        bupot_rows.append([
            "",
            "TOTAL",
            "",
            "",
            self._money_text(total_bruto),
            self._money_text(total_pengurang),
            self._money_text(total_bruto - total_pengurang),
        ])
        table(
            ["No", "Jenis", "NPWP Pemberi Kerja", "No Bupot", "Bruto", "Pengurang", "Netto"],
            bupot_rows,
            [0.45, 0.65, 1.55, 1.35, 1.0, 1.0, 1.0],
            row_height=20,
        )
        y += 8

        section("HARTA / SIMULASI I")
        harta_rows = []
        total_prev = total_now = 0.0
        for item in data.harta_current_rows:
            prev = self._money(item.nilai_tahun_sebelumnya)
            now = self._money(item.nilai_tahun_berjalan)
            total_prev += prev
            total_now += now
            harta_rows.append([
                item.nomor,
                item.kode_eform,
                item.kode_ct,
                item.nama_harta,
                item.tahun_perolehan,
                self._money_text(prev),
                self._money_text(now),
            ])
        harta_rows.append(["", "", "", "TOTAL HARTA", "", self._money_text(total_prev), self._money_text(total_now)])
        table(
            ["No", "EFORM", "CT", "Nama Harta", "Tahun", str(data.tahun_pajak - 1), str(data.tahun_pajak)],
            harta_rows,
            [0.4, 0.65, 0.65, 2.0, 0.7, 1.0, 1.0],
            row_height=20,
        )
        y += 8

        section("RINGKASAN PPh")
        table(
            ["Komponen", "Nilai"],
            [
                ["Status PTKP", data.status_ptkp],
                ["PTKP", f"Rp {self._money_text(data.pph_components.get('ptkp', 0))}"],
                ["PPh Terutang", f"Rp {self._money_text(data.pph_calc_result.get('pph_terutang', data.pph_components.get('pph_terutang', 0)))}"],
                ["Kredit Pajak", f"Rp {self._money_text(data.pph_components.get('kredit_pajak', 0))}"],
                ["PPh Pasal 25", f"Rp {self._money_text(data.pph_components.get('pph25', 0))}"],
            ],
            [2.2, 1.8],
            row_height=20,
        )

        if data.analisis_result is not None:
            y += 8
            section("ANALISIS PENGHASILAN VS HARTA")
            analysis = data.analisis_result
            table(
                ["Komponen", "Nilai"],
                [
                    ["Total Pengeluaran", f"Rp {self._money_text(getattr(analysis, 'total_pengeluaran', 0))}"],
                    ["Penghasilan Netto", f"Rp {self._money_text(getattr(analysis, 'penghasilan_netto', 0))}"],
                    ["Selisih", f"Rp {self._money_text(getattr(analysis, 'selisih_pengeluaran_vs_penghasilan', 0))}"],
                ],
                [2.2, 1.8],
                row_height=20,
            )

        y += 10
        ensure(18)
        painter.setFont(QFont("Arial", 7))
        painter.drawText(
            QRectF(margin, y, usable_width, 18),
            Qt.AlignCenter | Qt.AlignVCenter,
            "Arsip Kertas Kerja — TAX_CONVERTER L-1 • Bukan Form 1770",
        )

    @staticmethod
    def _looks_numeric(value: str) -> bool:
        text = str(value).replace("Rp", "").replace(".", "").replace(",", "").replace("-", "").strip()
        return bool(text) and text.isdigit()

    @staticmethod
    def _clip_text(value: str, width: float) -> str:
        text = str(value)
        max_chars = max(4, int(width / 5.2))
        if len(text) <= max_chars:
            return text
        return text[: max_chars - 1] + "…"
