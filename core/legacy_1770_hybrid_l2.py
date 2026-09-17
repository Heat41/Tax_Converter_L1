from __future__ import annotations

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


class LegacyLampiranIIXlsxRenderer:
    """Renderer Form 1770-II format lama dari canonical Bupot."""

    THIN = Side(style="thin", color="000000")
    BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
    VALUE_FILL = PatternFill("solid", fgColor="FFF2CC")
    MONEY = '#,##0;[Red]-#,##0;-'

    @staticmethod
    def _money(value) -> int:
        try:
            return int(round(float(value or 0)))
        except (TypeError, ValueError):
            return 0

    @staticmethod
    def _jenis_pph(item) -> str:
        raw = str(item.jenis_pph or "").strip()
        if not raw:
            return ""
        if raw.upper().startswith("PPH"):
            return raw
        return f"PPh {raw}"

    def render(self, ws, data) -> None:
        self._header(ws, data)
        row = self._table(ws, data)
        self._finish(ws, row)

    def _header(self, ws, data) -> None:
        ws.sheet_view.showGridLines = False

        ws.merge_cells("A1:B3")
        ws["A1"] = "FORMULIR\n1770 - II\nKEMENTERIAN KEUANGAN RI\nDIREKTORAT JENDERAL PAJAK"
        ws["A1"].font = Font(size=8, bold=True)
        ws["A1"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        ws.merge_cells("C1:H1")
        ws["C1"] = "LAMPIRAN - II"
        ws["C1"].font = Font(size=9, bold=True)
        ws["C1"].alignment = Alignment(horizontal="center")

        ws.merge_cells("C2:H2")
        ws["C2"] = "SPT TAHUNAN PPh WAJIB PAJAK ORANG PRIBADI"
        ws["C2"].font = Font(size=11, bold=True)
        ws["C2"].alignment = Alignment(horizontal="center")

        ws.merge_cells("C3:H4")
        ws["C3"] = (
            "DAFTAR PEMOTONGAN/PEMUNGUTAN PPh OLEH PIHAK LAIN,\n"
            "PPh YANG DIBAYAR/DIPOTONG DI LUAR NEGERI DAN\n"
            "PPh DITANGGUNG PEMERINTAH"
        )
        ws["C3"].font = Font(size=8, bold=True)
        ws["C3"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        ws.merge_cells("I1:J1")
        ws["I1"] = "TAHUN PAJAK"
        ws["I1"].font = Font(size=8, bold=True)
        ws["I1"].alignment = Alignment(horizontal="center")
        ws.merge_cells("I2:J2")
        ws["I2"] = int(data.tahun_pajak or 0)
        ws["I2"].fill = self.VALUE_FILL
        ws["I2"].font = Font(size=10, bold=True)
        ws["I2"].alignment = Alignment(horizontal="center")

        ws.merge_cells("A6:J6")
        ws["A6"] = "PERHATIAN : SEBELUM MENGISI BACALAH PETUNJUK PENGISIAN"
        ws["A6"].font = Font(size=7)

        ws.merge_cells("A8:B8")
        ws["A8"] = "NPWP"
        ws.merge_cells("C8:J8")
        ws["C8"] = str(data.npwp or "")
        ws["C8"].fill = self.VALUE_FILL
        ws["C8"].font = Font(bold=True)

        ws.merge_cells("A9:B9")
        ws["A9"] = "NAMA WAJIB PAJAK"
        ws.merge_cells("C9:J9")
        ws["C9"] = str(data.nama_wp or "").upper()
        ws["C9"].fill = self.VALUE_FILL
        ws["C9"].font = Font(bold=True)

        for row in (8, 9):
            for col in range(1, 11):
                ws.cell(row, col).border = self.BORDER

        ws.merge_cells("A11:J11")
        ws["A11"] = (
            "BAGIAN A : DAFTAR PEMOTONGAN/PEMUNGUTAN PPh OLEH PIHAK LAIN, "
            "PPh YANG DIBAYAR/DIPOTONG DI LUAR NEGERI DAN PPh DITANGGUNG PEMERINTAH"
        )
        ws["A11"].font = Font(size=8, bold=True)
        ws["A11"].alignment = Alignment(wrap_text=True)

    def _table(self, ws, data) -> int:
        row = 13
        headers = (
            "NO",
            "NAMA\nPEMOTONG/PEMUNGUT\nPAJAK",
            "NPWP\nPEMOTONG/PEMUNGUT\nPAJAK",
            "BUKTI\nPEMOTONGAN/PEMUNGUTAN\nNOMOR",
            "BUKTI\nPEMOTONGAN/PEMUNGUTAN\nTANGGAL",
            "JENIS PAJAK : PPh PASAL\n21/22/23/24/26/DTP *)",
            "JUMLAH PPh YANG DIPOTONG /\nDIPUNGUT\n(Rupiah)",
        )
        spans = ((1, 1), (2, 3), (4, 5), (6, 6), (7, 7), (8, 8), (9, 10))
        for value, (c1, c2) in zip(headers, spans):
            if c1 != c2:
                ws.merge_cells(start_row=row, start_column=c1, end_row=row, end_column=c2)
            cell = ws.cell(row, c1)
            cell.value = value
            cell.font = Font(size=7, bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            for col in range(c1, c2 + 1):
                ws.cell(row, col).border = self.BORDER
        ws.row_dimensions[row].height = 48
        row += 1

        # Baris legacy menyediakan 15 slot sebelum 'dst'. Jika data lebih banyak,
        # renderer memperpanjang tabel tanpa membuang canonical Bupot.
        items = list(data.bupot_rows or [])
        visible_rows = max(15, len(items))
        total_pph = 0

        for index in range(visible_rows):
            item = items[index] if index < len(items) else None
            if item is None:
                values = (index + 1 if index < 14 else "15\ndst", "", "", "", "", "", 0)
            else:
                amount = self._money(item.pph_dipotong)
                total_pph += amount
                values = (
                    index + 1 if index < 14 else ("15\ndst" if index == 14 else index + 1),
                    str(item.nama_pemotong or ""),
                    str(item.npwp_pemotong or item.npwp_pemberi_kerja or ""),
                    str(item.no_bupot or ""),
                    str(item.tanggal_pemotongan or item.tanggal_bukti or ""),
                    self._jenis_pph(item),
                    amount,
                )

            for value, (c1, c2) in zip(values, spans):
                if c1 != c2:
                    ws.merge_cells(start_row=row, start_column=c1, end_row=row, end_column=c2)
                cell = ws.cell(row, c1)
                cell.value = value
                cell.alignment = Alignment(
                    horizontal="right" if c1 == 9 else ("center" if c1 in {1, 4, 6, 7, 8} else "left"),
                    vertical="center",
                    wrap_text=True,
                )
                if c1 == 9:
                    cell.number_format = self.MONEY
                    if item is not None:
                        cell.fill = self.VALUE_FILL
                for col in range(c1, c2 + 1):
                    ws.cell(row, col).border = self.BORDER
            ws.row_dimensions[row].height = 26
            row += 1

        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
        ws.cell(row, 1).value = "JUMLAH BAGIAN A"
        ws.cell(row, 1).font = Font(bold=True)
        ws.cell(row, 1).alignment = Alignment(horizontal="center")
        ws.merge_cells(start_row=row, start_column=9, end_row=row, end_column=10)
        ws.cell(row, 9).value = total_pph
        ws.cell(row, 9).number_format = self.MONEY
        ws.cell(row, 9).font = Font(bold=True)
        ws.cell(row, 9).fill = self.VALUE_FILL
        ws.cell(row, 9).alignment = Alignment(horizontal="right")
        for col in range(1, 11):
            ws.cell(row, col).border = self.BORDER

        row += 1
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=10)
        ws.cell(row, 1).value = "Pindahkan Jumlah Bagian A Kolom (7) ke Formulir 1770 kredit pajak."
        ws.cell(row, 1).font = Font(size=7, italic=True)
        ws.cell(row, 1).alignment = Alignment(horizontal="right")
        return row

    def _finish(self, ws, row: int) -> None:
        widths = {
            "A": 4.0, "B": 9.0, "C": 9.0, "D": 9.0, "E": 9.0,
            "F": 12.0, "G": 11.0, "H": 14.0, "I": 11.0, "J": 11.0,
        }
        for col, width in widths.items():
            ws.column_dimensions[col].width = width
        ws.page_setup.orientation = "portrait"
        ws.page_setup.paperSize = ws.PAPERSIZE_LEGAL
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 1
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_margins.left = 0.2
        ws.page_margins.right = 0.2
        ws.page_margins.top = 0.25
        ws.page_margins.bottom = 0.25
        ws.print_area = f"A1:J{row}"
