from __future__ import annotations

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


class LegacyLampiranIVXlsxRenderer:
    """Renderer Form 1770-IV format lama dari snapshot FINAL.

    Bagian A diisi dari canonical Harta. Bagian B (Utang) dan Bagian C
    (Susunan Anggota Keluarga) tetap disediakan dalam format legacy, tetapi
    tidak diisi bila domain kerja belum memiliki sumber data yang sah.
    """

    THIN = Side(style="thin", color="000000")
    BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
    VALUE_FILL = PatternFill("solid", fgColor="FFF2CC")
    MONEY = '#,##0;[Red]-#,##0;-'

    @staticmethod
    def _money(value) -> int:
        try:
            return int(round(float(value or 0)))
        except (TypeError, ValueError):
            return 0

    @staticmethod
    def _text(value) -> str:
        if value is None:
            return ""
        return " ".join(str(value).strip().split())

    def render(self, ws, data) -> None:
        self._header(ws, data)
        row = self._bagian_a(ws, data)
        row = self._bagian_b(ws, row)
        row = self._bagian_c(ws, row)
        self._finish(ws, row)

    def _header(self, ws, data) -> None:
        ws.sheet_view.showGridLines = False

        ws.merge_cells("A1:B3")
        ws["A1"] = (
            "FORMULIR\n1770 - IV\n"
            "KEMENTERIAN KEUANGAN RI\nDIREKTORAT JENDERAL PAJAK"
        )
        ws["A1"].font = Font(size=8, bold=True)
        ws["A1"].alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )

        ws.merge_cells("C1:H1")
        ws["C1"] = "LAMPIRAN - IV"
        ws["C1"].font = Font(size=9, bold=True)
        ws["C1"].alignment = Alignment(horizontal="center")

        ws.merge_cells("C2:H2")
        ws["C2"] = "SPT TAHUNAN PPh WAJIB PAJAK ORANG PRIBADI"
        ws["C2"].font = Font(size=11, bold=True)
        ws["C2"].alignment = Alignment(horizontal="center")

        ws.merge_cells("C3:H5")
        ws["C3"] = (
            "HARTA PADA AKHIR TAHUN\n"
            "KEWAJIBAN/UTANG PADA AKHIR TAHUN\n"
            "DAFTAR SUSUNAN ANGGOTA KELUARGA"
        )
        ws["C3"].font = Font(size=8, bold=True)
        ws["C3"].alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )

        ws.merge_cells("I1:J1")
        ws["I1"] = "TAHUN PAJAK"
        ws["I1"].font = Font(size=8, bold=True)
        ws["I1"].alignment = Alignment(horizontal="center")
        ws.merge_cells("I2:J2")
        ws["I2"] = int(data.tahun_pajak or 0)
        ws["I2"].fill = self.VALUE_FILL
        ws["I2"].font = Font(size=10, bold=True)
        ws["I2"].alignment = Alignment(horizontal="center")

        ws.merge_cells("A7:B7")
        ws["A7"] = "NPWP"
        ws.merge_cells("C7:J7")
        ws["C7"] = str(data.npwp or "")
        ws["C7"].fill = self.VALUE_FILL
        ws["C7"].font = Font(bold=True)

        ws.merge_cells("A8:B8")
        ws["A8"] = "NAMA WAJIB PAJAK"
        ws.merge_cells("C8:J8")
        ws["C8"] = str(data.nama_wp or "").upper()
        ws["C8"].fill = self.VALUE_FILL
        ws["C8"].font = Font(bold=True)

        for row in (7, 8):
            for col in range(1, 11):
                ws.cell(row, col).border = self.BORDER

    def _section(self, ws, row: int, title: str) -> int:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=10)
        cell = ws.cell(row, 1)
        cell.value = title
        cell.font = Font(size=8, bold=True)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        for col in range(1, 11):
            ws.cell(row, col).border = self.BORDER
        ws.row_dimensions[row].height = 24
        return row + 1

    def _bagian_a(self, ws, data) -> int:
        row = 10
        row = self._section(ws, row, "BAGIAN A : HARTA PADA AKHIR TAHUN")

        headers = (
            "NO.",
            "KODE HARTA",
            "NAMA HARTA",
            "TAHUN PEROLEHAN",
            "HARGA PEROLEHAN (Rupiah)",
            "KETERANGAN",
        )
        spans = ((1, 1), (2, 2), (3, 4), (5, 5), (6, 7), (8, 10))
        for value, (c1, c2) in zip(headers, spans):
            if c1 != c2:
                ws.merge_cells(start_row=row, start_column=c1, end_row=row, end_column=c2)
            cell = ws.cell(row, c1)
            cell.value = value
            cell.font = Font(size=7, bold=True)
            cell.alignment = Alignment(
                horizontal="center", vertical="center", wrap_text=True
            )
            for col in range(c1, c2 + 1):
                ws.cell(row, col).border = self.BORDER
        ws.row_dimensions[row].height = 36
        row += 1

        items = list(data.harta_current_rows or [])
        visible_rows = max(10, len(items))
        total = 0

        for index in range(visible_rows):
            item = items[index] if index < len(items) else None
            if item is None:
                values = (index + 1, "", "", "", 0, "")
            else:
                amount = self._money(item.nilai_tahun_berjalan)
                total += amount
                note_parts = [
                    self._text(getattr(item, "nomor_akun_keterangan", "")),
                    self._text(getattr(item, "atas_nama", "")),
                    self._text(getattr(item, "nama_bank", "")),
                ]
                note = " | ".join(value for value in note_parts if value and value != "-")
                values = (
                    index + 1,
                    self._text(getattr(item, "kode_eform", "")),
                    self._text(getattr(item, "nama_harta", "")),
                    int(getattr(item, "tahun_perolehan", 0) or 0) or "",
                    amount,
                    note,
                )

            for value, (c1, c2) in zip(values, spans):
                if c1 != c2:
                    ws.merge_cells(start_row=row, start_column=c1, end_row=row, end_column=c2)
                cell = ws.cell(row, c1)
                cell.value = value
                cell.alignment = Alignment(
                    horizontal="right" if c1 == 6 else ("center" if c1 in {1, 2, 5} else "left"),
                    vertical="center",
                    wrap_text=True,
                )
                if c1 == 6:
                    cell.number_format = self.MONEY
                    if item is not None:
                        cell.fill = self.VALUE_FILL
                elif item is not None and c1 in {2, 3, 5, 8}:
                    cell.fill = self.VALUE_FILL
                for col in range(c1, c2 + 1):
                    ws.cell(row, col).border = self.BORDER
            ws.row_dimensions[row].height = 24
            row += 1

        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
        ws.cell(row, 1).value = "JUMLAH BAGIAN A"
        ws.cell(row, 1).font = Font(bold=True)
        ws.cell(row, 1).alignment = Alignment(horizontal="center")
        ws.merge_cells(start_row=row, start_column=6, end_row=row, end_column=7)
        ws.cell(row, 6).value = total
        ws.cell(row, 6).number_format = self.MONEY
        ws.cell(row, 6).font = Font(bold=True)
        ws.cell(row, 6).fill = self.VALUE_FILL
        ws.cell(row, 6).alignment = Alignment(horizontal="right")
        ws.merge_cells(start_row=row, start_column=8, end_row=row, end_column=10)
        for col in range(1, 11):
            ws.cell(row, col).border = self.BORDER
        return row + 2

    def _bagian_b(self, ws, row: int) -> int:
        row = self._section(ws, row, "BAGIAN B : KEWAJIBAN/UTANG PADA AKHIR TAHUN")
        headers = (
            "NO.",
            "NAMA PEMBERI PINJAMAN",
            "ALAMAT PEMBERI PINJAMAN",
            "TAHUN PEMINJAMAN",
            "JUMLAH (Rupiah)",
        )
        spans = ((1, 1), (2, 4), (5, 7), (8, 8), (9, 10))
        for value, (c1, c2) in zip(headers, spans):
            if c1 != c2:
                ws.merge_cells(start_row=row, start_column=c1, end_row=row, end_column=c2)
            cell = ws.cell(row, c1)
            cell.value = value
            cell.font = Font(size=7, bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            for col in range(c1, c2 + 1):
                ws.cell(row, col).border = self.BORDER
        ws.row_dimensions[row].height = 32
        row += 1

        # Domain aktif belum mempunyai canonical daftar Utang. Pertahankan slot
        # legacy agar user tidak kehilangan bentuk form dan jangan menebak data.
        for index in range(8):
            values = (index + 1, "", "", "", 0)
            for value, (c1, c2) in zip(values, spans):
                if c1 != c2:
                    ws.merge_cells(start_row=row, start_column=c1, end_row=row, end_column=c2)
                cell = ws.cell(row, c1)
                cell.value = value
                cell.alignment = Alignment(horizontal="center" if c1 in {1, 8} else "left", vertical="center", wrap_text=True)
                if c1 == 9:
                    cell.number_format = self.MONEY
                for col in range(c1, c2 + 1):
                    ws.cell(row, col).border = self.BORDER
            ws.row_dimensions[row].height = 22
            row += 1

        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
        ws.cell(row, 1).value = "JUMLAH BAGIAN B"
        ws.cell(row, 1).font = Font(bold=True)
        ws.cell(row, 1).alignment = Alignment(horizontal="center")
        ws.merge_cells(start_row=row, start_column=9, end_row=row, end_column=10)
        ws.cell(row, 9).value = 0
        ws.cell(row, 9).number_format = self.MONEY
        ws.cell(row, 9).font = Font(bold=True)
        for col in range(1, 11):
            ws.cell(row, col).border = self.BORDER
        return row + 2

    def _bagian_c(self, ws, row: int) -> int:
        row = self._section(ws, row, "BAGIAN C : DAFTAR SUSUNAN ANGGOTA KELUARGA")
        headers = ("NO.", "NAMA", "NIK", "HUBUNGAN KELUARGA", "PEKERJAAN")
        spans = ((1, 1), (2, 4), (5, 6), (7, 8), (9, 10))
        for value, (c1, c2) in zip(headers, spans):
            if c1 != c2:
                ws.merge_cells(start_row=row, start_column=c1, end_row=row, end_column=c2)
            cell = ws.cell(row, c1)
            cell.value = value
            cell.font = Font(size=7, bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            for col in range(c1, c2 + 1):
                ws.cell(row, col).border = self.BORDER
        ws.row_dimensions[row].height = 30
        row += 1

        # Domain aktif belum mempunyai canonical susunan keluarga.
        for index in range(10):
            values = (index + 1, "", "", "", "")
            for value, (c1, c2) in zip(values, spans):
                if c1 != c2:
                    ws.merge_cells(start_row=row, start_column=c1, end_row=row, end_column=c2)
                cell = ws.cell(row, c1)
                cell.value = value
                cell.alignment = Alignment(horizontal="center" if c1 == 1 else "left", vertical="center", wrap_text=True)
                for col in range(c1, c2 + 1):
                    ws.cell(row, col).border = self.BORDER
            ws.row_dimensions[row].height = 22
            row += 1
        return row

    def _finish(self, ws, row: int) -> None:
        widths = {
            "A": 4.0,
            "B": 8.0,
            "C": 10.0,
            "D": 10.0,
            "E": 10.0,
            "F": 11.0,
            "G": 11.0,
            "H": 11.0,
            "I": 11.0,
            "J": 11.0,
        }
        for col, width in widths.items():
            ws.column_dimensions[col].width = width

        ws.page_setup.orientation = "portrait"
        ws.page_setup.paperSize = ws.PAPERSIZE_LEGAL
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 1
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_margins.left = 0.2
        ws.page_margins.right = 0.2
        ws.page_margins.top = 0.25
        ws.page_margins.bottom = 0.25
        ws.print_area = f"A1:J{row}"
