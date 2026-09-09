from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from PySide6.QtCore import QMarginsF, QRectF
from PySide6.QtGui import QFont, QPainter, QPageLayout, QPageSize, QPdfWriter

from core.finalization import FinalizationInput


@dataclass(frozen=True)
class WorksheetArchiveExportResult:
    output_path: Path
    format_name: str
    bupot_count: int
    harta_count: int


class WorksheetArchiveExporter:
    """Stage 8B.2: ekspor Kertas Kerja menjadi Excel atau PDF arsip.

    Excel sengaja memakai label/header yang kompatibel dengan importer Stage 8B.1,
    sehingga file hasil export dapat diimpor kembali ke aplikasi. PDF adalah arsip
    baca-saja Kertas Kerja dan BUKAN PDF Form 1770 (Stage 8C).
    """

    @staticmethod
    def _ensure_exportable(data: FinalizationInput) -> None:
        if not data.npwp or not data.tahun_pajak:
            raise ValueError("Worksheet belum memiliki identitas NPWP dan Tahun Pajak.")
        dirty = []
        if data.is_harta_dirty:
            dirty.append("Harta")
        if data.is_pph_dirty:
            dirty.append("Penghasilan/PPh")
        if data.is_analisis_dirty:
            dirty.append("Analisis")
        if dirty:
            raise ValueError(
                "Simpan perubahan sebelum export Kertas Kerja: " + ", ".join(dirty)
            )

    def export_excel(self, data: FinalizationInput, output_path: str | Path) -> WorksheetArchiveExportResult:
        self._ensure_exportable(data)
        path = Path(output_path)
        if path.suffix.lower() != ".xlsx":
            path = path.with_suffix(".xlsx")
        path.parent.mkdir(parents=True, exist_ok=True)

        wb = Workbook()
        annual = wb.active
        annual.title = str(data.tahun_pajak)
        simulasi = wb.create_sheet("SIMULASI I")
        ref = wb.create_sheet("REF")

        self._write_annual_sheet(annual, data)
        self._write_simulasi_sheet(simulasi, data)
        self._write_ref_sheet(ref)
        wb.save(path)

        return WorksheetArchiveExportResult(path, "Excel", len(data.bupot_rows), len(data.harta_current_rows))

    def export_pdf(self, data: FinalizationInput, output_path: str | Path) -> WorksheetArchiveExportResult:
        self._ensure_exportable(data)
        path = Path(output_path)
        if path.suffix.lower() != ".pdf":
            path = path.with_suffix(".pdf")
        path.parent.mkdir(parents=True, exist_ok=True)

        writer = QPdfWriter(str(path))
        writer.setPageSize(QPageSize(QPageSize.A4))
        writer.setResolution(96)
        writer.setPageMargins(QMarginsF(10, 10, 10, 10), QPageLayout.Millimeter)
        writer.setTitle(f"Kertas Kerja SPT Tahunan {data.tahun_pajak} - {data.nama_wp}")
        writer.setCreator("TAX_CONVERTER L-1")

        painter = QPainter(writer)
        if not painter.isActive():
            raise RuntimeError("PDF Kertas Kerja tidak dapat dibuat.")
        try:
            self._paint_pdf(painter, writer, data)
        finally:
            painter.end()

        return WorksheetArchiveExportResult(path, "PDF", len(data.bupot_rows), len(data.harta_current_rows))

    @staticmethod
    def _money(value: Any) -> float:
        try:
            return float(value or 0)
        except (TypeError, ValueError):
            return 0.0

    @classmethod
    def _money_text(cls, value: Any) -> str:
        return f"{cls._money(value):,.0f}".replace(",", ".")

    @staticmethod
    def _style_title(cell) -> None:
        cell.font = Font(bold=True, size=14)

    @staticmethod
    def _style_header_row(ws, row: int, start: int, end: int) -> None:
        fill = PatternFill("solid", fgColor="D9EAF7")
        for col in range(start, end + 1):
            cell = ws.cell(row, col)
            cell.font = Font(bold=True)
            cell.fill = fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    def _write_annual_sheet(self, ws, data: FinalizationInput) -> None:
        year = data.tahun_pajak
        ws["A1"] = "KERTAS KERJA SPT TAHUNAN"
        self._style_title(ws["A1"])
        ws["A2"] = f"TAHUN {year}"
        ws["A4"], ws["B4"] = "NAMA", data.nama_wp
        ws["A5"], ws["B5"] = "NPWP", data.npwp

        headers = ["NO", "JENIS", "NPWP PEMBERI KERJA", "NO BUPOT", "BRUTO", "PENGURANG", "NETTO"]
        for col, text in enumerate(headers, start=1):
            ws.cell(8, col, text)
        self._style_header_row(ws, 8, 1, len(headers))

        row = 9
        total_bruto = total_pengurang = 0.0
        for no, item in enumerate(data.bupot_rows, start=1):
            bruto = self._money(item.bruto)
            pengurang = self._money(item.pengurang)
            values = [no, item.jenis, item.npwp_pemberi_kerja, item.no_bupot, bruto, pengurang, bruto - pengurang]
            for col, value in enumerate(values, start=1):
                ws.cell(row, col, value)
            total_bruto += bruto
            total_pengurang += pengurang
            row += 1
        ws.cell(row, 1, "TOTAL")
        ws.cell(row, 5, total_bruto)
        ws.cell(row, 6, total_pengurang)
        ws.cell(row, 7, total_bruto - total_pengurang)
        for col in range(1, 8):
            ws.cell(row, col).font = Font(bold=True)

        row += 3
        ws.cell(row, 1, "PEREDARAN BRUTO UMKM")
        ws.cell(row, 1).font = Font(bold=True)
        row += 1
        for col, text in enumerate(["NO", "MASA", "KETERANGAN", "BRUTO", "TARIF", "PPh SETOR"], start=1):
            ws.cell(row, col, text)
        self._style_header_row(ws, row, 1, 6)
        bruto_bulanan = list(data.umkm_state.get("bruto_bulanan", []) or [])
        setor_bulanan = list(data.umkm_state.get("pph_setor_bulanan", []) or [])
        for month in range(12):
            row += 1
            ws.cell(row, 1, month + 1)
            ws.cell(row, 2, f"M{month + 1}")
            ws.cell(row, 4, self._money(bruto_bulanan[month] if month < len(bruto_bulanan) else 0))
            ws.cell(row, 6, self._money(setor_bulanan[month] if month < len(setor_bulanan) else 0))

        other = dict(data.penghasilan_lainnya or {})
        row += 3
        ws.cell(row, 1, "PENGHASILAN LAINNYA")
        ws.cell(row, 1).font = Font(bold=True)
        other_rows = [
            ("Penghasilan Dalam Negeri Lainnya", "ADA" if other.get("domestic_other_enabled") else "TIDAK", other.get("domestic_other_dpp", 0), 0, ""),
            ("Sewa atas Tanah dan/atau Bangunan", "", other.get("sewa_dpp", 0), other.get("sewa_pph", 0), ""),
            ("Honor", "", other.get("honor_dpp", 0), other.get("honor_pph", 0), ""),
            ("Pekerjaan bebas", "", other.get("pekerjaan_bebas_dpp", 0), 0, ""),
            ("Prive", "", other.get("prive_dpp", 0), 0, ""),
            ("Hibah / Warisan", "", other.get("hibah_warisan_dpp", 0), 0, other.get("hibah_warisan_note", "")),
        ]
        for label, status, dpp, pph, note in other_rows:
            row += 1
            ws.cell(row, 2, label)
            ws.cell(row, 4, status)
            ws.cell(row, 5, self._money(dpp))
            ws.cell(row, 6, self._money(pph))
            ws.cell(row, 7, note)

        final_rows = list(other.get("final_other_rows", []) or [])
        row += 1
        ws.cell(row, 2, "Penghasilan Final Lainnya")
        for item in final_rows:
            row += 1
            if isinstance(item, dict):
                dpp = self._money(item.get("dpp", 0))
                tarif = self._money(item.get("tarif", 0))
                pph = self._money(item.get("pph", dpp * tarif))
                ws.cell(row, 4, item.get("keterangan", "Final Lainnya"))
                ws.cell(row, 5, dpp)
                ws.cell(row, 6, pph)

        row += 2
        ws.cell(row, 1, "PENGURANG PENGHASILAN NETO")
        ws.cell(row, 4, "Zakat")
        ws.cell(row, 5, self._money(data.zakat))
        row += 2
        ws.cell(row, 2, "STATUS PTKP")
        ws.cell(row, 5, data.status_ptkp)
        ws.cell(row + 1, 2, "PTKP")
        ws.cell(row + 1, 5, self._money(data.pph_components.get("ptkp", 0)))
        ws.cell(row + 2, 2, "PPh21 Terutang")
        ws.cell(row + 2, 5, self._money(data.pph_calc_result.get("pph_terutang", data.pph_components.get("pph_terutang", 0))))
        ws.cell(row + 3, 2, "PPh21 sudah dipotong/ dipungut")
        ws.cell(row + 3, 5, self._money(data.pph_components.get("kredit_pajak", 0)))
        ws.cell(row + 4, 2, "Angsuran PPh Pasal 25")
        ws.cell(row + 4, 5, self._money(data.pph_components.get("pph25", 0)))

        widths = {1: 9, 2: 34, 3: 22, 4: 24, 5: 18, 6: 18, 7: 24}
        for col, width in widths.items():
            ws.column_dimensions[get_column_letter(col)].width = width
        ws.freeze_panes = "A8"

    def _write_simulasi_sheet(self, ws, data: FinalizationInput) -> None:
        year = data.tahun_pajak
        prev = year - 1
        ws["A1"] = "SIMULASI I — HARTA DAN ANALISIS"
        self._style_title(ws["A1"])
        ws["A3"], ws["C3"] = "NAMA :", data.nama_wp
        ws["A4"], ws["C4"] = "NPWP :", data.npwp

        headers = ["NO", "KODE EFORM", "KODE CT", "NAMA HARTA", "NOMOR AKUN / KETERANGAN", "ATAS NAMA", "NAMA BANK", "TH PEROLEHAN", str(prev), str(year)]
        for col, text in enumerate(headers, start=1):
            ws.cell(8, col, text)
        self._style_header_row(ws, 8, 1, 10)

        row = 9
        total_prev = total_now = 0.0
        for item in data.harta_current_rows:
            values = [
                item.nomor, item.kode_eform, item.kode_ct, item.nama_harta,
                item.nomor_akun_keterangan, item.atas_nama, item.nama_bank,
                item.tahun_perolehan, self._money(item.nilai_tahun_sebelumnya), self._money(item.nilai_tahun_berjalan),
            ]
            for col, value in enumerate(values, start=1):
                ws.cell(row, col, value)
            total_prev += self._money(item.nilai_tahun_sebelumnya)
            total_now += self._money(item.nilai_tahun_berjalan)
            row += 1
        ws.cell(row, 1, "TOTAL HARTA")
        ws.cell(row, 9, total_prev)
        ws.cell(row, 10, total_now)
        for col in range(1, 11):
            ws.cell(row, col).font = Font(bold=True)

        analysis = data.analisis_result
        row += 3
        ws.cell(row, 1, "UTANG")
        ws.cell(row, 1).font = Font(bold=True)
        row += 1
        ws.cell(row, 1, "TOTAL")
        if analysis is not None:
            ws.cell(row, 9, self._money(getattr(analysis, "total_utang_sebelumnya", 0)))
            ws.cell(row, 10, self._money(getattr(analysis, "total_utang_berjalan", 0)))

        if analysis is not None:
            row += 3
            ws.cell(row, 1, "PERHITUNGAN PENGHASILAN VS KENAIKAN HARTA")
            ws.cell(row, 1).font = Font(bold=True)
            items = [
                ("a", "Naik/Turun Harta dan Utang", getattr(analysis, "naik_turun_harta_utang", 0)),
                ("b", "Biaya Hidup Setahun", getattr(analysis, "biaya_hidup_setahun", 0)),
                ("c", "Pajak-pajak", getattr(analysis, "pajak_pajak", 0)),
                ("d", "Pengeluaran lain-lain", getattr(analysis, "pengeluaran_lain_lain", 0)),
                ("e", "Kerugian (keuntungan) penjualan aset", getattr(analysis, "kerugian_keuntungan_penjualan_aset", 0)),
                ("f", "Utang baru atas kredit", getattr(analysis, "utang_baru_atas_kredit", 0)),
                ("g", "Harta baru dari kredit", getattr(analysis, "harta_baru_dari_kredit", 0)),
                ("", "TOTAL PENGELUARAN PER TAHUN", getattr(analysis, "total_pengeluaran", 0)),
                ("", "Penghasilan Netto", getattr(analysis, "penghasilan_netto", 0)),
                ("", "Selisih Total Pengeluaran vs Penghasilan Netto", getattr(analysis, "selisih_pengeluaran_vs_penghasilan", 0)),
            ]
            for code, label, value in items:
                row += 1
                ws.cell(row, 1, code)
                ws.cell(row, 2, label)
                ws.cell(row, 10, self._money(value))

        widths = {1: 9, 2: 14, 3: 12, 4: 30, 5: 30, 6: 20, 7: 20, 8: 14, 9: 18, 10: 18}
        for col, width in widths.items():
            ws.column_dimensions[get_column_letter(col)].width = width
        ws.freeze_panes = "A8"

    @staticmethod
    def _write_ref_sheet(ws) -> None:
        ws["A1"] = "REF"
        ws["A1"].font = Font(bold=True, size=14)
        ws["A3"] = "Catatan"
        ws["B3"] = "File arsip dibuat oleh TAX_CONVERTER L-1 dan dapat diimpor kembali melalui Stage 8B.1."
        ws.column_dimensions["A"].width = 18
        ws.column_dimensions["B"].width = 90

    def _paint_pdf(self, painter: QPainter, writer: QPdfWriter, data: FinalizationInput) -> None:
        page = writer.pageLayout().paintRectPixels(writer.resolution())
        margin = 36.0
        width = float(page.width()) - margin * 2
        y = margin

        def new_page():
            nonlocal y
            writer.newPage()
            y = margin

        def text(value: str, size=9, bold=False, height=18, indent=0):
            nonlocal y
            if y + height > page.height() - margin:
                new_page()
            font = QFont("Arial", size)
            font.setBold(bold)
            painter.setFont(font)
            painter.drawText(QRectF(margin + indent, y, width - indent, height), str(value))
            y += height

        text("KERTAS KERJA SPT TAHUNAN", 15, True, 24)
        text(f"Tahun Pajak {data.tahun_pajak}", 11, True, 20)
        text(f"Nama WP : {data.nama_wp}")
        text(f"NPWP    : {data.npwp}")
        y += 8

        text("BUKTI POTONG / PENGHASILAN PEKERJAAN", 11, True, 22)
        text("No | Jenis | NPWP Pemberi Kerja | No Bupot | Bruto | Pengurang | Netto", 8, True, 18)
        total_bruto = total_pengurang = 0.0
        for no, item in enumerate(data.bupot_rows, start=1):
            bruto = self._money(item.bruto)
            pengurang = self._money(item.pengurang)
            line = f"{no} | {item.jenis} | {item.npwp_pemberi_kerja} | {item.no_bupot} | {self._money_text(bruto)} | {self._money_text(pengurang)} | {self._money_text(bruto-pengurang)}"
            text(line, 7, False, 16)
            total_bruto += bruto
            total_pengurang += pengurang
        text(f"TOTAL BRUTO {self._money_text(total_bruto)} | PENGURANG {self._money_text(total_pengurang)} | NETTO {self._money_text(total_bruto-total_pengurang)}", 8, True, 20)

        y += 6
        text("HARTA / SIMULASI I", 11, True, 22)
        text(f"No | Kode EFORM | Kode CT | Nama Harta | Th Perolehan | {data.tahun_pajak-1} | {data.tahun_pajak}", 8, True, 18)
        total_prev = total_now = 0.0
        for item in data.harta_current_rows:
            prev = self._money(item.nilai_tahun_sebelumnya)
            now = self._money(item.nilai_tahun_berjalan)
            line = f"{item.nomor} | {item.kode_eform} | {item.kode_ct} | {item.nama_harta} | {item.tahun_perolehan} | {self._money_text(prev)} | {self._money_text(now)}"
            text(line, 7, False, 16)
            total_prev += prev
            total_now += now
        text(f"TOTAL HARTA {data.tahun_pajak-1}: {self._money_text(total_prev)} | {data.tahun_pajak}: {self._money_text(total_now)}", 8, True, 20)

        y += 6
        text("RINGKASAN PPh", 11, True, 22)
        text(f"Status PTKP : {data.status_ptkp}")
        text(f"PTKP        : Rp {self._money_text(data.pph_components.get('ptkp', 0))}")
        text(f"PPh Terutang: Rp {self._money_text(data.pph_calc_result.get('pph_terutang', data.pph_components.get('pph_terutang', 0)))}")
        text(f"Kredit Pajak: Rp {self._money_text(data.pph_components.get('kredit_pajak', 0))}")
        text(f"PPh Pasal 25: Rp {self._money_text(data.pph_components.get('pph25', 0))}")

        if data.analisis_result is not None:
            y += 6
            text("ANALISIS PENGHASILAN VS HARTA", 11, True, 22)
            analysis = data.analisis_result
            text(f"Total Pengeluaran : Rp {self._money_text(getattr(analysis, 'total_pengeluaran', 0))}")
            text(f"Penghasilan Netto : Rp {self._money_text(getattr(analysis, 'penghasilan_netto', 0))}")
            text(f"Selisih            : Rp {self._money_text(getattr(analysis, 'selisih_pengeluaran_vs_penghasilan', 0))}")

        y += 12
        text("Arsip dibuat oleh TAX_CONVERTER L-1. PDF ini adalah arsip Kertas Kerja, bukan Form 1770.", 7, False, 16)
