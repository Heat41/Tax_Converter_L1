from __future__ import annotations

import hashlib
import json
from dataclasses import asdict, dataclass, field
from datetime import datetime, timezone
from pathlib import Path
from typing import List, Optional

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from core.finalization import FinalizationInput
from core.worksheet_pph_state import WorksheetBupotRow
from core.mapping.worksheet_harta_mapper import WorksheetHartaRow
from core.legacy_1770_hybrid_l1h2 import LegacyLampiranIH2XlsxRenderer


ROUNDTRIP_SCHEMA = "TAX_CONVERTER_L1_LEGACY_XLSX_V1"
META_SHEET = "_TC_META"
DATA_HARTA_SHEET = "DATA REVISI - HARTA"
DATA_BUPOT_SHEET = "DATA REVISI - BUPOT"


@dataclass(frozen=True)
class LegacyXlsxIssue:
    code: str
    severity: str
    message: str


@dataclass
class LegacyXlsxExportResult:
    output_path: Path
    revision: int = 0
    snapshot_hash: str = ""
    sha256: str = ""
    issues: List[LegacyXlsxIssue] = field(default_factory=list)

    @property
    def errors(self):
        return [item for item in self.issues if item.severity == "ERROR"]

    @property
    def ok(self) -> bool:
        return not self.errors and self.output_path.is_file()


@dataclass
class LegacyXlsxRoundTripResult:
    source_path: Path
    npwp: str = ""
    nama_wp: str = ""
    tahun_pajak: int = 0
    base_revision: int = 0
    base_snapshot_hash: str = ""
    harta_rows: List[WorksheetHartaRow] = field(default_factory=list)
    bupot_rows: List[WorksheetBupotRow] = field(default_factory=list)
    issues: List[LegacyXlsxIssue] = field(default_factory=list)

    @property
    def errors(self):
        return [item for item in self.issues if item.severity == "ERROR"]

    @property
    def ok(self) -> bool:
        return bool(self.npwp and self.tahun_pajak) and not self.errors


class Legacy1770HybridXlsxService:
    """XLSX Format Lama HYBRID yang aman untuk round-trip revisi.

    Visual workbook:
    - 01-02: area e-Form baru.
    - 03 dst: area legacy/format lama.
    - DATA REVISI: canonical editable tables untuk Harta dan Bupot.
    - _TC_META: metadata tersembunyi untuk traceability snapshot/revision.

    Import revisi TIDAK menebak posisi field dari sheet visual; hanya canonical
    DATA REVISI yang dibaca. Dengan demikian perubahan user/client tetap stabil
    walau layout form visual berkembang.
    """

    HARTA_HEADERS = (
        "NO", "KODE EFORM", "KODE CT", "NAMA HARTA",
        "NOMOR AKUN / KETERANGAN", "ATAS NAMA", "NAMA BANK",
        "TH PEROLEHAN", "TAHUN SEBELUMNYA", "TAHUN BERJALAN",
    )
    BUPOT_HEADERS = (
        "NO", "JENIS BUPOT", "NO BUKPOT", "MASA", "TAHUN", "SIFAT", "STATUS",
        "NPWP PENERIMA", "NAMA PENERIMA", "FASILITAS", "JENIS PPH", "KOP",
        "BRUTO", "DPP PERSEN", "TARIF", "PENGURANG BRUTO", "PPH", "BUKTI",
        "NO BUKTI", "TANGGAL BUKTI", "NPWP PEMOTONG", "NAMA PEMOTONG",
        "TANGGAL PEMOTONGAN", "MEKANISME SP2D", "NO SP2D",
    )

    FORM_SHEETS = (
        ("01 eForm Induk H1", "NEW_EFORM_H1", "SPT TAHUNAN PPh WAJIB PAJAK ORANG PRIBADI"),
        ("02 eForm Induk H2", "NEW_EFORM_H2", "SPT TAHUNAN PPh WAJIB PAJAK ORANG PRIBADI"),
        ("03 Legacy Lamp I H2", "LEGACY_L1_H2", "LAMPIRAN I - HALAMAN 2"),
        ("04 Legacy Lamp II", "LEGACY", "LAMPIRAN II"),
        ("05 Legacy Lamp III", "LEGACY", "LAMPIRAN III"),
        ("06 Legacy Lamp IV", "LEGACY", "LAMPIRAN IV"),
    )

    EFORM_DARK_FILL = PatternFill("solid", fgColor="1F4E78")
    EFORM_SECTION_FILL = PatternFill("solid", fgColor="D9EAF7")
    EFORM_LIGHT_FILL = PatternFill("solid", fgColor="F4F8FB")
    EFORM_VALUE_FILL = PatternFill("solid", fgColor="FFF2CC")
    EFORM_THIN = Side(style="thin", color="7F8C8D")
    EFORM_BORDER = Border(
        left=EFORM_THIN,
        right=EFORM_THIN,
        top=EFORM_THIN,
        bottom=EFORM_THIN,
    )

    @staticmethod
    def _text(value) -> str:
        if value is None:
            return ""
        return " ".join(str(value).strip().split())

    @staticmethod
    def _number(value) -> float:
        if value in (None, ""):
            return 0.0
        if isinstance(value, (int, float)):
            return float(value)
        text = str(value).strip().replace("Rp", "").replace(" ", "")
        if "," in text and "." in text:
            text = text.replace(".", "").replace(",", ".")
        elif text.count(".") > 1:
            text = text.replace(".", "")
        elif "," in text:
            text = text.replace(",", ".")
        return float(text or 0)

    @staticmethod
    def _digits(value) -> str:
        return "".join(ch for ch in str(value or "") if ch.isdigit())

    @staticmethod
    def _style_header(ws, row: int, max_col: int) -> None:
        fill = PatternFill("solid", fgColor="D9EAF7")
        for col in range(1, max_col + 1):
            cell = ws.cell(row, col)
            cell.font = Font(bold=True)
            cell.fill = fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    @staticmethod
    def _setup_print(ws, orientation="portrait") -> None:
        ws.sheet_view.showGridLines = False
        ws.page_setup.orientation = orientation
        ws.page_setup.paperSize = ws.PAPERSIZE_LEGAL
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.freeze_panes = "A2"

    def export(
        self,
        data: FinalizationInput,
        output_path: str | Path,
        *,
        revision: int,
        snapshot_hash: str,
    ) -> LegacyXlsxExportResult:
        target = Path(output_path)
        if target.suffix.lower() != ".xlsx":
            target = target.with_suffix(".xlsx")

        result = LegacyXlsxExportResult(
            output_path=target,
            revision=int(revision or 0),
            snapshot_hash=str(snapshot_hash or ""),
        )

        if not data.npwp or not data.tahun_pajak:
            result.issues.append(
                LegacyXlsxIssue("LX_001", "ERROR", "Identitas snapshot FINAL belum lengkap.")
            )
            return result

        wb = Workbook()
        wb.remove(wb.active)

        for title, mode, heading in self.FORM_SHEETS:
            ws = wb.create_sheet(title)
            self._render_form_sheet(ws, data, mode, heading, revision)

        self._write_harta_sheet(wb.create_sheet(DATA_HARTA_SHEET), data)
        self._write_bupot_sheet(wb.create_sheet(DATA_BUPOT_SHEET), data)
        self._write_meta_sheet(
            wb.create_sheet(META_SHEET),
            data,
            revision=revision,
            snapshot_hash=snapshot_hash,
        )

        wb[META_SHEET].sheet_state = "veryHidden"
        target.parent.mkdir(parents=True, exist_ok=True)
        wb.save(target)
        result.sha256 = hashlib.sha256(target.read_bytes()).hexdigest()
        return result

    @staticmethod
    def _safe_float(value) -> float:
        try:
            return float(value or 0)
        except (TypeError, ValueError):
            return 0.0

    @staticmethod
    def _money_excel(value) -> int:
        try:
            return int(round(float(value or 0)))
        except (TypeError, ValueError):
            return 0

    @staticmethod
    def _yes_no_code(condition: bool) -> int:
        # Acuan e-Form memakai 1 = Tidak, 2 = Ya.
        return 2 if condition else 1

    def _pph_value(self, data: FinalizationInput, key: str, default=0.0) -> float:
        calc = data.pph_calc_result or {}
        components = data.pph_components or {}
        if key in calc:
            return self._safe_float(calc.get(key))
        return self._safe_float(components.get(key, default))

    def _setup_eform_sheet(self, ws, *, page_no: int) -> None:
        ws.sheet_view.showGridLines = False
        widths = {
            "A": 4.5,
            "B": 7.0,
            "C": 34.0,
            "D": 17.0,
            "E": 17.0,
            "F": 17.0,
            "G": 17.0,
            "H": 17.0,
            "I": 17.0,
            "J": 17.0,
        }
        for col, width in widths.items():
            ws.column_dimensions[col].width = width

        ws.page_setup.orientation = "portrait"
        ws.page_setup.paperSize = ws.PAPERSIZE_LEGAL
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 1
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_margins.left = 0.25
        ws.page_margins.right = 0.25
        ws.page_margins.top = 0.35
        ws.page_margins.bottom = 0.35
        ws.print_area = f"A1:J{92 if page_no == 1 else 94}"

    def _eform_banner(self, ws, data: FinalizationInput, *, page_no: int, revision: int) -> int:
        self._setup_eform_sheet(ws, page_no=page_no)

        ws.merge_cells("A1:J1")
        ws["A1"] = "KEMENTERIAN KEUANGAN REPUBLIK INDONESIA • DIREKTORAT JENDERAL PAJAK"
        ws["A1"].font = Font(size=10, bold=True, color="FFFFFF")
        ws["A1"].fill = self.EFORM_DARK_FILL
        ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 24

        ws.merge_cells("A2:J2")
        ws["A2"] = "SPT TAHUNAN PAJAK PENGHASILAN (PPh) WAJIB PAJAK ORANG PRIBADI"
        ws["A2"].font = Font(size=13, bold=True, color="FFFFFF")
        ws["A2"].fill = self.EFORM_DARK_FILL
        ws["A2"].alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[2].height = 27

        ws.merge_cells("A3:H3")
        ws["A3"] = "INDUK"
        ws["A3"].font = Font(size=12, bold=True)
        ws["A3"].alignment = Alignment(horizontal="left", vertical="center")
        ws.merge_cells("I3:J3")
        ws["I3"] = f"HALAMAN {page_no}"
        ws["I3"].font = Font(size=11, bold=True)
        ws["I3"].alignment = Alignment(horizontal="right", vertical="center")

        header_items = [
            ("A5", "TAHUN PAJAK", int(data.tahun_pajak or 0)),
            ("D5", "PERIODE", "1 s.d 12"),
            ("G5", "STATUS", "NORMAL"),
            ("I5", "REVISION", int(revision or 0)),
        ]
        for coord, label, value in header_items:
            col = ws[coord].column
            row = ws[coord].row
            ws.cell(row, col).value = label
            ws.cell(row, col).font = Font(size=8, bold=True)
            ws.cell(row, col).fill = self.EFORM_SECTION_FILL
            ws.cell(row, col).border = self.EFORM_BORDER
            ws.cell(row, col).alignment = Alignment(horizontal="center", vertical="center")
            ws.merge_cells(start_row=row + 1, start_column=col, end_row=row + 1, end_column=min(col + 1, 10))
            value_cell = ws.cell(row + 1, col)
            value_cell.value = value
            value_cell.font = Font(size=9, bold=True)
            value_cell.fill = self.EFORM_VALUE_FILL
            value_cell.border = self.EFORM_BORDER
            value_cell.alignment = Alignment(horizontal="center", vertical="center")

        return 8

    def _section_header(self, ws, row: int, title: str) -> int:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=10)
        cell = ws.cell(row, 1)
        cell.value = title
        cell.font = Font(size=9, bold=True)
        cell.fill = self.EFORM_SECTION_FILL
        cell.border = self.EFORM_BORDER
        cell.alignment = Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[row].height = 22
        return row + 1

    def _field_row(
        self,
        ws,
        row: int,
        no,
        label: str,
        value="",
        *,
        choice=None,
        money: bool = False,
        wrap: bool = True,
    ) -> int:
        ws.cell(row, 1).value = no
        ws.cell(row, 1).alignment = Alignment(horizontal="center", vertical="top")
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=7)
        label_cell = ws.cell(row, 2)
        label_cell.value = label
        label_cell.alignment = Alignment(vertical="top", wrap_text=wrap)

        if choice is not None:
            ws.cell(row, 8).value = choice
            ws.cell(row, 8).alignment = Alignment(horizontal="center", vertical="center")
            ws.cell(row, 8).fill = self.EFORM_VALUE_FILL
        ws.merge_cells(start_row=row, start_column=9, end_row=row, end_column=10)
        value_cell = ws.cell(row, 9)
        value_cell.value = value
        value_cell.fill = self.EFORM_VALUE_FILL
        value_cell.alignment = Alignment(horizontal="right" if money else "left", vertical="top", wrap_text=wrap)
        if money:
            value_cell.number_format = '#,##0;[Red]-#,##0;-'

        for col in range(1, 11):
            ws.cell(row, col).border = self.EFORM_BORDER
        ws.row_dimensions[row].height = 29 if wrap else 22
        return row + 1

    def _render_eform_induk_h1(self, ws, data: FinalizationInput, revision: int) -> None:
        row = self._eform_banner(ws, data, page_no=1, revision=revision)

        row = self._section_header(ws, row, "A. IDENTITAS WAJIB PAJAK")
        row = self._field_row(ws, row, "1", "NIK/NPWP", str(data.npwp or ""))
        row = self._field_row(ws, row, "2", "Nama", str(data.nama_wp or "").upper())
        row = self._field_row(ws, row, "3", "Jenis ID", "")
        row = self._field_row(ws, row, "4", "No. ID", str(data.npwp or ""))
        row = self._field_row(ws, row, "5", "No. Telepon", "")
        row = self._field_row(ws, row, "6", "Email", "")
        row = self._field_row(
            ws, row, "7",
            "Status Kewajiban Perpajakan Suami dan Istri",
            "",
        )
        row = self._field_row(ws, row, "8", "NIK/NPWP Suami/Istri", "")

        row = self._section_header(ws, row, "B. IKHTISAR PENGHASILAN NETO")
        total_netto_bupot = sum(
            self._safe_float(item.bruto) - self._safe_float(item.pengurang)
            for item in (data.bupot_rows or [])
        )
        domestic_other = self._pph_value(data, "penghasilan_neto_lainnya")
        foreign_income = 0.0
        business_net = 0.0

        row = self._field_row(
            ws, row, "1a",
            "Apakah Anda menerima penghasilan dalam negeri dari pekerjaan?",
            self._money_excel(total_netto_bupot),
            choice=self._yes_no_code(abs(total_netto_bupot) > 0.5),
            money=True,
        )
        row = self._field_row(
            ws, row, "1b",
            "Penghasilan neto dari usaha dan/atau pekerjaan bebas",
            self._money_excel(business_net),
            choice=self._yes_no_code(abs(business_net) > 0.5),
            money=True,
        )
        row = self._field_row(
            ws, row, "1c",
            "Apakah Anda menerima penghasilan dalam negeri lainnya?",
            self._money_excel(domestic_other),
            choice=self._yes_no_code(abs(domestic_other) > 0.5),
            money=True,
        )
        row = self._field_row(
            ws, row, "1d",
            "Apakah Anda menerima penghasilan luar negeri?",
            self._money_excel(foreign_income),
            choice=1,
            money=True,
        )

        row = self._section_header(ws, row, "C. PERHITUNGAN PPh TERUTANG")
        neto_setahun = self._pph_value(
            data,
            "penghasilan_neto_sebelum_pengurang",
            total_netto_bupot + domestic_other,
        )
        pengurang = self._pph_value(data, "pengurang_penghasilan_neto")
        neto_after = self._pph_value(
            data,
            "penghasilan_neto_gabungan",
            neto_setahun - pengurang,
        )
        ptkp = self._pph_value(data, "ptkp")
        pkp = self._pph_value(data, "pkp")
        pph_terutang = self._pph_value(data, "pph_terutang")

        row = self._field_row(ws, row, "2", "Penghasilan neto setahun (1a+1b+1c+1d)", self._money_excel(neto_setahun), money=True)
        row = self._field_row(
            ws, row, "3",
            "Apakah terdapat pengurang penghasilan neto seperti kompensasi kerugian atau zakat/sumbangan keagamaan yang bersifat wajib?",
            self._money_excel(pengurang),
            choice=self._yes_no_code(abs(pengurang) > 0.5),
            money=True,
        )
        row = self._field_row(ws, row, "4", "Penghasilan neto setelah pengurang penghasilan neto (2-3)", self._money_excel(neto_after), money=True)
        row = self._field_row(ws, row, "5", f"Penghasilan tidak kena pajak ({data.status_ptkp or '-'})", self._money_excel(ptkp), money=True)
        row = self._field_row(ws, row, "6", "Penghasilan kena pajak (4-5)", self._money_excel(pkp), money=True)
        row = self._field_row(ws, row, "7", "PPh terutang", self._money_excel(pph_terutang), money=True)
        row = self._field_row(ws, row, "8", "Apakah terdapat pengurang PPh terutang?", 0, choice=1, money=True)
        row = self._field_row(ws, row, "9", "PPh terutang setelah pengurang PPh terutang (7-8)", self._money_excel(pph_terutang), money=True)

        row = self._section_header(ws, row, "D. KREDIT PAJAK")
        kredit = self._pph_value(data, "kredit_pajak")
        pph25 = self._pph_value(data, "pph25")
        row = self._field_row(
            ws, row, "10a",
            "Apakah terdapat PPh yang telah dipotong/dipungut oleh pihak lain?",
            self._money_excel(kredit),
            choice=self._yes_no_code(abs(kredit) > 0.5),
            money=True,
        )
        row = self._field_row(ws, row, "10b", "Angsuran PPh Pasal 25", self._money_excel(pph25), money=True)
        row = self._field_row(ws, row, "10c", "STP PPh Pasal 25 (Hanya pokok pajak)", 0, money=True)
        self._field_row(
            ws, row, "10d",
            "Apakah Anda menerima pengembalian/pengurangan kredit PPh luar negeri yang telah dikreditkan?",
            0,
            choice=1,
            money=True,
        )

    def _render_eform_induk_h2(self, ws, data: FinalizationInput, revision: int) -> None:
        row = self._eform_banner(ws, data, page_no=2, revision=revision)

        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
        ws.cell(row, 1).value = f"NIK/NPWP  {data.npwp}"
        ws.merge_cells(start_row=row, start_column=6, end_row=row, end_column=10)
        ws.cell(row, 6).value = f"TAHUN PAJAK/BAGIAN TAHUN PAJAK  {data.tahun_pajak}"
        for col in (1, 6):
            ws.cell(row, col).font = Font(bold=True)
            ws.cell(row, col).fill = self.EFORM_LIGHT_FILL
            ws.cell(row, col).border = self.EFORM_BORDER
        row += 2

        kurang_lebih = self._pph_value(data, "kurang_lebih_bayar")
        if abs(kurang_lebih) < 0.5:
            pph_terutang = self._pph_value(data, "pph_terutang")
            kredit = self._pph_value(data, "kredit_pajak")
            pph25 = self._pph_value(data, "pph25")
            kurang_lebih = pph_terutang - kredit - pph25

        row = self._section_header(ws, row, "E. PPh KURANG/LEBIH BAYAR")
        row = self._field_row(ws, row, "11a", "PPh kurang/lebih bayar (9-10a-10b-10c+10d)", self._money_excel(kurang_lebih), money=True)
        row = self._field_row(
            ws, row, "11b",
            "Apakah terdapat Surat Keputusan Persetujuan Pengangsuran atau Penundaan Pembayaran Pajak?",
            0,
            choice=1,
            money=True,
        )
        row = self._field_row(ws, row, "11c", "PPh yang masih harus dibayar (11a-11b)", self._money_excel(kurang_lebih), money=True)

        row = self._section_header(ws, row, "F. PEMBETULAN (DIISI JIKA STATUS SPT ADALAH PEMBETULAN)")
        row = self._field_row(ws, row, "12a", "PPh kurang/lebih bayar pada SPT yang dibetulkan", "")
        row = self._field_row(ws, row, "12b", "PPh kurang/lebih bayar karena pembetulan (11a-12a)", "")

        row = self._section_header(ws, row, "G. PERMOHONAN PENGEMBALIAN PPh LEBIH BAYAR")
        row = self._field_row(ws, row, "", "PPh lebih bayar pada 11a atau 12b mohon", "")
        row = self._field_row(ws, row, "", "Nomor rekening / Nama bank / Nama pemilik rekening", "")

        row = self._section_header(ws, row, "H. ANGSURAN PPh PASAL 25 TAHUN PAJAK BERIKUTNYA")
        row = self._field_row(
            ws, row, "13a",
            "Apakah Anda hanya menerima penghasilan teratur dan berkewajiban membayar angsuran PPh Pasal 25 Tahun Pajak berikutnya?",
            0,
            choice=1,
            money=True,
        )
        row = self._field_row(
            ws, row, "13b",
            "Apakah Anda menyusun perhitungan tersendiri angsuran PPh Pasal 25 Tahun Pajak berikutnya?",
            "",
            choice=1,
        )
        row = self._field_row(
            ws, row, "13c",
            "Apakah Anda membayar angsuran PPh Pasal 25 OPPT Tahun Pajak berikutnya?",
            "",
            choice=1,
        )

        row = self._section_header(ws, row, "I. PERNYATAAN TRANSAKSI LAINNYA")
        total_harta = sum(
            self._safe_float(item.nilai_tahun_berjalan)
            for item in (data.harta_current_rows or [])
        )
        other = data.penghasilan_lainnya or {}
        umkm = data.umkm_state or {}
        final_income = sum(
            self._safe_float(value)
            for value in (umkm.get("bruto_bulanan") or [])
        )
        final_income += sum(
            self._safe_float(item.get("dpp"))
            for item in (other.get("final_other_rows") or [])
            if isinstance(item, dict)
        )
        bukan_objek = (
            self._safe_float(other.get("prive_dpp"))
            + self._safe_float(other.get("hibah_warisan_dpp"))
        )

        row = self._field_row(ws, row, "14a", "Harta pada akhir Tahun Pajak", self._money_excel(total_harta), money=True)
        row = self._field_row(ws, row, "14b", "Apakah Anda memiliki utang pada akhir tahun pajak?", 0, choice=1, money=True)
        row = self._field_row(
            ws, row, "14c",
            "Apakah Anda menerima penghasilan yang dikenakan pajak penghasilan bersifat final?",
            self._money_excel(final_income),
            choice=self._yes_no_code(abs(final_income) > 0.5),
            money=True,
        )
        row = self._field_row(
            ws, row, "14d",
            "Apakah Anda menerima penghasilan yang tidak termasuk objek pajak?",
            self._money_excel(bukan_objek),
            choice=self._yes_no_code(abs(bukan_objek) > 0.5),
            money=True,
        )
        row = self._field_row(ws, row, "14e", "Apakah Anda melaporkan biaya penyusutan dan/atau amortisasi fiskal?", "", choice=1)
        row = self._field_row(
            ws, row, "14f",
            "Apakah Anda melaporkan biaya entertainment, biaya promosi, natura/kenikmatan, serta piutang yang nyata-nyata tidak dapat ditagih?",
            "",
            choice=1,
        )
        row = self._field_row(
            ws, row, "14g",
            "Apakah Anda menerima dividen dan/atau penghasilan lain dari luar negeri dan melaporkannya sebagai penghasilan tidak termasuk objek pajak?",
            "",
            choice=1,
        )
        row = self._field_row(
            ws, row, "14h",
            "Kelebihan PPh Final atas penghasilan dari usaha dengan peredaran bruto tertentu yang dapat dimintakan pengembalian",
            0,
            money=True,
        )

        row = self._section_header(ws, row, "J. LAMPIRAN TAMBAHAN")
        for no, label in (
            ("15a", "Laporan keuangan/laporan keuangan yang telah diaudit"),
            ("15b", "Bukti pembayaran zakat/sumbangan keagamaan"),
            ("15c", "Bukti pemotongan/pemungutan sehubungan dengan kredit pajak luar negeri"),
            ("15d", "Surat kuasa khusus"),
            ("15e", "Dokumen lainnya"),
        ):
            row = self._field_row(ws, row, no, label, "", choice=1)

        row = self._section_header(ws, row, "K. PERNYATAAN")
        row = self._field_row(
            ws, row, "",
            "Saya menyatakan bahwa apa yang telah diberitahukan di atas beserta lampirannya adalah benar, lengkap, dan jelas.",
            "",
            choice=1,
        )
        row = self._field_row(ws, row, "", "Penandatangan", "Wajib Pajak")
        row = self._field_row(ws, row, "", "NIK/NPWP", str(data.npwp or ""))
        self._field_row(ws, row, "", "Nama", str(data.nama_wp or "").upper())

    def _render_legacy_placeholder(self, ws, data, heading: str, revision: int) -> None:
        ws.merge_cells("A1:H1")
        ws["A1"] = heading
        ws["A1"].font = Font(size=14, bold=True)
        ws["A1"].alignment = Alignment(horizontal="center")

        ws.merge_cells("A2:H2")
        ws["A2"] = "FORMAT LAMA / LEGACY DJP"
        ws["A2"].font = Font(bold=True)
        ws["A2"].alignment = Alignment(horizontal="center")

        rows = [
            ("Nama Wajib Pajak", data.nama_wp),
            ("NPWP", str(data.npwp)),
            ("Tahun Pajak", data.tahun_pajak),
            ("Revision", revision),
            ("Status", "FINAL"),
        ]
        for offset, (label, value) in enumerate(rows, start=4):
            ws.cell(offset, 1).value = label
            ws.cell(offset, 2).value = value
            ws.cell(offset, 1).font = Font(bold=True)

        ws["A11"] = "CATATAN"
        ws["A11"].font = Font(bold=True)
        ws.merge_cells("A12:H14")
        ws["A12"] = (
            "Mulai Lampiran I Halaman 2 dan seluruh lampiran berikutnya "
            "menggunakan format lama/legacy yang sudah dikunci."
        )
        ws["A12"].alignment = Alignment(wrap_text=True, vertical="top")

        for col in range(1, 9):
            ws.column_dimensions[get_column_letter(col)].width = 17
        self._setup_print(ws)

    def _render_form_sheet(self, ws, data, mode: str, heading: str, revision: int) -> None:
        if mode == "NEW_EFORM_H1":
            self._render_eform_induk_h1(ws, data, revision)
            return
        if mode == "NEW_EFORM_H2":
            self._render_eform_induk_h2(ws, data, revision)
            return
        if mode == "LEGACY_L1_H2":
            LegacyLampiranIH2XlsxRenderer().render(ws, data)
            return
        self._render_legacy_placeholder(ws, data, heading, revision)

    def _write_harta_sheet(self, ws, data: FinalizationInput) -> None:
        ws.append(self.HARTA_HEADERS)
        self._style_header(ws, 1, len(self.HARTA_HEADERS))
        for index, row in enumerate(data.harta_current_rows, start=1):
            ws.append([
                index,
                row.kode_eform,
                row.kode_ct,
                row.nama_harta,
                row.nomor_akun_keterangan,
                row.atas_nama,
                row.nama_bank,
                row.tahun_perolehan,
                float(row.nilai_tahun_sebelumnya or 0),
                float(row.nilai_tahun_berjalan or 0),
            ])
        for column in range(1, len(self.HARTA_HEADERS) + 1):
            ws.column_dimensions[get_column_letter(column)].width = 20
        ws.column_dimensions["D"].width = 34
        ws.column_dimensions["E"].width = 30
        self._setup_print(ws, "landscape")

    def _write_bupot_sheet(self, ws, data: FinalizationInput) -> None:
        ws.append(self.BUPOT_HEADERS)
        self._style_header(ws, 1, len(self.BUPOT_HEADERS))
        for index, row in enumerate(data.bupot_rows, start=1):
            ws.append([
                index,
                row.jenis,
                row.no_bupot,
                row.masa,
                row.tahun,
                row.sifat,
                row.status,
                row.npwp_penerima,
                row.nama_penerima,
                row.fasilitas,
                row.jenis_pph,
                row.kop,
                float(row.bruto or 0),
                float(row.dpp_persen or 0),
                float(row.tarif or 0),
                float(row.pengurang or 0),
                float(row.pph_dipotong or 0),
                row.bukti,
                row.no_bukti,
                row.tanggal_bukti,
                row.npwp_pemotong or row.npwp_pemberi_kerja,
                row.nama_pemotong,
                row.tanggal_pemotongan,
                row.mekanisme_sp2d,
                row.no_sp2d,
            ])
        for column in range(1, len(self.BUPOT_HEADERS) + 1):
            ws.column_dimensions[get_column_letter(column)].width = 18
        self._setup_print(ws, "landscape")

    def _write_meta_sheet(
        self,
        ws,
        data: FinalizationInput,
        *,
        revision: int,
        snapshot_hash: str,
    ) -> None:
        meta = {
            "schema": ROUNDTRIP_SCHEMA,
            "npwp": str(data.npwp),
            "nama_wp": str(data.nama_wp),
            "tahun_pajak": int(data.tahun_pajak),
            "base_revision": int(revision or 0),
            "base_snapshot_hash": str(snapshot_hash or ""),
            "exported_at": datetime.now(timezone.utc).isoformat(),
            "harta_original_hash": str(data.harta_original_hash or ""),
        }
        ws["A1"] = "KEY"
        ws["B1"] = "VALUE"
        for row, (key, value) in enumerate(meta.items(), start=2):
            ws.cell(row, 1).value = key
            ws.cell(row, 2).value = value
        ws["D1"] = "JSON"
        ws["D2"] = json.dumps(meta, ensure_ascii=False, sort_keys=True)

    def import_revision(
        self,
        file_path: str | Path,
        *,
        expected_npwp: Optional[str] = None,
        expected_year: Optional[int] = None,
        expected_snapshot_hash: Optional[str] = None,
    ) -> LegacyXlsxRoundTripResult:
        path = Path(file_path)
        result = LegacyXlsxRoundTripResult(source_path=path)
        if not path.is_file():
            result.issues.append(LegacyXlsxIssue("LX_101", "ERROR", "File Excel revisi tidak ditemukan."))
            return result

        try:
            wb = load_workbook(path, data_only=False)
        except Exception as exc:
            result.issues.append(LegacyXlsxIssue("LX_102", "ERROR", f"Workbook tidak dapat dibaca: {exc}"))
            return result

        required = {META_SHEET, DATA_HARTA_SHEET, DATA_BUPOT_SHEET}
        missing = sorted(required - set(wb.sheetnames))
        if missing:
            result.issues.append(
                LegacyXlsxIssue(
                    "LX_103", "ERROR",
                    "Workbook bukan hasil round-trip Tax Converter atau sheet canonical hilang: "
                    + ", ".join(missing),
                )
            )
            return result

        meta_ws = wb[META_SHEET]
        meta = {}
        for row in range(2, meta_ws.max_row + 1):
            key = self._text(meta_ws.cell(row, 1).value)
            if key:
                meta[key] = meta_ws.cell(row, 2).value

        if self._text(meta.get("schema")) != ROUNDTRIP_SCHEMA:
            result.issues.append(LegacyXlsxIssue("LX_104", "ERROR", "Versi schema round-trip tidak dikenali."))
            return result

        result.npwp = self._digits(meta.get("npwp"))
        result.nama_wp = self._text(meta.get("nama_wp"))
        result.tahun_pajak = int(self._number(meta.get("tahun_pajak")))
        result.base_revision = int(self._number(meta.get("base_revision")))
        result.base_snapshot_hash = self._text(meta.get("base_snapshot_hash"))

        if expected_npwp and self._digits(expected_npwp) != result.npwp:
            result.issues.append(LegacyXlsxIssue("LX_105", "ERROR", "NPWP workbook revisi tidak sesuai dengan WP aktif."))
        if expected_year and int(expected_year) != result.tahun_pajak:
            result.issues.append(LegacyXlsxIssue("LX_106", "ERROR", "Tahun Pajak workbook revisi tidak sesuai dengan Worksheet aktif."))
        if expected_snapshot_hash and str(expected_snapshot_hash) != result.base_snapshot_hash:
            result.issues.append(
                LegacyXlsxIssue(
                    "LX_107", "ERROR",
                    "Workbook revisi bukan berasal dari snapshot FINAL aktif yang diharapkan.",
                )
            )
        if result.errors:
            return result

        self._read_harta(wb[DATA_HARTA_SHEET], result)
        self._read_bupot(wb[DATA_BUPOT_SHEET], result)
        return result

    def _read_harta(self, ws, result: LegacyXlsxRoundTripResult) -> None:
        headers = {
            self._text(ws.cell(1, col).value).upper(): col
            for col in range(1, ws.max_column + 1)
        }
        for required in self.HARTA_HEADERS:
            if required not in headers:
                result.issues.append(LegacyXlsxIssue("LX_108", "ERROR", f"Kolom Harta '{required}' hilang."))
                return

        for row in range(2, ws.max_row + 1):
            name = self._text(ws.cell(row, headers["NAMA HARTA"]).value)
            code_ct = self._text(ws.cell(row, headers["KODE CT"]).value)
            if not name and not code_ct:
                continue
            try:
                result.harta_rows.append(
                    WorksheetHartaRow(
                        nomor=len(result.harta_rows) + 1,
                        kode_eform=self._text(ws.cell(row, headers["KODE EFORM"]).value),
                        kode_ct=code_ct,
                        nama_harta=name,
                        nomor_akun_keterangan=self._text(ws.cell(row, headers["NOMOR AKUN / KETERANGAN"]).value),
                        atas_nama=self._text(ws.cell(row, headers["ATAS NAMA"]).value),
                        nama_bank=self._text(ws.cell(row, headers["NAMA BANK"]).value),
                        tahun_perolehan=int(self._number(ws.cell(row, headers["TH PEROLEHAN"]).value)),
                        nilai_tahun_sebelumnya=self._number(ws.cell(row, headers["TAHUN SEBELUMNYA"]).value),
                        nilai_tahun_berjalan=self._number(ws.cell(row, headers["TAHUN BERJALAN"]).value),
                    )
                )
            except (TypeError, ValueError) as exc:
                result.issues.append(
                    LegacyXlsxIssue("LX_109", "ERROR", f"Baris Harta {row} tidak valid: {exc}")
                )

    def _read_bupot(self, ws, result: LegacyXlsxRoundTripResult) -> None:
        headers = {
            self._text(ws.cell(1, col).value).upper(): col
            for col in range(1, ws.max_column + 1)
        }
        for required in self.BUPOT_HEADERS:
            if required not in headers:
                result.issues.append(LegacyXlsxIssue("LX_110", "ERROR", f"Kolom Bupot '{required}' hilang."))
                return

        def text(row, key):
            return self._text(ws.cell(row, headers[key]).value)

        def number(row, key):
            return self._number(ws.cell(row, headers[key]).value)

        for row in range(2, ws.max_row + 1):
            jenis = text(row, "JENIS BUPOT")
            no_bupot = text(row, "NO BUKPOT")
            if not any((jenis, no_bupot, number(row, "BRUTO"), number(row, "PPH"))):
                continue
            npwp_pemotong = self._digits(text(row, "NPWP PEMOTONG"))
            try:
                result.bupot_rows.append(
                    WorksheetBupotRow(
                        jenis=jenis,
                        no_bupot=no_bupot,
                        masa=text(row, "MASA"),
                        tahun=text(row, "TAHUN"),
                        sifat=text(row, "SIFAT"),
                        status=text(row, "STATUS"),
                        npwp_penerima=self._digits(text(row, "NPWP PENERIMA")),
                        nama_penerima=text(row, "NAMA PENERIMA"),
                        fasilitas=text(row, "FASILITAS"),
                        jenis_pph=text(row, "JENIS PPH"),
                        kop=text(row, "KOP"),
                        bruto=number(row, "BRUTO"),
                        dpp_persen=number(row, "DPP PERSEN"),
                        tarif=number(row, "TARIF"),
                        pengurang=number(row, "PENGURANG BRUTO"),
                        pph_dipotong=number(row, "PPH"),
                        bukti=text(row, "BUKTI"),
                        no_bukti=text(row, "NO BUKTI"),
                        tanggal_bukti=text(row, "TANGGAL BUKTI"),
                        npwp_pemotong=npwp_pemotong,
                        nama_pemotong=text(row, "NAMA PEMOTONG"),
                        tanggal_pemotongan=text(row, "TANGGAL PEMOTONGAN"),
                        mekanisme_sp2d=text(row, "MEKANISME SP2D"),
                        no_sp2d=text(row, "NO SP2D"),
                        npwp_pemberi_kerja=npwp_pemotong,
                    )
                )
            except (TypeError, ValueError) as exc:
                result.issues.append(
                    LegacyXlsxIssue("LX_111", "ERROR", f"Baris Bupot {row} tidak valid: {exc}")
                )
