from __future__ import annotations

from collections import defaultdict

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


class LegacyLampiranIIIXlsxRenderer:
    """Renderer Form 1770-III format lama dari snapshot FINAL.

    Mapping dibuat konservatif: data yang hanya tersedia sebagai agregat tidak
    ditebak ke kategori khusus. Penghasilan final yang tak terklasifikasi dan
    UMKM masuk ke baris 16; bukan-objek agregat masuk ke Bagian B baris 6.
    """

    THIN = Side(style="thin", color="000000")
    BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
    VALUE_FILL = PatternFill("solid", fgColor="FFF2CC")
    MONEY = '#,##0;[Red]-#,##0;-'

    FINAL_LABELS = (
        "BUNGA DEPOSITO, TABUNGAN, DISKONTO SBI DAN SURAT BERHARGA NEGARA",
        "BUNGA/DISKONTO OBLIGASI",
        "PENJUALAN SAHAM DI BURSA EFEK",
        "HADIAH UNDIAN",
        "PESANGON, TUNJANGAN HARI TUA DAN TEBUSAN PENSIUN YANG DIBAYAR SEKALIGUS",
        "HONORARIUM ATAS BEBAN APBN/APBD",
        "PENGALIHAN HAK ATAS TANAH DAN/ATAU BANGUNAN",
        "BANGUNAN YANG DITERIMA DALAM RANGKA BANGUN GUNA SERAH",
        "SEWA ATAS TANAH DAN/ATAU BANGUNAN",
        "USAHA JASA KONSTRUKSI",
        "PENYALUR/DEALER/AGEN PRODUK BBM",
        "BUNGA SIMPANAN YANG DIBAYARKAN OLEH KOPERASI KEPADA ANGGOTA",
        "PENGHASILAN DARI TRANSAKSI DERIVATIF",
        "DIVIDEN",
        "PENGHASILAN ISTERI DARI SATU PEMBERI KERJA",
        "PENGHASILAN LAIN YANG DIKENAKAN PAJAK FINAL DAN/ATAU BERSIFAT FINAL",
    )

    NON_OBJECT_LABELS = (
        "BANTUAN/SUMBANGAN/HIBAH",
        "WARISAN",
        "BAGIAN LABA ANGGOTA PERSEROAN KOMANDITER TIDAK ATAS SAHAM, PERSEKUTUAN, FIRMA, DAN KONGSI",
        "KLAIM ASURANSI KESEHATAN, KECELAKAAN, JIWA, DWIGUNA, DAN BEASISWA",
        "BEASISWA",
        "PENGHASILAN LAIN YANG TIDAK TERMASUK OBJEK PAJAK",
    )

    @staticmethod
    def _number(value) -> float:
        if value in (None, ""):
            return 0.0
        if isinstance(value, (int, float)):
            return float(value)
        text = str(value).strip().replace("Rp", "").replace(" ", "")
        try:
            if "," in text and "." in text:
                text = text.replace(".", "").replace(",", ".")
            elif text.count(".") > 1:
                text = text.replace(".", "")
            elif text.count(".") == 1 and len(text.rsplit(".", 1)[1]) == 3:
                text = text.replace(".", "")
            elif "," in text:
                tail = text.rsplit(",", 1)[1]
                text = text.replace(",", "" if len(tail) == 3 else ".")
            return float(text or 0)
        except (TypeError, ValueError):
            return 0.0

    @classmethod
    def _money(cls, value) -> int:
        return int(round(cls._number(value)))

    @staticmethod
    def _value(item, *names):
        for name in names:
            if isinstance(item, dict) and name in item:
                return item.get(name)
            if hasattr(item, name):
                return getattr(item, name)
        return None

    @staticmethod
    def _classify_final_income(keterangan) -> int:
        text = " ".join(str(keterangan or "").upper().split())
        if any(key in text for key in ("DEPOSITO", "TABUNGAN", "DISKONTO SBI", "SURAT BERHARGA NEGARA", " SBN")):
            return 1
        if "OBLIGASI" in text:
            return 2
        if "SAHAM" in text and ("BURSA" in text or "EFEK" in text):
            return 3
        if "HADIAH" in text and "UNDIAN" in text:
            return 4
        if any(key in text for key in ("PESANGON", "TUNJANGAN HARI TUA", "TEBUSAN PENSIUN")):
            return 5
        if "HONORARIUM" in text and any(key in text for key in ("APBN", "APBD")):
            return 6
        if "PENGALIHAN" in text and "TANAH" in text:
            return 7
        if "BANGUN" in text and "GUNA" in text and "SERAH" in text:
            return 8
        if "SEWA" in text and ("TANAH" in text or "BANGUNAN" in text):
            return 9
        if "KONSTRUKSI" in text:
            return 10
        if any(key in text for key in ("BBM", "PENYALUR", "DEALER", "AGEN PRODUK")):
            return 11
        if "KOPERASI" in text and "BUNGA" in text:
            return 12
        if "DERIVATIF" in text:
            return 13
        if "DIVIDEN" in text:
            return 14
        if ("ISTERI" in text or "SUAMI" in text) and "PEMBERI KERJA" in text:
            return 15
        return 16

    @classmethod
    def _sum_values(cls, value) -> float:
        if isinstance(value, dict):
            return sum(cls._number(v) for v in value.values())
        if isinstance(value, (list, tuple)):
            return sum(cls._number(v) for v in value)
        return cls._number(value)

    @classmethod
    def _first_nonzero(cls, mapping, *keys) -> float:
        if not isinstance(mapping, dict):
            return 0.0
        for key in keys:
            if key not in mapping:
                continue
            value = cls._sum_values(mapping.get(key))
            if abs(value) > 0.000001:
                return value
        return 0.0

    def _map_final_rows(self, data):
        buckets = defaultdict(lambda: [0.0, 0.0])
        other = data.penghasilan_lainnya or {}
        rows = []
        if isinstance(other, dict):
            for key in ("final_other_rows", "penghasilan_final_lainnya"):
                value = other.get(key)
                if isinstance(value, (list, tuple)):
                    rows.extend(value)

        for item in rows:
            label = self._value(item, "keterangan", "jenis", "uraian", "nama")
            dpp = self._number(self._value(item, "dpp", "bruto", "penghasilan_bruto"))
            pph = self._number(self._value(item, "pph", "pph_final", "pph_terutang"))
            no = self._classify_final_income(label)
            buckets[no][0] += dpp
            buckets[no][1] += pph

        umkm = data.umkm_state or {}
        if isinstance(umkm, dict):
            bruto = self._first_nonzero(
                umkm,
                "bruto_bulanan",
                "total_bruto",
                "penghasilan_bruto",
                "jumlah_penghasilan_bruto_final",
                "bruto",
            )
            pph = self._first_nonzero(
                umkm,
                "pph_setor",
                "pph_final",
                "total_pph",
                "jumlah_pph_final",
            )
            buckets[16][0] += bruto
            buckets[16][1] += pph

        return buckets

    def _non_object_total(self, data) -> float:
        analysis = getattr(data, "analisis_result", None)
        value = self._number(getattr(analysis, "jumlah_penghasilan_bukan_objek", 0))
        if abs(value) > 0.000001:
            return value
        other = data.penghasilan_lainnya or {}
        return self._first_nonzero(
            other,
            "jumlah_penghasilan_bukan_objek",
            "penghasilan_bukan_objek",
            "non_object_total",
            "bukan_objek_total",
        )

    def _spouse_separate(self, data) -> float:
        other = data.penghasilan_lainnya or {}
        return self._first_nonzero(
            other,
            "penghasilan_pasangan_terpisah",
            "penghasilan_neto_pasangan_terpisah",
            "penghasilan_isteri_suami_terpisah",
        )

    def render(self, ws, data) -> None:
        self._header(ws, data)
        row = self._bagian_a(ws, data)
        row = self._bagian_b(ws, row, data)
        row = self._bagian_c(ws, row, data)
        self._finish(ws, row)

    def _header(self, ws, data) -> None:
        ws.sheet_view.showGridLines = False
        ws.merge_cells("A1:B3")
        ws["A1"] = "FORMULIR\n1770 - III\nKEMENTERIAN KEUANGAN RI\nDIREKTORAT JENDERAL PAJAK"
        ws["A1"].font = Font(size=8, bold=True)
        ws["A1"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        ws.merge_cells("C1:H1")
        ws["C1"] = "LAMPIRAN - III"
        ws["C1"].font = Font(size=9, bold=True)
        ws["C1"].alignment = Alignment(horizontal="center")

        ws.merge_cells("C2:H2")
        ws["C2"] = "SPT TAHUNAN PPh WAJIB PAJAK ORANG PRIBADI"
        ws["C2"].font = Font(size=11, bold=True)
        ws["C2"].alignment = Alignment(horizontal="center")

        ws.merge_cells("C3:H5")
        ws["C3"] = (
            "PENGHASILAN YANG DIKENAKAN PAJAK FINAL DAN/ATAU BERSIFAT FINAL\n"
            "PENGHASILAN YANG TIDAK TERMASUK OBJEK PAJAK\n"
            "PENGHASILAN ISTERI/SUAMI YANG DIKENAKAN PAJAK SECARA TERPISAH"
        )
        ws["C3"].font = Font(size=7, bold=True)
        ws["C3"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        ws.merge_cells("I1:J1")
        ws["I1"] = "TAHUN PAJAK"
        ws["I1"].font = Font(size=8, bold=True)
        ws["I1"].alignment = Alignment(horizontal="center")
        ws.merge_cells("I2:J2")
        ws["I2"] = int(data.tahun_pajak or 0)
        ws["I2"].fill = self.VALUE_FILL
        ws["I2"].font = Font(size=10, bold=True)
        ws["I2"].alignment = Alignment(horizontal="center")

        ws.merge_cells("A7:B7")
        ws["A7"] = "NPWP"
        ws.merge_cells("C7:J7")
        ws["C7"] = str(data.npwp or "")
        ws["C7"].fill = self.VALUE_FILL
        ws["C7"].font = Font(bold=True)

        ws.merge_cells("A8:B8")
        ws["A8"] = "NAMA WAJIB PAJAK"
        ws.merge_cells("C8:J8")
        ws["C8"] = str(data.nama_wp or "").upper()
        ws["C8"].fill = self.VALUE_FILL
        ws["C8"].font = Font(bold=True)
        for row in (7, 8):
            for col in range(1, 11):
                ws.cell(row, col).border = self.BORDER

    def _section(self, ws, row: int, title: str) -> int:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=10)
        cell = ws.cell(row, 1)
        cell.value = title
        cell.font = Font(size=8, bold=True)
        cell.alignment = Alignment(wrap_text=True)
        for col in range(1, 11):
            ws.cell(row, col).border = self.BORDER
        return row + 1

    def _bagian_a(self, ws, data) -> int:
        row = 10
        row = self._section(
            ws,
            row,
            "BAGIAN A : PENGHASILAN YANG DIKENAKAN PAJAK FINAL DAN/ATAU BERSIFAT FINAL",
        )
        spans = ((1, 1), (2, 6), (7, 8), (9, 10))
        for value, (c1, c2) in zip(
            ("NO.", "JENIS PENGHASILAN", "DASAR PENGENAAN PAJAK/PENGHASILAN BRUTO (Rupiah)", "PPh TERUTANG (Rupiah)"),
            spans,
        ):
            if c1 != c2:
                ws.merge_cells(start_row=row, start_column=c1, end_row=row, end_column=c2)
            cell = ws.cell(row, c1)
            cell.value = value
            cell.font = Font(size=7, bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            for col in range(c1, c2 + 1):
                ws.cell(row, col).border = self.BORDER
        ws.row_dimensions[row].height = 36
        row += 1

        buckets = self._map_final_rows(data)
        total_dpp = total_pph = 0.0
        for no, label in enumerate(self.FINAL_LABELS, start=1):
            dpp, pph = buckets[no]
            total_dpp += dpp
            total_pph += pph
            values = (no, label, self._money(dpp), self._money(pph))
            for value, (c1, c2) in zip(values, spans):
                if c1 != c2:
                    ws.merge_cells(start_row=row, start_column=c1, end_row=row, end_column=c2)
                cell = ws.cell(row, c1)
                cell.value = value
                cell.alignment = Alignment(
                    horizontal="right" if c1 in {7, 9} else ("center" if c1 == 1 else "left"),
                    vertical="center",
                    wrap_text=True,
                )
                if c1 in {7, 9}:
                    cell.number_format = self.MONEY
                    if abs(self._number(value)) > 0.000001:
                        cell.fill = self.VALUE_FILL
                for col in range(c1, c2 + 1):
                    ws.cell(row, col).border = self.BORDER
            ws.row_dimensions[row].height = 24
            row += 1

        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
        ws.cell(row, 1).value = "17. JUMLAH (1 s.d. 16)"
        ws.cell(row, 1).font = Font(bold=True)
        ws.cell(row, 1).alignment = Alignment(horizontal="right")
        for c1, c2, value in ((7, 8, total_dpp), (9, 10, total_pph)):
            ws.merge_cells(start_row=row, start_column=c1, end_row=row, end_column=c2)
            ws.cell(row, c1).value = self._money(value)
            ws.cell(row, c1).number_format = self.MONEY
            ws.cell(row, c1).font = Font(bold=True)
            ws.cell(row, c1).fill = self.VALUE_FILL
            ws.cell(row, c1).alignment = Alignment(horizontal="right")
        for col in range(1, 11):
            ws.cell(row, col).border = self.BORDER
        return row + 2

    def _bagian_b(self, ws, row: int, data) -> int:
        row = self._section(ws, row, "BAGIAN B : PENGHASILAN YANG TIDAK TERMASUK OBJEK PAJAK")
        spans = ((1, 1), (2, 7), (8, 10))
        for value, (c1, c2) in zip(("NO.", "JENIS PENGHASILAN", "PENGHASILAN BRUTO (Rupiah)"), spans):
            if c1 != c2:
                ws.merge_cells(start_row=row, start_column=c1, end_row=row, end_column=c2)
            ws.cell(row, c1).value = value
            ws.cell(row, c1).font = Font(size=7, bold=True)
            ws.cell(row, c1).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            for col in range(c1, c2 + 1):
                ws.cell(row, col).border = self.BORDER
        row += 1

        total = self._non_object_total(data)
        for no, label in enumerate(self.NON_OBJECT_LABELS, start=1):
            value = total if no == 6 else 0
            values = (no, label, self._money(value))
            for item, (c1, c2) in zip(values, spans):
                if c1 != c2:
                    ws.merge_cells(start_row=row, start_column=c1, end_row=row, end_column=c2)
                cell = ws.cell(row, c1)
                cell.value = item
                cell.alignment = Alignment(
                    horizontal="right" if c1 == 8 else ("center" if c1 == 1 else "left"),
                    vertical="center",
                    wrap_text=True,
                )
                if c1 == 8:
                    cell.number_format = self.MONEY
                    if abs(self._number(item)) > 0.000001:
                        cell.fill = self.VALUE_FILL
                for col in range(c1, c2 + 1):
                    ws.cell(row, col).border = self.BORDER
            ws.row_dimensions[row].height = 24
            row += 1

        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=7)
        ws.cell(row, 1).value = "JUMLAH BAGIAN B"
        ws.cell(row, 1).font = Font(bold=True)
        ws.cell(row, 1).alignment = Alignment(horizontal="right")
        ws.merge_cells(start_row=row, start_column=8, end_row=row, end_column=10)
        ws.cell(row, 8).value = self._money(total)
        ws.cell(row, 8).number_format = self.MONEY
        ws.cell(row, 8).font = Font(bold=True)
        ws.cell(row, 8).fill = self.VALUE_FILL
        ws.cell(row, 8).alignment = Alignment(horizontal="right")
        for col in range(1, 11):
            ws.cell(row, col).border = self.BORDER
        return row + 2

    def _bagian_c(self, ws, row: int, data) -> int:
        row = self._section(
            ws,
            row,
            "BAGIAN C : PENGHASILAN ISTERI/SUAMI YANG DIKENAKAN PAJAK SECARA TERPISAH",
        )
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=7)
        ws.cell(row, 1).value = "PENGHASILAN NETO ISTERI/SUAMI YANG DIKENAKAN PAJAK SECARA TERPISAH"
        ws.cell(row, 1).font = Font(size=7, bold=True)
        ws.cell(row, 1).alignment = Alignment(wrap_text=True)
        ws.merge_cells(start_row=row, start_column=8, end_row=row, end_column=10)
        ws.cell(row, 8).value = self._money(self._spouse_separate(data))
        ws.cell(row, 8).number_format = self.MONEY
        ws.cell(row, 8).fill = self.VALUE_FILL
        ws.cell(row, 8).alignment = Alignment(horizontal="right")
        for col in range(1, 11):
            ws.cell(row, col).border = self.BORDER
        return row

    def _finish(self, ws, row: int) -> None:
        widths = {
            "A": 4.0, "B": 7.0, "C": 11.0, "D": 11.0, "E": 11.0,
            "F": 11.0, "G": 11.0, "H": 11.0, "I": 11.0, "J": 11.0,
        }
        for col, width in widths.items():
            ws.column_dimensions[col].width = width
        ws.page_setup.orientation = "portrait"
        ws.page_setup.paperSize = ws.PAPERSIZE_LEGAL
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 1
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_margins.left = 0.18
        ws.page_margins.right = 0.18
        ws.page_margins.top = 0.22
        ws.page_margins.bottom = 0.22
        ws.print_area = f"A1:J{row}"
