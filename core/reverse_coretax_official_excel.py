from __future__ import annotations

import re
import shutil
from dataclasses import dataclass, field
from pathlib import Path
from typing import Dict, List, Optional, Tuple

from openpyxl import load_workbook

from core.coretax_official_schema import get_official_schema
from core.reverse_coretax_mapping import (
    CATEGORY_ORDER,
    ReverseCoretaxMappingService,
    ReverseCoretaxPackage,
    ReverseCoretaxRow,
)

DATA_START_ROW = 4
HEADER_ROW = 3
METADATA_ROWS = (1, 2)


@dataclass(frozen=True)
class OfficialExcelExportIssue:
    code: str
    severity: str
    message: str
    category: Optional[str] = None


@dataclass
class OfficialExcelExportResult:
    output_dir: Path
    files: Dict[str, Path] = field(default_factory=dict)
    row_counts: Dict[str, int] = field(default_factory=dict)
    template_files: Dict[str, Path] = field(default_factory=dict)
    issues: List[OfficialExcelExportIssue] = field(default_factory=list)

    @property
    def errors(self) -> List[OfficialExcelExportIssue]:
        return [item for item in self.issues if item.severity == "ERROR"]

    @property
    def warnings(self) -> List[OfficialExcelExportIssue]:
        return [item for item in self.issues if item.severity == "WARNING"]

    @property
    def ok(self) -> bool:
        return not self.errors and len(self.files) == 6


class OfficialCoretaxExcelExporter:
    """Stage 8D.4C - export 6 Excel dengan template Coretax asli."""

    def __init__(self, template_dir: str | Path):
        self.template_dir = Path(template_dir)

    @staticmethod
    def _normalize_filename(value: str) -> str:
        value = str(value or "").lower()
        value = re.sub(r"[^a-z0-9]+", " ", value)
        return " ".join(value.split())

    def _find_template(self, category: str) -> Optional[Path]:
        schema = get_official_schema(category)
        hint = self._normalize_filename(schema.excel_filename_hint)
        candidates: List[Path] = []

        for path in self.template_dir.rglob("*.xlsx"):
            normalized = self._normalize_filename(path.stem)
            if hint in normalized:
                candidates.append(path)

        if not candidates:
            return None

        candidates.sort(
            key=lambda path: (
                abs(len(self._normalize_filename(path.stem)) - len(hint)),
                str(path).lower(),
            )
        )
        return candidates[0]

    def _resolve_templates(
        self,
        result: OfficialExcelExportResult,
    ) -> Dict[str, Path]:
        resolved: Dict[str, Path] = {}

        if not self.template_dir.exists():
            result.issues.append(
                OfficialExcelExportIssue(
                    "RCX4C_001",
                    "ERROR",
                    f"Folder template tidak ditemukan: {self.template_dir}",
                )
            )
            return resolved

        for category in CATEGORY_ORDER:
            template = self._find_template(category)
            if template is None:
                schema = get_official_schema(category)
                result.issues.append(
                    OfficialExcelExportIssue(
                        "RCX4C_002",
                        "ERROR",
                        (
                            f"Template resmi tidak ditemukan untuk {category}. "
                            f"Hint filename: {schema.excel_filename_hint}"
                        ),
                        category,
                    )
                )
                continue

            resolved[category] = template
            result.template_files[category] = template

        return resolved

    @staticmethod
    def _clean_header(value: object) -> str:
        return str(value or "").strip()

    def _validate_template(
        self,
        template_path: Path,
        category: str,
        result: OfficialExcelExportResult,
    ) -> bool:
        schema = get_official_schema(category)

        try:
            wb = load_workbook(template_path, data_only=False)
        except Exception as exc:
            result.issues.append(
                OfficialExcelExportIssue(
                    "RCX4C_003",
                    "ERROR",
                    f"Template tidak dapat dibuka: {template_path.name}: {exc}",
                    category,
                )
            )
            return False

        try:
            if schema.excel_sheet not in wb.sheetnames:
                result.issues.append(
                    OfficialExcelExportIssue(
                        "RCX4C_004",
                        "ERROR",
                        (
                            f"Sheet '{schema.excel_sheet}' tidak ditemukan "
                            f"pada {template_path.name}."
                        ),
                        category,
                    )
                )
                return False

            ws = wb[schema.excel_sheet]
            actual_headers = tuple(
                self._clean_header(ws.cell(HEADER_ROW, column_index).value)
                for column_index in range(1, len(schema.excel_headers) + 1)
            )
            expected_headers = tuple(
                self._clean_header(header)
                for header in schema.excel_headers
            )

            if actual_headers != expected_headers:
                result.issues.append(
                    OfficialExcelExportIssue(
                        "RCX4C_005",
                        "ERROR",
                        (
                            f"Header template {category} tidak sama dengan "
                            "kontrak resmi Stage 8D.4A."
                        ),
                        category,
                    )
                )
                return False

            return True
        finally:
            wb.close()

    @staticmethod
    def _normalize_label(value: object) -> str:
        text = str(value or "").strip().upper()
        text = re.sub(r"[^A-Z0-9]+", " ", text)
        return " ".join(text.split())

    @classmethod
    def _set_label_value(
        cls,
        ws,
        label_terms: Tuple[str, ...],
        value: object,
    ) -> bool:
        normalized_terms = tuple(cls._normalize_label(term) for term in label_terms)

        for row_index in METADATA_ROWS:
            for column_index in range(1, 21):
                cell = ws.cell(row_index, column_index)
                label = cls._normalize_label(cell.value)
                if not label:
                    continue
                if not any(term in label for term in normalized_terms):
                    continue

                ws.cell(row_index, column_index + 1).value = value
                return True

        return False

    def _write_metadata(
        self,
        ws,
        package: ReverseCoretaxPackage,
        category: str,
        result: OfficialExcelExportResult,
    ) -> None:
        if not self._set_label_value(
            ws,
            ("NPWP", "NPWP WAJIB PAJAK", "TIN"),
            package.npwp,
        ):
            result.issues.append(
                OfficialExcelExportIssue(
                    "RCX4C_101",
                    "WARNING",
                    "Label NPWP tidak ditemukan pada metadata baris 1-2 template.",
                    category,
                )
            )

        if not self._set_label_value(
            ws,
            ("TAHUN PAJAK", "TAX YEAR"),
            package.tahun_pajak,
        ):
            result.issues.append(
                OfficialExcelExportIssue(
                    "RCX4C_102",
                    "WARNING",
                    "Label Tahun Pajak tidak ditemukan pada metadata baris 1-2 template.",
                    category,
                )
            )

    @staticmethod
    def _meta(
        row: ReverseCoretaxRow,
        key: str,
        fallback: object = "",
    ) -> object:
        value = row.official_metadata.get(key)
        if value is None or value == "":
            return fallback
        return value

    @classmethod
    def _official_row_values(
        cls,
        row: ReverseCoretaxRow,
    ) -> Tuple[object, ...]:
        category = row.kategori

        if category == "KAS":
            return (
                row.kode_harta,
                cls._meta(row, "account_number", row.nomor_akun_keterangan),
                cls._meta(row, "account_on_behalf_of", row.atas_nama),
                cls._meta(row, "bank_name", row.nama_bank),
                cls._meta(row, "country"),
                cls._meta(row, "year", row.tahun_perolehan),
                cls._meta(row, "balance", row.nilai),
                cls._meta(row, "remarks"),
            )

        if category == "PIUTANG":
            return (
                row.kode_harta,
                cls._meta(row, "country"),
                cls._meta(row, "identity_number"),
                cls._meta(row, "receivable_name", row.atas_nama),
                cls._meta(row, "receivable_value", row.nilai),
                cls._meta(row, "year", row.tahun_perolehan),
                cls._meta(row, "receivable_balance", row.nilai),
                cls._meta(row, "remarks"),
            )

        if category == "INVESTASI":
            return (
                row.kode_harta,
                cls._meta(row, "country"),
                cls._meta(row, "institution_tin"),
                cls._meta(row, "institution_name", row.atas_nama or row.nama_bank),
                cls._meta(row, "account_number", row.nomor_akun_keterangan),
                cls._meta(row, "cost_of_acquisition", row.nilai),
                cls._meta(row, "year", row.tahun_perolehan),
                cls._meta(row, "current_balance", row.nilai),
                cls._meta(row, "remarks"),
            )

        if category == "BERGERAK":
            return (
                row.kode_harta,
                cls._meta(row, "asset_model", row.nama_harta),
                cls._meta(
                    row,
                    "police_registration_number",
                    row.nomor_akun_keterangan,
                ),
                cls._meta(row, "ownership_type"),
                cls._meta(row, "ownership_tin"),
                cls._meta(row, "ownership_name", row.atas_nama),
                cls._meta(row, "year", row.tahun_perolehan),
                cls._meta(row, "cost_of_acquisition", row.nilai),
                cls._meta(row, "fair_market_value", row.nilai),
                cls._meta(row, "remarks"),
            )

        if category == "HTB":
            return (
                row.kode_harta,
                cls._meta(row, "location_of_asset", row.nomor_akun_keterangan),
                cls._meta(row, "property_size_land"),
                cls._meta(row, "property_size_building"),
                cls._meta(row, "source_of_ownership"),
                cls._meta(row, "certificate_number"),
                cls._meta(row, "year", row.tahun_perolehan),
                cls._meta(row, "cost_of_acquisition", row.nilai),
                cls._meta(row, "fair_market_value", row.nilai),
                cls._meta(row, "remarks"),
            )

        if category == "LAINNYA":
            return (
                row.kode_harta,
                cls._meta(row, "year", row.tahun_perolehan),
                cls._meta(row, "account_number", row.nomor_akun_keterangan),
                cls._meta(row, "additional_information", row.nama_harta),
                cls._meta(row, "cost_of_acquisition", row.nilai),
                cls._meta(row, "current_value", row.nilai),
                cls._meta(row, "remarks"),
            )

        raise ValueError(f"Kategori Coretax tidak dikenal: {category}")

    @staticmethod
    def _clear_existing_data(ws, column_count: int) -> None:
        max_row = max(ws.max_row, DATA_START_ROW)
        for row_index in range(DATA_START_ROW, max_row + 1):
            for column_index in range(1, column_count + 1):
                ws.cell(row_index, column_index).value = None

    @staticmethod
    def _write_rows(
        ws,
        rows: List[ReverseCoretaxRow],
        row_value_factory,
    ) -> None:
        for offset, source_row in enumerate(rows):
            excel_row = DATA_START_ROW + offset
            for column_index, value in enumerate(
                row_value_factory(source_row),
                start=1,
            ):
                ws.cell(excel_row, column_index).value = value

    def _export_category(
        self,
        package: ReverseCoretaxPackage,
        category: str,
        template_path: Path,
        output_dir: Path,
        result: OfficialExcelExportResult,
    ) -> Optional[Path]:
        schema = get_official_schema(category)
        target = output_dir / template_path.name
        shutil.copy2(template_path, target)

        try:
            wb = load_workbook(target, data_only=False)
            ws = wb[schema.excel_sheet]

            self._write_metadata(ws, package, category, result)
            self._clear_existing_data(ws, len(schema.excel_headers))

            rows = package.rows_by_category.get(category, [])
            self._write_rows(ws, rows, self._official_row_values)

            wb.save(target)
            wb.close()
            return target

        except Exception as exc:
            result.issues.append(
                OfficialExcelExportIssue(
                    "RCX4C_201",
                    "ERROR",
                    f"Gagal menulis template {template_path.name}: {exc}",
                    category,
                )
            )
            try:
                target.unlink(missing_ok=True)
            except OSError:
                pass
            return None

    def export_package(
        self,
        package: ReverseCoretaxPackage,
        output_dir: str | Path,
    ) -> OfficialExcelExportResult:
        target_dir = Path(output_dir)
        result = OfficialExcelExportResult(output_dir=target_dir)

        if not package.can_export:
            result.issues.append(
                OfficialExcelExportIssue(
                    "RCX4C_301",
                    "ERROR",
                    "Paket reverse Coretax belum valid untuk export resmi.",
                )
            )
            return result

        templates = self._resolve_templates(result)
        if result.errors:
            return result

        if not all(
            self._validate_template(templates[category], category, result)
            for category in CATEGORY_ORDER
        ):
            return result

        target_dir.mkdir(parents=True, exist_ok=True)

        for category in CATEGORY_ORDER:
            rows = package.rows_by_category.get(category, [])
            target = self._export_category(
                package,
                category,
                templates[category],
                target_dir,
                result,
            )
            if target is not None:
                result.files[category] = target
                result.row_counts[category] = len(rows)

        if result.ok:
            result.issues.append(
                OfficialExcelExportIssue(
                    "RCX4C_INFO",
                    "INFO",
                    (
                        "Enam file Excel Coretax resmi berhasil dibuat dari "
                        "snapshot FINAL menggunakan template asli."
                    ),
                )
            )

        return result

    def export_active_final(
        self,
        npwp: str,
        tahun_pajak: int,
        output_dir: str | Path,
        *,
        db_path: Optional[str | Path] = None,
    ) -> OfficialExcelExportResult:
        package = ReverseCoretaxMappingService(
            db_path=db_path
        ).build_active_final(npwp, tahun_pajak)

        result = self.export_package(package, output_dir)

        for issue in package.issues:
            if issue.severity in {"ERROR", "WARNING"}:
                result.issues.append(
                    OfficialExcelExportIssue(
                        issue.code,
                        issue.severity,
                        issue.message,
                    )
                )

        return result
