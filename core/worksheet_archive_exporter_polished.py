from __future__ import annotations

from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from PySide6.QtCore import Qt, QRectF
from PySide6.QtGui import QBrush, QColor, QFont, QFontMetricsF, QPainter, QPen, QPdfWriter

from core.finalization import FinalizationInput
from core.worksheet_archive_exporter_styled import StyledWorksheetArchiveExporter


class PolishedWorksheetArchiveExporter(StyledWorksheetArchiveExporter):
    """Final visual polish Stage 8B.2 untuk arsip Excel/PDF.

    Fokus:
    - identifier pajak/rekening selalu diperlakukan sebagai TEXT di Excel;
    - proporsi kolom, wrapping, subtotal/total dan print layout lebih rapi;
    - PDF memakai tabel dengan wrapping dinamis, spacing section dan nomor halaman.
    """

    TEXT_FORMAT = "@"
    TOTAL_FILL_STRONG = PatternFill("solid", fgColor="DCE6EF")

    def _style_excel(self, path: Path, data: FinalizationInput) -> None:
        super()._style_excel(path, data)
        wb = load_workbook(path)
        annual = wb[str(data.tahun_pajak)]
        simulasi = wb["SIMULASI I"]

        self._polish_annual(annual, data)
        self._polish_simulasi(simulasi, data)
        wb.save(path)

    def _polish_annual(self, ws, data: FinalizationInput) -> None:
        # Identitas: paksa string agar tidak pernah berubah menjadi scientific notation.
        ws["B4"] = str(data.nama_wp or "")
        ws["B5"] = str(data.npwp or "")
        ws["B5"].number_format = self.TEXT_FORMAT
        ws["B5"].alignment = Alignment(horizontal="left", vertical="center")

        bupot_header = 8
        bupot_count = len(data.bupot_rows)
        for index, item in enumerate(data.bupot_rows, start=bupot_header + 1):
            ws.cell(index, 2).value = str(item.jenis or "")
            ws.cell(index, 3).value = str(item.npwp_pemberi_kerja or "")
            ws.cell(index, 4).value = str(item.no_bupot or "")
            for col in (2, 3, 4):
                ws.cell(index, col).number_format = self.TEXT_FORMAT
                ws.cell(index, col).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            ws.row_dimensions[index].height = 22

        total_row = bupot_header + bupot_count + 1
        for col in range(1, 8):
            ws.cell(total_row, col).fill = self.TOTAL_FILL_STRONG
            ws.cell(total_row, col).font = Font(bold=True)
            ws.cell(total_row, col).alignment = Alignment(vertical="center")
        ws.row_dimensions[total_row].height = 23

        # Lebar dibuat untuk penggunaan nyata di layar dan print.
        widths = {
            "A": 8,
            "B": 13,
            "C": 24,
            "D": 23,
            "E": 18,
            "F": 18,
            "G": 18,
        }
        for col, width in widths.items():
            ws.column_dimensions[col].width = width

        # Section setelah Bupot dibuat lebih konsisten dan tidak terlalu padat.
        for row in range(1, ws.max_row + 1):
            label = str(ws.cell(row, 1).value or ws.cell(row, 2).value or "").strip().upper()
            if label in {
                "PEREDARAN BRUTO UMKM",
                "PENGHASILAN LAINNYA",
                "PENGURANG PENGHASILAN NETO",
            }:
                ws.row_dimensions[row].height = 24
            for col in range(1, 8):
                ws.cell(row, col).alignment = Alignment(
                    horizontal=ws.cell(row, col).alignment.horizontal,
                    vertical="center",
                    wrap_text=True,
                )

        ws.page_margins.left = 0.25
        ws.page_margins.right = 0.25
        ws.page_margins.top = 0.4
        ws.page_margins.bottom = 0.4
        ws.print_options.horizontalCentered = True

    def _polish_simulasi(self, ws, data: FinalizationInput) -> None:
        ws["C3"] = str(data.nama_wp or "")
        ws["C4"] = str(data.npwp or "")
        ws["C4"].number_format = self.TEXT_FORMAT
        ws["C4"].alignment = Alignment(horizontal="left", vertical="center")

        header = 8
        for idx, item in enumerate(data.harta_current_rows, start=header + 1):
            # Kode dan identifier harus tetap utuh sebagai teks.
            ws.cell(idx, 2).value = str(item.kode_eform or "")
            ws.cell(idx, 3).value = str(item.kode_ct or "")
            ws.cell(idx, 5).value = str(item.nomor_akun_keterangan or "")
            ws.cell(idx, 6).value = str(item.atas_nama or "")
            ws.cell(idx, 7).value = str(item.nama_bank or "")
            for col in (2, 3, 5, 6, 7):
                ws.cell(idx, col).number_format = self.TEXT_FORMAT
            for col in range(1, 11):
                ws.cell(idx, col).alignment = Alignment(vertical="center", wrap_text=True)
            ws.row_dimensions[idx].height = 30

        total_row = header + len(data.harta_current_rows) + 1
        for col in range(1, 11):
            ws.cell(total_row, col).fill = self.TOTAL_FILL_STRONG
            ws.cell(total_row, col).font = Font(bold=True)
        ws.row_dimensions[total_row].height = 24

        widths = {
            "A": 7,
            "B": 11,
            "C": 10,
            "D": 36,
            "E": 30,
            "F": 20,
            "G": 20,
            "H": 13,
            "I": 18,
            "J": 18,
        }
        for col, width in widths.items():
            ws.column_dimensions[col].width = width

        # Analisis dibuat lebih compact.
        for row in range(total_row + 1, ws.max_row + 1):
            if any(ws.cell(row, col).value not in (None, "") for col in range(1, 11)):
                ws.row_dimensions[row].height = 22

        ws.page_margins.left = 0.2
        ws.page_margins.right = 0.2
        ws.page_margins.top = 0.35
        ws.page_margins.bottom = 0.35
        ws.print_options.horizontalCentered = True

    # ------------------------------------------------------------------ PDF
    def _paint_pdf(self, painter: QPainter, writer: QPdfWriter, data: FinalizationInput) -> None:
        page = writer.pageLayout().paintRectPixels(writer.resolution())
        margin = 26.0
        usable_width = float(page.width()) - margin * 2
        bottom = float(page.height()) - margin - 20
        y = margin
        page_no = 1

        painter.setPen(QPen(QColor(86, 96, 106), 0.65))

        def draw_footer() -> None:
            font = QFont("Arial", 7)
            painter.setFont(font)
            painter.setPen(QPen(QColor(110, 116, 122), 0.5))
            footer = QRectF(margin, float(page.height()) - margin - 12, usable_width, 12)
            painter.drawText(footer, Qt.AlignCenter | Qt.AlignVCenter, f"Arsip Kertas Kerja — TAX_CONVERTER L-1 • Halaman {page_no}")
            painter.setPen(QPen(QColor(86, 96, 106), 0.65))

        def new_page() -> None:
            nonlocal y, page_no
            draw_footer()
            writer.newPage()
            page_no += 1
            y = margin

        def ensure(height: float) -> None:
            if y + height > bottom:
                new_page()

        def title(text: str, size=14, height=27) -> None:
            nonlocal y
            ensure(height)
            font = QFont("Arial", size)
            font.setBold(True)
            painter.setFont(font)
            painter.drawText(QRectF(margin, y, usable_width, height), Qt.AlignCenter | Qt.AlignVCenter, text)
            y += height

        def section(text: str) -> None:
            nonlocal y
            ensure(27)
            rect = QRectF(margin, y + 3, usable_width, 23)
            painter.fillRect(rect, QBrush(QColor(226, 234, 241)))
            painter.drawRect(rect)
            font = QFont("Arial", 9)
            font.setBold(True)
            painter.setFont(font)
            painter.drawText(rect.adjusted(7, 0, -5, 0), Qt.AlignLeft | Qt.AlignVCenter, text)
            y += 29

        def wrapped_lines(text: str, width: float, font: QFont) -> list[str]:
            value = str(text or "")
            if not value:
                return [""]
            metrics = QFontMetricsF(font)
            available = max(12.0, width - 7.0)
            words = value.split()
            lines: list[str] = []
            current = ""
            for word in words:
                candidate = word if not current else f"{current} {word}"
                if metrics.horizontalAdvance(candidate) <= available:
                    current = candidate
                else:
                    if current:
                        lines.append(current)
                    current = word
            if current:
                lines.append(current)
            return lines or [value]

        def table(
            headers: Iterable[str],
            rows: Iterable[Iterable[Any]],
            widths: Iterable[float],
            *,
            font_size: int = 7,
            min_height: float = 20.0,
            numeric_cols: set[int] | None = None,
            total_predicate=None,
        ) -> None:
            nonlocal y
            headers = [str(v) for v in headers]
            rows = [["" if v is None else str(v) for v in row] for row in rows]
            fractions = list(widths)
            total = sum(fractions) or 1.0
            widths_px = [usable_width * value / total for value in fractions]
            numeric_cols = set(numeric_cols or set())
            body_font = QFont("Arial", font_size)
            header_font = QFont("Arial", font_size)
            header_font.setBold(True)
            line_height = QFontMetricsF(body_font).height() + 2

            def draw_header() -> None:
                nonlocal y
                header_h = 22.0
                ensure(header_h)
                x = margin
                painter.setFont(header_font)
                for idx, value in enumerate(headers):
                    rect = QRectF(x, y, widths_px[idx], header_h)
                    painter.fillRect(rect, QBrush(QColor(211, 229, 242)))
                    painter.drawRect(rect)
                    painter.drawText(rect.adjusted(3, 1, -3, -1), Qt.AlignCenter | Qt.AlignVCenter | Qt.TextWordWrap, value)
                    x += widths_px[idx]
                y += header_h

            draw_header()
            for row in rows:
                line_sets = [wrapped_lines(value, widths_px[idx], body_font) for idx, value in enumerate(row)]
                max_lines = max(len(lines) for lines in line_sets) if line_sets else 1
                row_h = max(min_height, max_lines * line_height + 5)
                if y + row_h > bottom:
                    new_page()
                    draw_header()

                is_total = bool(total_predicate(row)) if total_predicate else False
                x = margin
                painter.setFont(body_font)
                for idx, value in enumerate(row):
                    rect = QRectF(x, y, widths_px[idx], row_h)
                    if is_total:
                        painter.fillRect(rect, QBrush(QColor(226, 234, 241)))
                    painter.drawRect(rect)
                    if is_total:
                        total_font = QFont(body_font)
                        total_font.setBold(True)
                        painter.setFont(total_font)
                    else:
                        painter.setFont(body_font)
                    align = Qt.AlignRight if idx in numeric_cols else Qt.AlignLeft
                    painter.drawText(rect.adjusted(4, 2, -4, -2), align | Qt.AlignVCenter | Qt.TextWordWrap, value)
                    x += widths_px[idx]
                y += row_h

        title("KERTAS KERJA SPT TAHUNAN")
        title(f"TAHUN PAJAK {data.tahun_pajak}", 10, 20)
        table(
            ["IDENTITAS", "KETERANGAN"],
            [["Nama Wajib Pajak", data.nama_wp], ["NPWP", data.npwp]],
            [1.15, 3.85],
            font_size=8,
            min_height=20,
        )
        y += 7

        section("BUKTI POTONG / PENGHASILAN PEKERJAAN")
        bupot_rows = []
        total_bruto = total_pengurang = 0.0
        for no, item in enumerate(data.bupot_rows, start=1):
            bruto = self._money(item.bruto)
            pengurang = self._money(item.pengurang)
            total_bruto += bruto
            total_pengurang += pengurang
            bupot_rows.append([
                no, item.jenis, item.npwp_pemberi_kerja, item.no_bupot,
                self._money_text(bruto), self._money_text(pengurang), self._money_text(bruto - pengurang),
            ])
        bupot_rows.append(["", "TOTAL", "", "", self._money_text(total_bruto), self._money_text(total_pengurang), self._money_text(total_bruto - total_pengurang)])
        table(
            ["No", "Jenis", "NPWP Pemberi Kerja", "No Bupot", "Bruto", "Pengurang", "Netto"],
            bupot_rows,
            [0.38, 0.55, 1.55, 1.25, 1.0, 1.0, 1.0],
            font_size=6,
            min_height=18,
            numeric_cols={0, 4, 5, 6},
            total_predicate=lambda row: len(row) > 1 and row[1] == "TOTAL",
        )
        y += 7

        section("HARTA / SIMULASI I")
        harta_rows = []
        total_prev = total_now = 0.0
        for item in data.harta_current_rows:
            prev = self._money(item.nilai_tahun_sebelumnya)
            now = self._money(item.nilai_tahun_berjalan)
            total_prev += prev
            total_now += now
            harta_rows.append([
                item.nomor, item.kode_eform, item.kode_ct, item.nama_harta,
                item.tahun_perolehan, self._money_text(prev), self._money_text(now),
            ])
        harta_rows.append(["", "", "", "TOTAL HARTA", "", self._money_text(total_prev), self._money_text(total_now)])
        table(
            ["No", "EFORM", "CT", "Nama Harta", "Tahun", str(data.tahun_pajak - 1), str(data.tahun_pajak)],
            harta_rows,
            [0.32, 0.55, 0.52, 2.75, 0.60, 1.08, 1.08],
            font_size=6,
            min_height=18,
            numeric_cols={0, 4, 5, 6},
            total_predicate=lambda row: len(row) > 3 and row[3] == "TOTAL HARTA",
        )
        y += 7

        section("RINGKASAN PPh")
        table(
            ["Komponen", "Nilai"],
            [
                ["Status PTKP", data.status_ptkp],
                ["PTKP", f"Rp {self._money_text(data.pph_components.get('ptkp', 0))}"],
                ["PPh Terutang", f"Rp {self._money_text(data.pph_calc_result.get('pph_terutang', data.pph_components.get('pph_terutang', 0)))}"],
                ["Kredit Pajak", f"Rp {self._money_text(data.pph_components.get('kredit_pajak', 0))}"],
                ["PPh Pasal 25", f"Rp {self._money_text(data.pph_components.get('pph25', 0))}"],
            ],
            [1.25, 1.0],
            font_size=7,
            min_height=19,
            numeric_cols={1},
        )
        y += 7

        if data.analisis_result is not None:
            analysis = data.analisis_result
            section("ANALISIS PENGHASILAN VS HARTA")
            table(
                ["Komponen", "Nilai"],
                [
                    ["Total Pengeluaran", f"Rp {self._money_text(getattr(analysis, 'total_pengeluaran', 0))}"],
                    ["Penghasilan Netto", f"Rp {self._money_text(getattr(analysis, 'penghasilan_netto', 0))}"],
                    ["Selisih", f"Rp {self._money_text(getattr(analysis, 'selisih_pengeluaran_vs_penghasilan', 0))}"],
                ],
                [1.25, 1.0],
                font_size=7,
                min_height=19,
                numeric_cols={1},
            )

        draw_footer()
