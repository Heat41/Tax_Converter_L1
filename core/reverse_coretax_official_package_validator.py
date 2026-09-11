from __future__ import annotations

import hashlib
import json
from dataclasses import dataclass, field
from pathlib import Path
from typing import Dict, List, Optional
from xml.etree import ElementTree as ET

from openpyxl import load_workbook

from core.coretax_official_schema import get_official_schema


@dataclass(frozen=True)
class OfficialPackageValidationIssue:
    code: str
    severity: str
    message: str
    category: Optional[str] = None


@dataclass
class OfficialPackageValidationResult:
    package_dir: Path
    checked_files: Dict[str, Path] = field(default_factory=dict)
    issues: List[OfficialPackageValidationIssue] = field(default_factory=list)
    manifest: dict = field(default_factory=dict)

    @property
    def errors(self) -> List[OfficialPackageValidationIssue]:
        return [item for item in self.issues if item.severity == "ERROR"]

    @property
    def warnings(self) -> List[OfficialPackageValidationIssue]:
        return [item for item in self.issues if item.severity == "WARNING"]

    @property
    def ok(self) -> bool:
        return not self.errors


class OfficialCoretaxPackageValidator:
    """Stage 8D.4F - validasi paket resmi hasil Stage 8D.4E."""

    MANIFEST_VERSION = 1
    EXPECTED_EXCEL_CATEGORIES = (
        "KAS",
        "PIUTANG",
        "INVESTASI",
        "BERGERAK",
        "HTB",
        "LAINNYA",
    )
    EXPECTED_XML_CATEGORIES = (
        "KAS",
        "INVESTASI",
        "BERGERAK",
        "HTB",
    )
    EXPECTED_UNSUPPORTED_XML = (
        "PIUTANG",
        "LAINNYA",
    )

    @staticmethod
    def _sha256(path: Path) -> str:
        digest = hashlib.sha256()
        with path.open("rb") as stream:
            for chunk in iter(lambda: stream.read(1024 * 1024), b""):
                digest.update(chunk)
        return digest.hexdigest()

    @staticmethod
    def _error(
        result: OfficialPackageValidationResult,
        code: str,
        message: str,
        category: Optional[str] = None,
    ) -> None:
        result.issues.append(
            OfficialPackageValidationIssue(
                code=code,
                severity="ERROR",
                message=message,
                category=category,
            )
        )

    def _load_manifest(
        self,
        result: OfficialPackageValidationResult,
    ) -> bool:
        manifest_path = result.package_dir / "manifest.json"

        if not manifest_path.is_file():
            self._error(
                result,
                "RCX4F_001",
                "manifest.json tidak ditemukan.",
            )
            return False

        try:
            manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError) as exc:
            self._error(
                result,
                "RCX4F_002",
                f"manifest.json tidak dapat dibaca: {exc}",
            )
            return False

        if not isinstance(manifest, dict):
            self._error(
                result,
                "RCX4F_003",
                "Root manifest harus berupa object JSON.",
            )
            return False

        result.manifest = manifest
        result.checked_files["MANIFEST"] = manifest_path
        return True

    def _validate_manifest_contract(
        self,
        result: OfficialPackageValidationResult,
    ) -> None:
        manifest = result.manifest

        if manifest.get("manifest_version") != self.MANIFEST_VERSION:
            self._error(
                result,
                "RCX4F_010",
                (
                    "Versi manifest tidak didukung: "
                    f"{manifest.get('manifest_version')!r}."
                ),
            )

        if manifest.get("stage") != "8D.4E":
            self._error(
                result,
                "RCX4F_011",
                f"Stage manifest tidak sesuai: {manifest.get('stage')!r}.",
            )

        for field_name in (
            "npwp",
            "tahun_pajak",
            "revision",
            "snapshot_hash",
            "total_rows",
            "excel",
            "xml",
        ):
            if field_name not in manifest:
                self._error(
                    result,
                    "RCX4F_012",
                    f"Field manifest wajib tidak ditemukan: {field_name}",
                )

    def _validate_hash(
        self,
        result: OfficialPackageValidationResult,
        path: Path,
        expected_hash: object,
        category: str,
        code_prefix: str,
    ) -> None:
        expected = str(expected_hash or "").strip().lower()
        actual = self._sha256(path).lower()

        if not expected or actual != expected:
            self._error(
                result,
                f"{code_prefix}_HASH",
                (
                    f"SHA-256 file {path.name} tidak sama dengan manifest. "
                    f"expected={expected or '-'} actual={actual}"
                ),
                category,
            )

    def _validate_excel_file(
        self,
        result: OfficialPackageValidationResult,
        category: str,
        item: dict,
    ) -> None:
        relative_path = str(item.get("path") or "")
        path = result.package_dir / relative_path

        if not path.is_file():
            self._error(
                result,
                "RCX4F_EXCEL_MISSING",
                f"File Excel tidak ditemukan: {relative_path}",
                category,
            )
            return

        result.checked_files[f"EXCEL:{category}"] = path
        self._validate_hash(
            result,
            path,
            item.get("sha256"),
            category,
            "RCX4F_EXCEL",
        )

        schema = get_official_schema(category)

        try:
            wb = load_workbook(path, data_only=False, read_only=True)
        except Exception as exc:
            self._error(
                result,
                "RCX4F_EXCEL_OPEN",
                f"File Excel tidak dapat dibuka: {path.name}: {exc}",
                category,
            )
            return

        try:
            if schema.excel_sheet not in wb.sheetnames:
                self._error(
                    result,
                    "RCX4F_EXCEL_SHEET",
                    (
                        f"Sheet {schema.excel_sheet!r} tidak ditemukan "
                        f"pada {path.name}."
                    ),
                    category,
                )
                return

            ws = wb[schema.excel_sheet]
            actual_headers = tuple(
                str(ws.cell(3, column).value or "").strip()
                for column in range(1, len(schema.excel_headers) + 1)
            )
            expected_headers = tuple(
                str(value).strip()
                for value in schema.excel_headers
            )

            if actual_headers != expected_headers:
                self._error(
                    result,
                    "RCX4F_EXCEL_HEADER",
                    f"Header Excel {category} berubah dari kontrak resmi.",
                    category,
                )

            actual_rows = 0
            for row_index in range(4, ws.max_row + 1):
                if any(
                    ws.cell(row_index, column).value not in (None, "")
                    for column in range(1, len(schema.excel_headers) + 1)
                ):
                    actual_rows += 1

            expected_rows = int(item.get("rows") or 0)
            if actual_rows != expected_rows:
                self._error(
                    result,
                    "RCX4F_EXCEL_ROWS",
                    (
                        f"Jumlah baris Excel {category} berbeda: "
                        f"manifest={expected_rows}, actual={actual_rows}."
                    ),
                    category,
                )
        finally:
            wb.close()

    def _validate_xml_file(
        self,
        result: OfficialPackageValidationResult,
        category: str,
        item: dict,
    ) -> None:
        relative_path = str(item.get("path") or "")
        path = result.package_dir / relative_path

        if not path.is_file():
            self._error(
                result,
                "RCX4F_XML_MISSING",
                f"File XML tidak ditemukan: {relative_path}",
                category,
            )
            return

        result.checked_files[f"XML:{category}"] = path
        self._validate_hash(
            result,
            path,
            item.get("sha256"),
            category,
            "RCX4F_XML",
        )

        schema = get_official_schema(category)

        try:
            root = ET.parse(path).getroot()
        except (OSError, ET.ParseError) as exc:
            self._error(
                result,
                "RCX4F_XML_PARSE",
                f"XML tidak dapat dibaca: {path.name}: {exc}",
                category,
            )
            return

        if root.tag != schema.xml_root:
            self._error(
                result,
                "RCX4F_XML_ROOT",
                (
                    f"Root XML {category} tidak sesuai: "
                    f"expected={schema.xml_root}, actual={root.tag}."
                ),
                category,
            )
            return

        rows = root.findall(schema.xml_list)
        expected_rows = int(item.get("rows") or 0)

        if len(rows) != expected_rows:
            self._error(
                result,
                "RCX4F_XML_ROWS",
                (
                    f"Jumlah elemen XML {category} berbeda: "
                    f"manifest={expected_rows}, actual={len(rows)}."
                ),
                category,
            )

        for row_index, row in enumerate(rows, start=1):
            actual_fields = tuple(child.tag for child in list(row))
            if actual_fields != schema.xml_fields:
                self._error(
                    result,
                    "RCX4F_XML_FIELDS",
                    (
                        f"Urutan field XML {category} baris {row_index} "
                        "berubah dari kontrak resmi."
                    ),
                    category,
                )
                break

    def _validate_file_sections(
        self,
        result: OfficialPackageValidationResult,
    ) -> None:
        manifest = result.manifest
        excel = manifest.get("excel")
        xml = manifest.get("xml")

        if not isinstance(excel, dict):
            self._error(
                result,
                "RCX4F_020",
                "Bagian excel pada manifest tidak valid.",
            )
            return

        if not isinstance(xml, dict):
            self._error(
                result,
                "RCX4F_021",
                "Bagian xml pada manifest tidak valid.",
            )
            return

        excel_files = excel.get("files")
        xml_files = xml.get("files")

        if not isinstance(excel_files, dict):
            self._error(
                result,
                "RCX4F_022",
                "excel.files pada manifest tidak valid.",
            )
            excel_files = {}

        if not isinstance(xml_files, dict):
            self._error(
                result,
                "RCX4F_023",
                "xml.files pada manifest tidak valid.",
            )
            xml_files = {}

        if tuple(excel_files.keys()) != self.EXPECTED_EXCEL_CATEGORIES:
            self._error(
                result,
                "RCX4F_024",
                (
                    "Daftar kategori Excel tidak lengkap/berubah: "
                    f"{list(excel_files)}"
                ),
            )

        if tuple(xml_files.keys()) != self.EXPECTED_XML_CATEGORIES:
            self._error(
                result,
                "RCX4F_025",
                (
                    "Daftar kategori XML tidak lengkap/berubah: "
                    f"{list(xml_files)}"
                ),
            )

        unsupported = tuple(xml.get("unsupported_categories") or ())
        if unsupported != self.EXPECTED_UNSUPPORTED_XML:
            self._error(
                result,
                "RCX4F_026",
                (
                    "Kategori XML unsupported berubah: "
                    f"{list(unsupported)}"
                ),
            )

        if int(excel.get("file_count") or 0) != len(excel_files):
            self._error(
                result,
                "RCX4F_027",
                "excel.file_count tidak sama dengan jumlah file di manifest.",
            )

        if int(xml.get("file_count") or 0) != len(xml_files):
            self._error(
                result,
                "RCX4F_028",
                "xml.file_count tidak sama dengan jumlah file di manifest.",
            )

        for category in self.EXPECTED_EXCEL_CATEGORIES:
            item = excel_files.get(category)
            if not isinstance(item, dict):
                self._error(
                    result,
                    "RCX4F_EXCEL_ENTRY",
                    f"Entry Excel {category} tidak tersedia.",
                    category,
                )
                continue
            self._validate_excel_file(result, category, item)

        for category in self.EXPECTED_XML_CATEGORIES:
            item = xml_files.get(category)
            if not isinstance(item, dict):
                self._error(
                    result,
                    "RCX4F_XML_ENTRY",
                    f"Entry XML {category} tidak tersedia.",
                    category,
                )
                continue
            self._validate_xml_file(result, category, item)

    def validate(
        self,
        package_dir: str | Path,
    ) -> OfficialPackageValidationResult:
        result = OfficialPackageValidationResult(
            package_dir=Path(package_dir)
        )

        if not result.package_dir.is_dir():
            self._error(
                result,
                "RCX4F_000",
                f"Folder paket tidak ditemukan: {result.package_dir}",
            )
            return result

        if not self._load_manifest(result):
            return result

        self._validate_manifest_contract(result)
        self._validate_file_sections(result)

        if result.ok:
            result.issues.append(
                OfficialPackageValidationIssue(
                    "RCX4F_INFO",
                    "INFO",
                    (
                        "Paket resmi Coretax valid: manifest, hash, "
                        "Excel, dan XML konsisten."
                    ),
                )
            )

        return result
