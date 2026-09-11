from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import Dict, List, Optional, Sequence, Tuple

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from core.reverse_coretax_mapping import (
    CATEGORY_ORDER,
    ReverseCoretaxMappingService,
    ReverseCoretaxPackage,
    ReverseCoretaxRow,
)


@dataclass(frozen=True)
class ReverseCoretaxExportIssue:
    code: str
    severity: str
    message: str


@dataclass
class ReverseCoretaxExcelExportResult:
    output_dir: Path
    files: Dict[str, Path] = field(default_factory=dict)
    row_counts: Dict[str, int] = field(default_factory=dict)
    issues: List[ReverseCoretaxExportIssue] = field(default_factory=list)

    @property
    def errors(self) -> List[ReverseCoretaxExportIssue]:
        return [item for item in self.issues if item.severity == "ERROR"]

    @property
    def ok(self) -> bool:
        return not self.errors and len(self.files) == 6


class ReverseCoretaxExcelExporter:
    """Stage 8D.2 - ekspor 6 file Excel L-1 dari snapshot FINAL.

    Schema fisik saat ini mengikuti kontrak kolom importer internal yang sudah
    digunakan proyek. Ini adalah baseline reverse-export yang dapat diuji round
    trip. Sebelum diberi label "siap upload Coretax", schema tetap harus
    dikalibrasi terhadap contoh/template import Coretax resmi yang digunakan
    kantor.
    """

    CATEGORY_CONFIG: Dict[str, dict] = {
        "KAS": {
            "filename": "01_Kas_Setara_Kas.xlsx",
            "sheet": "Kas",
            "headers": (
                "kode_harta",
                "nama_harta",
                "tahun_perolehan",
                "harga_perolehan",
                "keterangan",
            ),
        },
        "PIUTANG": {
            "filename": "02_Piutang.xlsx",
            "sheet": "Piutang",
            "headers": (
                "kode_harta",
                "nama_peminjam",
                "npwp_peminjam",
                "tahun_perolehan",
                "harga_perolehan",
                "keterangan",
            ),
        },
        "INVESTASI": {
            "filename": "03_Investasi.xlsx",
            "sheet": "Investasi",
            "headers": (
                "kode_harta",
                "nama_harta",
                "penerbit_saham",
                "tahun_perolehan",
                "harga_perolehan",
                "keterangan",
            ),
        },
        "BERGERAK": {
            "filename": "04_Harta_Bergerak.xlsx",
            "sheet": "Harta Bergerak",
            "headers": (
                "kode_harta",
                "nama_harta",
                "merek_type",
                "tahun_perolehan",
                "harga_perolehan",
                "keterangan",
            ),
        },
        "HTB": {
            "filename": "05_Harta_Tidak_Bergerak.xlsx",
            "sheet": "Harta Tidak Bergerak",
            "headers": (
                "kode_harta",
                "jenis_harta",
                "lokasi_alamat",
                "tahun_perolehan",
                "harga_perolehan",
                "keterangan",
            ),
        },
        "LAINNYA": {
            "filename": "06_Harta_Lainnya.xlsx",
            "sheet": "Harta Lainnya",
            "headers": (
                "kode_harta",
                "nama_harta",
                "tahun_perolehan",
                "harga_perolehan",
                "keterangan",
            ),
        },
    }

    HEADER_FILL = "1F4E78"
    HEADER_FONT = "FFFFFF"

    @staticmethod
    def _join_details(*values: object) -> str:
        return "; ".join(
            str(value).strip()
            for value in values
            if str(value or "").strip() not in {"", "-"}
        )

    @classmethod
    def _row_values(cls, row: ReverseCoretaxRow) -> Tuple[object, ...]:
        category = row.kategori
        detail = row.nomor_akun_keterangan
        owner = row.atas_nama
        bank = row.nama_bank

        if category == "KAS":
            return (
                row.kode_harta,
                row.nama_harta or bank or "Kas / Setara Kas",
                row.tahun_perolehan,
                row.nilai,
                cls._join_details(detail, owner, bank),
            )
        if category == "PIUTANG":
            return (
                row.kode_harta,
                owner or row.nama_harta,
                detail,
                row.tahun_perolehan,
                row.nilai,
                cls._join_details(row.nama_harta),
            )
        if category == "INVESTASI":
            return (
                row.kode_harta,
                row.nama_harta,
                owner or bank,
                row.tahun_perolehan,
                row.nilai,
                cls._join_details(detail),
            )
        if category == "BERGERAK":
            return (
                row.kode_harta,
                row.nama_harta,
                detail,
                row.tahun_perolehan,
                row.nilai,
                cls._join_details(owner),
            )
        if category == "HTB":
            return (
                row.kode_harta,
                row.nama_harta,
                detail,
                row.tahun_perolehan,
                row.nilai,
                cls._join_details(owner),
            )
        if category == "LAINNYA":
            return (
                row.kode_harta,
                row.nama_harta,
                row.tahun_perolehan,
                row.nilai,
                cls._join_details(detail, owner, bank),
            )
        raise ValueError(f"Kategori reverse Coretax tidak dikenal: {category}")

    @classmethod
    def _style_sheet(cls, ws, headers: Sequence[str], row_count: int) -> None:
        for cell in ws[1]:
            cell.font = Font(bold=True, color=cls.HEADER_FONT)
            cell.fill = PatternFill("solid", fgColor=cls.HEADER_FILL)
            cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(1, row_count + 1)}"

        for column_index, header in enumerate(headers, start=1):
            max_len = len(header)
            for row_index in range(2, row_count + 2):
                value = ws.cell(row_index, column_index).value
                max_len = max(max_len, len(str(value or "")))
            ws.column_dimensions[get_column_letter(column_index)].width = min(
                max(max_len + 2, 12),
                38,
            )

        for row_index in range(2, row_count + 2):
            for column_index, header in enumerate(headers, start=1):
                cell = ws.cell(row_index, column_index)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if header == "harga_perolehan":
                    cell.number_format = "#,##0"

    def _write_category(
        self,
        package: ReverseCoretaxPackage,
        category: str,
        output_dir: Path,
    ) -> Path:
        config = self.CATEGORY_CONFIG[category]
        headers = config["headers"]
        rows = package.rows_by_category.get(category, [])

        wb = Workbook()
        ws = wb.active
        ws.title = config["sheet"]
        ws.append(list(headers))
        for row in rows:
            ws.append(list(self._row_values(row)))

        self._style_sheet(ws, headers, len(rows))

        target = output_dir / config["filename"]
        wb.save(target)
        return target

    def export_package(
        self,
        package: ReverseCoretaxPackage,
        output_dir: str | Path,
    ) -> ReverseCoretaxExcelExportResult:
        target_dir = Path(output_dir)
        result = ReverseCoretaxExcelExportResult(output_dir=target_dir)

        if not package.can_export:
            result.issues.append(
                ReverseCoretaxExportIssue(
                    "RCX_001",
                    "ERROR",
                    "Paket reverse Coretax belum valid untuk diekspor.",
                )
            )
            return result

        target_dir.mkdir(parents=True, exist_ok=True)

        for category in CATEGORY_ORDER:
            target = self._write_category(package, category, target_dir)
            result.files[category] = target
            result.row_counts[category] = len(package.rows_by_category.get(category, []))

        result.issues.append(
            ReverseCoretaxExportIssue(
                "RCX_INFO",
                "INFO",
                "Enam file Excel berhasil dibuat dari snapshot FINAL. "
                "Schema mengikuti kontrak importer internal dan belum diberi status "
                "template Coretax resmi sampai dikalibrasi dengan file contoh DJP/kantor.",
            )
        )
        return result

    def export_active_final(
        self,
        npwp: str,
        tahun_pajak: int,
        output_dir: str | Path,
        *,
        db_path: Optional[str | Path] = None,
    ) -> ReverseCoretaxExcelExportResult:
        package = ReverseCoretaxMappingService(db_path=db_path).build_active_final(
            npwp,
            tahun_pajak,
        )
        result = self.export_package(package, output_dir)
        for issue in package.issues:
            if issue.severity in {"ERROR", "WARNING"}:
                result.issues.append(
                    ReverseCoretaxExportIssue(
                        issue.code,
                        issue.severity,
                        issue.message,
                    )
                )
        return result
